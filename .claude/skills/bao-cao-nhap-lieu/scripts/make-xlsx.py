"""Xuất bản EXCEL của báo cáo tình trạng nhập liệu — cùng số liệu với bộ ảnh PNG.

Dùng chung `collect()` của make-report.py nên hai bản không bao giờ lệch nhau: sửa luật ở
collect.sql là cả ảnh lẫn Excel đổi theo.

Khác ảnh ở chỗ Excel liệt kê ĐỦ MỌI ĐƠN VỊ (kể cả đơn vị nộp đủ) và có cột "Nhóm" để lọc — ảnh chỉ
nêu đơn vị có vấn đề cho gọn khi gửi Zalo.

Chạy:
    uv run --directory apps/api python .claude/skills/bao-cao-nhap-lieu/scripts/make-xlsx.py \
        --days 18 --until 2026-08-10 --purchase-days 222 --out plans/visuals/2026-08-10
"""

from __future__ import annotations

import argparse
import importlib.util
import pathlib
from datetime import date, timedelta

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Tên file có dấu gạch ngang nên không import thẳng được — nạp theo đường dẫn.
_spec = importlib.util.spec_from_file_location(
    'make_report', pathlib.Path(__file__).resolve().parent / 'make-report.py')
_mr = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_mr)
GRADE_LABEL, PLAN_COLS, ROOT, collect = _mr.GRADE_LABEL, _mr.PLAN_COLS, _mr.ROOT, _mr.collect

HDR_FILL = PatternFill('solid', fgColor='1E7A46')
GRP_FILL = PatternFill('solid', fgColor='E2EFDA')
BAD_FILL = PatternFill('solid', fgColor='FFC7CE')
NA_FILL = PatternFill('solid', fgColor='EDEDED')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def sheet(wb, title: str, intro: str, hdr: list[str], widths: list[int]):
    ws = wb.create_sheet(title) if wb.sheetnames != ['Sheet'] else wb.active
    ws.title = title
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = intro
    ws['A2'].font = Font(size=10, italic=True)
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(4, i, h)
        c.font = Font(bold=True, size=10, color='FFFFFF')
        c.fill = HDR_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BOX
    ws.row_dimensions[4].height = 42
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C5'
    return ws


def put(ws, r: int, vals: list, bad_cols: tuple = (), group_row: bool = False):
    for i, v in enumerate(vals, start=1):
        c = ws.cell(r, i, v)
        c.border = BOX
        if group_row:
            c.font = Font(bold=True)
            c.fill = GRP_FILL
        elif i in bad_cols:
            c.fill = BAD_FILL
        if isinstance(v, str) and v == 'không áp dụng':
            c.fill = NA_FILL
        if i > 2 and isinstance(v, str) and len(v) > 40:
            c.alignment = Alignment(wrap_text=True, vertical='center')


def dmy(d: str) -> str:
    # Hợp đồng có thể chưa có ngày nào (chưa ký, chưa giao) → để trống thay vì in '//'.
    return f'{d[8:10]}/{d[5:7]}/{d[:4]}' if d else ''


def build(g: dict, days: int, until: date, pdays: int, pg: dict | None) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ky = f'{(until - timedelta(days=days - 1)).strftime("%d/%m")} – {until.strftime("%d/%m/%Y")}'

    # ── A · Tình trạng nộp (mọi đơn vị) ──────────────────────────────────────
    ws = sheet(wb, 'A · Tình trạng nộp',
               f'Kỳ {ky}. Số = SỐ NGÀY đã nộp trong kỳ; 0 = chưa nhập; "không áp dụng" = đơn vị '
               f'không được giao kế hoạch thu mua nên không phải nộp biểu Thu mua.',
               ['#', 'Đơn vị', 'Nhóm', 'Báo cáo thu mua', 'Tiêu thụ – Tồn kho'], [5, 42, 20, 16, 16])
    miss = lambda v: v != '-' and int(v) == 0      # noqa: E731
    r = 4
    for i, (name, pur, con) in enumerate(g['A'], 1):
        cells = [pur, con]
        if all(miss(v) or v == '-' for v in cells):
            grp = 'Không nộp gì trong kỳ'
        elif any(miss(v) for v in cells):
            grp = 'Thiếu một phần'
        else:
            grp = 'Đã nộp đủ'
        r += 1
        show = ['không áp dụng' if v == '-' else int(v) for v in cells]
        bad = tuple(4 + j for j, v in enumerate(cells) if miss(v))
        put(ws, r, [i, name, grp, *show], bad_cols=bad)
    ws.auto_filter.ref = f'A4:E{r}'

    # ── A2 · Thiếu đơn giá ───────────────────────────────────────────────────
    ws = sheet(wb, 'A2 · Thiếu đơn giá',
               'Ngày CÓ tổ chức thu mua nhưng ô đơn giá còn trống (ngày tích "không tổ chức thu '
               'mua" đã loại ra).', ['#', 'Đơn vị', 'Ngày'], [5, 42, 14])
    for i, row in enumerate(g['D'], 1):
        put(ws, 4 + i, [i, row[0], dmy(row[1])], bad_cols=(3,))
    ws.auto_filter.ref = f'A4:C{4 + len(g["D"])}'

    # ── B · Sai đơn vị tính ──────────────────────────────────────────────────
    ws = sheet(wb, 'B · Sai đơn vị tính',
               'Không giới hạn kỳ — lỗi còn tồn trên hệ thống thì còn phải sửa. Giá mủ nguyên liệu '
               'phải nhập ĐỒNG/ĐỘ (mặt bằng 100–1.500); giá bán ở HỢP ĐỒNG TIÊU THỤ phải nhập '
               'TRIỆU ĐỒNG/TẤN (40–70) hoặc USD/TẤN (1.400–2.200).',
               ['#', 'Đơn vị', 'Chỉ tiêu', 'Giá lớn nhất đang lưu', 'Số ô sai',
                'Từ ngày', 'Đến ngày', 'Ghi chú'], [5, 42, 22, 20, 10, 12, 12, 30])
    r = 4
    r += 1
    put(ws, r, ['', 'GIÁ MỦ NGUYÊN LIỆU (phải là đồng/độ)', '', '', '', '', '', ''], group_row=True)
    for i, row in enumerate(g['B'], 1):
        r += 1
        put(ws, r, [i, row[0], GRADE_LABEL.get(row[1], row[1]), int(row[2]), int(row[3]),
                    None, None, 'Nhập nhầm đồng/kg hoặc đồng/tấn'], bad_cols=(4,))
    r += 1
    put(ws, r, ['', 'GIÁ BÁN Ở HỢP ĐỒNG TIÊU THỤ (phải là triệu đồng/tấn)', '', '', '', '', '', ''],
        group_row=True)
    for i, row in enumerate(g['C'], 1):
        r += 1
        # C: đơn vị|số dòng|từ ngày|đến ngày|giá lớn nhất|tiền|thiếu tỷ giá|mã hợp đồng
        note = f'HĐ {row[7]}' if len(row) > 7 and row[7] else ''
        note = ' · '.join(filter(None, [note, f'loại tiền {row[5]}',
                                        'có dòng thiếu tỷ giá' if row[6] == 't' else '']))
        put(ws, r, [i, row[0], 'Giá bán', float(row[4]), int(row[1]),
                    dmy(row[2]), dmy(row[3]), note], bad_cols=(4,))

    # ── C · Tồn kho theo ngày ────────────────────────────────────────────────
    dates = [(until - timedelta(days=i)).isoformat() for i in range(days - 1, -1, -1)]
    ws = sheet(wb, 'C · Tồn kho theo ngày',
               f'Kỳ {ky} ({days} ngày). Cột "Ngày còn thiếu" liệt kê đúng những ngày đơn vị chưa '
               f'nộp biểu Tồn kho.',
               ['#', 'Đơn vị', 'Khu vực', 'Đã nộp', 'Còn thiếu', 'Tỷ lệ nộp',
                'Ngày còn thiếu'], [5, 42, 20, 10, 10, 11, 80])
    r = 4
    for i, (name, region, done_csv) in enumerate(g['E'], 1):
        done = set(filter(None, done_csv.split(',')))
        gone = [d for d in dates if d not in done]
        r += 1
        put(ws, r, [i, name, region or '—', len(dates) - len(gone), len(gone),
                    round((len(dates) - len(gone)) / len(dates), 4),
                    ', '.join(f'{d[8:10]}/{d[5:7]}' for d in gone)],
            bad_cols=(5,) if gone else ())
        ws.cell(r, 6).number_format = '0%'
    ws.auto_filter.ref = f'A4:G{r}'

    # ── D · Thu mua theo tháng ───────────────────────────────────────────────
    rows_f = (pg or g)['F']
    p_until, p_days = until, pdays
    start = p_until - timedelta(days=p_days - 1)
    per_month: dict[str, int] = {}
    for i in range(p_days):
        m = (start + timedelta(days=i)).strftime('%Y-%m')
        per_month[m] = per_month.get(m, 0) + 1
    months = sorted(per_month)
    ws = sheet(wb, 'D · Thu mua theo tháng',
               f'Kỳ {start.strftime("%d/%m/%Y")} – {p_until.strftime("%d/%m/%Y")} ({p_days} ngày). '
               f'Ô tháng = số ngày đã nộp / số ngày của tháng đó TRONG KỲ. Ngày tích "không tổ chức '
               f'thu mua" vẫn tính là ĐÃ NỘP. Đơn vị không được giao kế hoạch thu mua không có tên.',
               ['#', 'Đơn vị', 'Khu vực', 'Đã nộp', 'Tổng ngày', 'Tỷ lệ nộp',
                *[f'T{int(m[5:7])}/{per_month[m]}' for m in months]],
               [5, 42, 20, 10, 10, 11, *[9] * len(months)])
    r = 4
    for name, region, applies, csv in rows_f:
        if applies != 't':
            continue
        got = {}
        for part in filter(None, csv.split(',')):
            mm, cc = part.split(':')
            got[mm] = int(cc)
        total = sum(got.get(m, 0) for m in months)
        r += 1
        put(ws, r, [r - 4, name, region or '—', total, p_days, round(total / p_days, 4),
                    *[got.get(m, 0) for m in months]],
            bad_cols=tuple(7 + j for j, m in enumerate(months) if got.get(m, 0) < per_month[m]))
        ws.cell(r, 6).number_format = '0%'
    ws.auto_filter.ref = f'A4:{get_column_letter(6 + len(months))}{r}'

    # ── E · Kế hoạch năm ─────────────────────────────────────────────────────
    plan_cols = [c.replace('<br>', ' ') for c in PLAN_COLS]
    ws = sheet(wb, 'E · Kế hoạch năm',
               f'Kế hoạch năm {until.year}. "Chưa khai" = ô còn TRỐNG; khai số 0 vẫn tính là ĐÃ '
               f'KHAI (nghĩa là "không có").',
               ['#', 'Đơn vị', 'Khu vực', 'Đã khai', *plan_cols],
               [5, 42, 20, 10, *[16] * len(plan_cols)])
    r = 4
    for i, row in enumerate(g['G'], 1):
        ok = [c == 't' for c in row[2:7]]
        r += 1
        put(ws, r, [i, row[0], row[1] or '—', f'{sum(ok)}/5',
                    *['Đã khai' if v else 'Chưa khai' for v in ok]],
            bad_cols=tuple(5 + j for j, v in enumerate(ok) if not v))
    ws.auto_filter.ref = f'A4:{get_column_letter(4 + len(plan_cols))}{r}'
    return wb


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--days', type=int, default=7, help='số ngày của kỳ xét (ảnh A, C)')
    ap.add_argument('--purchase-days', type=int, default=None,
                    help='số ngày của kỳ THU MUA (sheet D) nếu khác --days, vd 222 = từ đầu năm')
    ap.add_argument('--until', default=(date.today() - timedelta(days=1)).isoformat())
    ap.add_argument('--out', default='plans/visuals')
    ap.add_argument('--local', action='store_true')
    args = ap.parse_args()

    until = date.fromisoformat(args.until)
    g = collect(args.days, args.local, until)
    pdays = args.purchase_days or args.days
    # Kỳ thu mua khác kỳ chung thì phải hỏi lại DB đúng kỳ đó, không suy từ bộ số cũ.
    pg = collect(pdays, args.local, until) if pdays != args.days else None

    out_dir = pathlib.Path(args.out)
    if not out_dir.is_absolute():
        out_dir = ROOT / out_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f'tinh-trang-nhap-lieu-{until.strftime("%d-%m-%Y")}.xlsx'
    build(g, args.days, until, pdays, pg).save(path)
    print('đã tạo', path)


if __name__ == '__main__':
    main()
