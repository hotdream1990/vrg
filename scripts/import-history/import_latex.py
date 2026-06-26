"""Import giá mủ nước vào fact_price (source='vrg').

- Sheet 'GIÁ THU MUA MỦ NC'  -> price_type='purchase', grade = tên công ty.
- Sheet 'GIÁ MỦ NC THEO KV'   -> price_type='region',   grade = tên khu vực.

Định dạng wide: 1 hàng tiêu đề là ngày (dd/mm, năm ngầm KHÔNG tin -> dựng lại),
mỗi hàng dữ liệu là 1 đơn vị; giá có thể là khoảng '407-412' -> lấy trung điểm.
Mặc định đọc file master (đã kiểm chứng độ phủ). Đơn vị: đồng/độ TSC.

Chạy:
  uv run --with openpyxl --with "psycopg[binary]" python import_latex.py [--file X] [--dry-run]
"""
from __future__ import annotations

import argparse

import openpyxl

import _lib

DEFAULT_FILE = str(_lib.DOCS / "bieu-mau/Tâm/Mẫu file lấy giá các sàn, giá mủ nước...xlsx")
# (sheet, hàng tiêu đề ngày, cột tên đơn vị, cột ngày bắt đầu, price_type)
SHEETS = [
    ("GIÁ THU MUA MỦ NC", 2, 2, 4, "purchase"),
    ("GIÁ MỦ NC THEO KV", 3, 2, 3, "region"),
]


def _date_columns(ws, hdr_row: int, start_col: int, from_year: int) -> dict[int, str]:
    """{cột Excel -> 'YYYY-MM-DD'} đã dựng lại năm, loại ô ngày lạc, lọc >= from_year.

    Cột đi trái->phải nên ngày phải tăng dần; ô có ngày NHỎ hơn ngày lớn nhất đã thấy
    là ô lạc/typo trong nguồn -> bỏ.
    """
    header = list(ws.iter_rows(min_row=hdr_row, max_row=hdr_row, values_only=True))[0]
    cols, dms = [], []
    for i, v in enumerate(header, 1):
        if i < start_col:
            continue
        dm = _lib.daymonth(v)
        if dm:
            cols.append(i)
            dms.append(dm)
    years = _lib.reconstruct_years(dms, last_year=2026)
    out: dict[int, str] = {}
    running = ""
    for col, (d, mo), yr in zip(cols, dms, years):
        ds = f"{yr:04d}-{mo:02d}-{d:02d}"
        if ds < running:  # ngày lùi so với mốc lớn nhất -> ô lạc, bỏ
            continue
        running = ds
        if yr >= from_year:
            out[col] = ds
    return out


def run(path: str, from_year: int, dry: bool) -> int:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    total = 0
    for sheet, hdr_row, name_col, start_col, ptype in SHEETS:
        if sheet not in wb.sheetnames:
            print(f"[latex] bỏ qua (không có sheet) '{sheet}'")
            continue
        ws = wb[sheet]
        date_cols = _date_columns(ws, hdr_row, start_col, from_year)
        rows: list[dict] = []
        for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
            name = r[name_col - 1] if name_col - 1 < len(r) else None
            if not name or not str(name).strip():
                continue
            grade = str(name).strip()
            for col, as_of in date_cols.items():
                v = _lib.parse_number(r[col - 1]) if col - 1 < len(r) else None
                if v is None:
                    continue
                rows.append(dict(as_of=as_of, source="vrg", grade=grade, contract="",
                                 price_type=ptype, price=v, currency="VND",
                                 unit="đồng/độ TSC", source_ts=None))
        n = _lib.upsert(rows, dry_run=dry)
        total += n
        tag = "(dry-run) " if dry else ""
        print(f"[latex] '{sheet}' ({ptype}): {n} bản ghi {tag}| {len(date_cols)} ngày >= {from_year}")
    wb.close()
    return total


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", default=DEFAULT_FILE)
    ap.add_argument("--from-year", type=int, default=_lib.MIN_YEAR)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    run(a.file, a.from_year, a.dry_run)
