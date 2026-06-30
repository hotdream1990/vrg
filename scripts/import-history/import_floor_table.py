"""Import Giá sàn Tập đoàn từ bảng theo năm 'Giá sàn 2024-2026.xlsx' → vrg_floor_price.

Nguồn ĐẦY ĐỦ hơn 'Bảng tính giá' cũ (77 lần ban hành vs 30). Mỗi sheet = 1 năm:
  STT | Ngày | <khối USD/tấn theo grade> | <khối VNĐ/tấn theo grade>
- 2 khối có thứ tự grade KHÁC nhau (vd LATEX cột 13 ở USD nhưng cột 29 ở VNĐ)
  → map theo TÊN header từng khối, KHÔNG theo offset cố định.
- Bỏ dòng trung bình (ô 'Ngày' không phải date) + grade 'SKIM.B' (DB không có).
- Ghi đè toàn bộ vrg_floor_price; lan đánh số toàn cục theo ngày (chronological);
  GIỮ ngày đã có trong DB mà file không có (không mất dữ liệu).

Thay cho import_floor.py (nguồn cũ). Chạy:
  uv run --with openpyxl --with "psycopg[binary]" python import_floor_table.py [--dry-run]
"""
from __future__ import annotations

import argparse
import datetime

import openpyxl

import _lib

FILE = _lib.DOCS / "bieu-mau" / "Giá sàn 2024-2026.xlsx"

# tên grade trong file -> tên canonical trong DB (gồm cả biến thể có/không dấu cách)
CANON = {
    "CV50": "SVR CV 50", "CV60": "SVR CV60", "SVRL": "SVR L",
    "SVR3L Mix": "SVR 3L Mix", "SVR 3L Mix": "SVR 3L Mix", "SVR3L": "SVR 3L",
    "SVR5S": "SVR 5S", "SVR5": "SVR 5", "SVR10 Mix": "SVR 10 Mix", "SVR10": "SVR 10",
    "SVR20": "SVR 20", "LATEX": "LATEX", "RSS3": "RSS 3", "RSS1": "RSS 1",
}


def _canon(g) -> str | None:
    return CANON.get(" ".join(str(g).split())) if g else None


def _blocks(hdr) -> tuple[list[int], list[int]]:
    """Tách header thành [cột khối USD], [cột khối VNĐ] theo khoảng trống giữa 2 khối."""
    cols = [i for i, c in enumerate(hdr) if c and " ".join(str(c).split()) not in ("STT", "Ngày")]
    groups, cur = [], [cols[0]]
    for a, b in zip(cols, cols[1:]):
        if b - a > 1:
            groups.append(cur)
            cur = [b]
        else:
            cur.append(b)
    groups.append(cur)
    return groups[0], groups[1]


def parse(path: str) -> tuple[dict, set]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    recs: dict[tuple[str, str], tuple] = {}
    unmapped: set[str] = set()
    for sh in wb.sheetnames:
        rows = list(wb[sh].iter_rows(values_only=True))
        hdr = rows[1]
        usd_b, vnd_b = _blocks(hdr)
        usd_col = {_canon(hdr[c]): c for c in usd_b if _canon(hdr[c])}
        vnd_col = {_canon(hdr[c]): c for c in vnd_b if _canon(hdr[c])}
        for c in usd_b + vnd_b:
            if hdr[c] and not _canon(hdr[c]):
                unmapped.add(str(hdr[c]).strip())
        for r in rows[2:]:
            if len(r) < 2 or not isinstance(r[1], datetime.datetime):
                continue  # bỏ dòng trung bình
            d = r[1].date().isoformat()
            for g in usd_col.keys() & vnd_col.keys():
                fob = r[usd_col[g]]
                if fob is None:
                    continue
                vnd = r[vnd_col[g]]
                recs[(d, g)] = (float(fob), float(vnd) if vnd is not None else None)
    return recs, unmapped


def run(from_year: int, dry: bool) -> int:
    import psycopg

    recs, unmapped = parse(str(FILE))
    recs = {k: v for k, v in recs.items() if int(k[0][:4]) >= from_year}
    # giữ ngày đã có trong DB mà file không có
    with psycopg.connect(_lib.get_dsn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT as_of::text, grade, fob_usd, domestic_vnd FROM vrg_floor_price")
        db_rows = cur.fetchall()
    file_dates = {d for d, _ in recs}
    for d, g, fob, vnd in db_rows:
        if d not in file_dates:
            recs.setdefault((d, g), (fob, vnd))
    dates = sorted({d for d, _ in recs})
    lan_of = {d: i + 1 for i, d in enumerate(dates)}  # lan toàn cục theo ngày
    payload = [{"lan": lan_of[d], "as_of": d, "grade": g, "fob_usd": fob, "domestic_vnd": vnd}
               for (d, g), (fob, vnd) in recs.items()]
    print(f"[floor-table] {len(dates)} lần ban hành, {len(payload)} dòng | {dates[0]} -> {dates[-1]}"
          + (f" | bỏ grade: {unmapped}" if unmapped else "")
          + (" (dry-run)" if dry else ""))
    if dry:
        return len(payload)
    with psycopg.connect(_lib.get_dsn()) as conn, conn.cursor() as cur:
        cur.execute("DELETE FROM vrg_floor_price")
        cur.executemany(
            "INSERT INTO vrg_floor_price (lan, as_of, grade, fob_usd, domestic_vnd) "
            "VALUES (%(lan)s, %(as_of)s, %(grade)s, %(fob_usd)s, %(domestic_vnd)s)", payload)
        conn.commit()
    print(f"[floor-table] đã ghi {len(payload)} dòng vào vrg_floor_price")
    return len(payload)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--from-year", type=int, default=_lib.MIN_YEAR)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    run(a.from_year, a.dry_run)
