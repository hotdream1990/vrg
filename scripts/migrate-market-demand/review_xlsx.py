"""Bảng đối chiếu (Excel) cho chủ dự án duyệt TRƯỚC khi ghi dữ liệu chuyển đổi lên prod."""
from __future__ import annotations

from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

UNIT = {"ton": "tấn", "container": "container"}
PRICE_UNIT = {"VND": "triệu đ/tấn", "USD": "USD/tấn"}
HEAD = PatternFill("solid", fgColor="0A9E48")
FLAG = PatternFill("solid", fgColor="FEF3C7")
MANUAL = PatternFill("solid", fgColor="EEF6FF")

COLUMNS = [
    ("STT", 5), ("Đơn vị", 28), ("Ngày nhận", 11), ("Khách hàng", 30), ("Chủng loại", 18),
    ("Số lượng", 10), ("ĐVT", 9), ("Đơn giá", 10), ("Đơn vị giá", 11), ("Giao tại", 18),
    ("Thời gian giao", 22), ("Kết quả", 32), ("Cách chuyển", 10), ("Cần anh xem", 40),
    ("Nội dung gốc", 70),
]


def _dmy(iso: str | None) -> str:
    return f"{iso[8:10]}/{iso[5:7]}/{iso[:4]}" if iso else ""


def _row(n: int, it: dict) -> list:
    return [
        n, it["company"], _dmy(it["as_of"]), it["customer"], it["grade"],
        it["qty"], UNIT[it["qty_unit"]] if it["qty"] is not None else "",
        it["price"], PRICE_UNIT[it["currency"]] if it["price"] is not None else "",
        it["delivery_place"], it["delivery_time"], it["result"], it["method"], it["flag"],
        " ".join(it["original"].split()),
    ]


def write_review(items: list[dict], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Đối chiếu"
    for c, (title, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "E2"
    for n, it in enumerate(sorted(items, key=lambda x: (x["company"], x["as_of"])), 1):
        ws.append(_row(n, it))
        fill = FLAG if it["flag"] else MANUAL if it["method"] == "Soạn tay" else None
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=n + 1, column=c)
            cell.alignment = Alignment(vertical="top", wrap_text=c in (4, 12, 13, 15, 16))
            if fill:
                cell.fill = fill
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(items) + 1}"

    s = wb.create_sheet("Tổng quan")
    lines = [
        ("Tổng số phiếu sau chuyển đổi", len(items)),
        ("Tách tự động", sum(i["method"] == "Tự động" for i in items)),
        ("Soạn tay (nền xanh nhạt)", sum(i["method"] == "Soạn tay" for i in items)),
        ("Cần anh xem (nền vàng)", sum(bool(i["flag"]) for i in items)),
        ("", ""),
        ("Đã có kết quả (đã ký)", sum(bool(i["result"]) for i in items)),
        ("", ""),
        *[(f"Đơn vị · {k}", v) for k, v in sorted(Counter(i["company"] for i in items).items())],
        ("", ""),
        ("Ghi chú", "Nguyên văn cũ được giữ trong ô Ghi chú của từng phiếu. Bản cũ không ghi kết quả "
                    "thì để trống ô Kết quả — đơn vị cập nhật lại được bất cứ lúc nào."),
    ]
    for label, value in lines:
        s.append([label, value])
    s.column_dimensions["A"].width, s.column_dimensions["B"].width = 42, 90
    wb.save(path)
