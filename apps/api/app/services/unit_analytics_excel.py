"""Xuất Excel cho các bảng Thống kê (thu mua · tiêu thụ · tồn kho) — cột động theo bộ lọc.

Khác `unit_period_excel` (bám cứng mẫu Biểu (1)/(2)), file này dựng bảng từ danh sách cột
truyền vào nên dùng lại được cho mọi cách nhóm (đơn vị · khu vực · chủng loại · ngày).
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HEAD_FILL = PatternFill("solid", fgColor="D9E7D5")
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="9AA5A0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NUM = "#,##0.00"

Col = tuple[str, str, str]        # (khoá dữ liệu, tiêu đề, đơn vị tính)

PURCHASE_COLS: list[Col] = [
    # Mủ nguyên liệu khai theo QUY KHÔ, thành phẩm theo số thực mua — ghi vào đơn vị tính để file
    # rời khỏi hệ thống vẫn tự nói được nó là số gì.
    ("qty_latex", "Sản lượng mủ nước", "tấn quy khô"),
    ("qty_cup", "Sản lượng mủ chén", "tấn quy khô"),
    # Tổng mủ nguyên liệu + kế hoạch đứng liền nhau, RỒI mới tới thành phẩm mua ngoài — giữ đúng
    # thứ tự của màn Thống kê để người đối chiếu file với màn hình không phải dò cột.
    ("qty_material", "Tổng mủ nguyên liệu (mủ nước + chén)", "tấn"),
    ("qty_finished", "Sản lượng thành phẩm", "tấn"),
    ("qty_total", "Tổng sản lượng (mủ NL + thành phẩm)", "tấn"),
    ("price_latex_avg", "Đơn giá BQ mủ nước", "đồng/độ"),
    ("price_cup_avg", "Đơn giá BQ mủ chén", "đồng/độ"),
    ("price_finished_avg", "Đơn giá BQ thành phẩm", "triệu đ/tấn"),
    ("days", "Số ngày có số liệu", "ngày"),
    ("no_purchase_days", "Số ngày không tổ chức thu mua", "ngày"),
]

CONSUMPTION_COLS: list[Col] = [
    ("qty", "Tổng sản lượng tiêu thụ", "tấn"),
    ("qty_long_term", "HĐ dài hạn", "tấn"),
    ("qty_spot", "HĐ chuyến", "tấn"),
    ("qty_principle", "HĐ nguyên tắc", "tấn"),
    ("qty_unknown_type", "HĐ chưa khai loại", "tấn"),
    ("qty_export", "XK / UTXK", "tấn"),
    ("qty_domestic", "Tiêu thụ trong nước", "tấn"),
    ("qty_internal", "Tiêu thụ nội bộ", "tấn"),
    ("revenue_ty", "Doanh thu", "tỷ đồng"),
    ("avg_price_trieu", "Giá bán bình quân", "triệu đ/tấn"),
    ("lines", "Số dòng bán", "dòng"),
]

#: Kế hoạch là chỉ tiêu NĂM của TỪNG ĐƠN VỊ → chỉ chèn khi nhóm theo đơn vị/khu vực;
#: nhóm theo ngày/chủng loại/loại HĐ thì mọi ô đều trống, thà bỏ cột còn hơn để cột rỗng.
_PLAN_COLS: list[Col] = [
    ("plan_sales_spot_tonnes", "KH tiêu thụ HĐ chuyến", "tấn (chỉ tiêu năm)"),
    ("pct_plan_sales_spot", "% thực hiện KH tiêu thụ", "% (= HĐ chuyến / KH)"),
]
#: % KH doanh thu tính trên RỔ đơn vị được giao KH (cùng luật Dashboard) — xem `_attach_plan`.
_REVENUE_PLAN_COLS: list[Col] = [
    ("plan_revenue_ty", "KH doanh thu năm", "tỷ đồng (chỉ tiêu năm)"),
    ("pct_plan_revenue", "% thực hiện KH doanh thu", "% (= DT đơn vị có KH / KH)"),
]
#: % KH thu mua so sản lượng mủ NGUYÊN LIỆU của kỳ đang xem với chỉ tiêu CẢ NĂM — nói rõ trong nhãn
#: để không ai đọc nhầm thành luỹ kế từ đầu năm (kỳ mặc định của màn Thống kê là TUẦN).
_PURCHASE_PLAN_COLS: list[Col] = [
    ("plan_tonnes", "KH thu mua năm", "tấn (chỉ tiêu năm)"),
    ("pct_plan", "% KH năm", "% (= mủ NL kỳ này / KH năm)"),
]


def _with_plan(cols: list[Col], after: str, plan_cols: list[Col], group_by: str) -> list[Col]:
    """Chèn cột kế hoạch ngay sau cột `after` — đúng chỗ người đọc cần, và chỉ khi có số để điền."""
    if group_by not in ("company", "region"):
        return cols
    i = [c[0] for c in cols].index(after) + 1
    return [*cols[:i], *plan_cols, *cols[i:]]


def consumption_cols(group_by: str) -> list[Col]:
    """Cột bảng Tiêu thụ — kế hoạch đứng ngay sau "HĐ chuyến" (tử số của % nằm cạnh mẫu số)."""
    cols = _with_plan(CONSUMPTION_COLS, "qty_spot", _PLAN_COLS, group_by)
    return _with_plan(cols, "revenue_ty", _REVENUE_PLAN_COLS, group_by)


def purchase_cols(group_by: str) -> list[Col]:
    """Cột bảng Thu mua — kế hoạch đứng ngay sau "Sản lượng mủ nguyên liệu" (chính là tử số)."""
    return _with_plan(PURCHASE_COLS, "qty_material", _PURCHASE_PLAN_COLS, group_by)


CONSUMPTION_DETAIL_COLS: list[Col] = [
    ("as_of", "Ngày", ""), ("company", "Đơn vị", ""), ("source_label", "Nguồn mủ", ""),
    ("code", "Số HĐ/PL", ""), ("contract_label", "Loại HĐ", ""), ("channel_label", "Hình thức", ""),
    ("grade", "Chủng loại", ""), ("qty", "Số lượng", "tấn"), ("price", "Đơn giá", ""),
    ("ccy", "Loại tiền", ""), ("revenue_vnd", "Doanh thu", "đồng"),
    ("warehouse_date", "Ngày xuất kho", ""), ("invoice_date", "Ngày hoá đơn", ""),
]

STOCK_COLS: list[Col] = [
    ("as_of", "Ngày lấy số", ""),
    ("age_days", "Số cũ so với ngày chốt", "ngày"),
    ("not_warehoused", "Tồn thành phẩm chưa nhập kho", "tấn"),
    ("warehoused", "Tồn thành phẩm đã nhập kho", "tấn"),
    ("total", "Tổng tồn kho thành phẩm", "tấn"),
    ("signed_undelivered", "Đã ký HĐ chưa giao", "tấn quy khô"),
    ("tradable", "Tồn có thể giao dịch", "tấn"),
    ("material", "Tồn kho nguyên liệu (quy khô)", "tấn"),
]

GROUP_LABELS = {"company": "Đơn vị", "region": "Khu vực", "grade": "Chủng loại", "day": "Ngày",
                "material": "Loại mủ", "contract": "Loại HĐ", "channel": "Hình thức",
                "source": "Nguồn mủ", "none": "Chi tiết"}


def _head(ws, row: int, col: int, value: Any, *, bold: bool = True, size: int = 10) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(bold=bold, size=size)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.fill = _HEAD_FILL
    c.border = _BORDER


#: Sheet PHỤ đi kèm file (bảng phẳng): {name, title, note, columns, rows}. Không có dòng Tổng cộng —
#: dòng tổng nằm dưới vùng dữ liệu sẽ lọt vào bộ lọc/pivot của người đọc.
Sheet = dict[str, Any]


def _table(ws, *, title: str, period_label: str, period: str, note: str, cols: list[Col],
           rows: list[dict], totals: dict | None, label_col: int, flat: bool = False) -> None:
    """Đổ tiêu đề + bảng vào MỘT sheet.

    `label_col` = số cột đầu làm nhãn nhóm (0 = bảng phẳng, không có dòng Tổng cộng).
    `flat=True` (sheet chi tiết): gộp đơn vị tính vào ngay tiêu đề để header chỉ còn MỘT dòng —
    có vậy `auto_filter` mới đúng, chứ để 2 dòng thì dòng đơn vị tính bị Excel coi là dữ liệu.
    """
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"{period_label}: {period}").font = Font(bold=True)
    ws.cell(row=3, column=1, value=note).font = Font(italic=True, size=9, color="666666")
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, min(len(cols), 8)))

    head = 5
    for i, (_, label, unit) in enumerate(cols, start=1):
        _head(ws, head, i, f"{label} ({unit})" if flat and unit else label)
        if not flat:
            _head(ws, head + 1, i, unit, bold=False, size=8)
    first = head + (1 if flat else 2)

    r = first
    for row in rows:
        for i, (key, *_rest) in enumerate(cols, start=1):
            c = ws.cell(row=r, column=i, value=row.get(key))
            if isinstance(row.get(key), (int, float)):
                c.number_format = _NUM
            c.border = _BORDER
        r += 1

    if totals is not None and label_col:
        # Nhãn "Tổng cộng" đặt ở cột NHÓM (sau cột Khu vực nếu có) cho khớp bảng trên web.
        for i in range(1, label_col + 1):
            c = ws.cell(row=r, column=i, value="Tổng cộng" if i == label_col else None)
            c.font = Font(bold=True)
            c.fill = _TOTAL_FILL
            c.border = _BORDER
        for i, (key, *_rest) in enumerate(cols[label_col:], start=label_col + 1):
            c = ws.cell(row=r, column=i, value=totals.get(key))
            c.number_format = _NUM
            c.font = Font(bold=True)
            c.fill = _TOTAL_FILL
            c.border = _BORDER

    if flat and rows:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(cols))}{first + len(rows) - 1}"
    ws.freeze_panes = ws.cell(row=first, column=2)
    ws.column_dimensions["A"].width = 26
    for i in range(2, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16


def build_xlsx(*, title: str, period: str, note: str, group_by: str, columns: list[Col],
               rows: list[dict], totals: dict | None, label_key: str = "label",
               period_label: str = "Kỳ báo cáo", sheets: list[Sheet] | None = None) -> bytes:
    """Dựng .xlsx: tiêu đề + kỳ + ghi chú bộ lọc, bảng dữ liệu, dòng Tổng cộng (nếu có).

    `sheets` = các sheet CHI TIẾT đi kèm (mỗi dòng một bản ghi, có sẵn bộ lọc của Excel) — bảng
    tổng hợp trả lời "bao nhiêu", sheet chi tiết trả lời "gồm những gì" mà không phải tải lại số.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = GROUP_LABELS.get(group_by, "Thống kê")[:31]
    detail = group_by == "none"
    # Nhóm theo đơn vị thì kèm cột Khu vực (giống bảng trên web) để lọc/pivot lại trong Excel.
    lead: list[Col] = [] if detail else [(label_key, GROUP_LABELS.get(group_by, "Nhóm"), "")]
    if group_by == "company":
        lead = [("region", "Khu vực", ""), *lead]
    _table(ws, title=title, period_label=period_label, period=period, note=note,
           cols=lead + columns, rows=rows, totals=None if detail else totals, label_col=len(lead))
    for s in sheets or []:
        _table(wb.create_sheet(str(s["name"])[:31]), title=s["title"], period_label=period_label,
               period=period, note=s["note"], cols=s["columns"], rows=s["rows"],
               totals=None, label_col=0, flat=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
