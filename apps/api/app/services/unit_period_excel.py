"""Xuất Excel báo cáo theo kỳ — bám đúng cột của 2 sheet mẫu trong
"Biểu mẫu-Báo cáo tuần-năm 2026": "(2) Biểu mẫu-Thu mua" và "(1) Biểu mẫu Tiêu thụ -Tồn kho".

Mỗi đơn vị 1 dòng (kèm Khu vực), cuối bảng là dòng Tổng cộng. Số liệu do
`unit_period_report.period_report()` tính sẵn theo quy tắc cộng dồn / thời điểm / bình quân.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_HEAD_FILL = PatternFill("solid", fgColor="D9E7D5")
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="9AA5A0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NUM = "#,##0.00"

# (khoá dữ liệu, tiêu đề cột, ghi chú công thức của mẫu, ĐƠN VỊ TÍNH)
# Đơn vị tính nằm CHUNG tuple với cột — trước đây tách thành list riêng nên thêm cột mà quên
# thêm đơn vị là cả bảng lệch nhãn mà không ai biết.
_PURCHASE_COLS: list[tuple[str, str, str, str]] = [
    ("latex_wet", "Sản lượng thu mua mủ nước", "cộng dồn", "tấn"),
    ("coagulum", "Sản lượng thu mua mủ chén", "cộng dồn", "tấn"),
    ("cup_raw", "Sản lượng mủ NL nước chưa cán vắt (chén)", "cộng dồn", "tấn"),
    ("price_cup_raw_avg", "Đơn giá BQ mủ NL chưa cán vắt", "bình quân gia quyền", "đồng/kg"),
    ("rss_pressed", "Sản lượng mủ NL đã cán vắt (RSS)", "cộng dồn", "tấn"),
    ("price_rss_pressed_avg", "Đơn giá BQ mủ NL đã cán vắt", "bình quân gia quyền", "đồng/kg"),
    ("finished_qty", "Sản lượng thu mua thành phẩm", "cộng dồn các chủng loại", "tấn"),
    ("total_purchase", "Tổng sản lượng thu mua", "= mủ nước + mủ chén", "tấn"),
    ("price_latex_avg", "Giá thu mua mủ nước BQ", "bình quân gia quyền", "đồng/độ TSC"),
    ("price_cup_avg", "Giá thu mua mủ chén BQ", "bình quân gia quyền", "đồng/độ"),
    ("no_purchase_days", "Số ngày không tổ chức thu mua", "đếm ngày", "ngày"),
    ("plan_tonnes", "Kế hoạch thu mua", "số liệu năm", "tấn"),
    ("pct_plan", "% thực hiện kế hoạch", "= thực hiện / kế hoạch", "%"),
    ("consumption", "Sản lượng tiêu thụ mủ thu mua", "cộng dồn", "tấn"),
    ("finished_sold_qty", "Sản lượng tiêu thụ mủ thành phẩm", "cộng dồn", "tấn"),
    ("revenue_ty", "Doanh thu tiêu thụ mủ thu mua", "cộng dồn", "tỷ đồng"),
    ("avg_sell_price", "Giá bán bình quân", "= doanh thu / sản lượng", "triệu đ/tấn"),
]

# TỒN KHO: khối 1 và khối 2 là HAI chỉ tiêu khác nhau → mỗi khối một dòng riêng, rồi mới tới tổng.
# Khối 3 (đã ký hợp đồng chưa giao) là phần NẰM TRONG tồn kho thành phẩm đã có hợp đồng →
# chỉ báo, không cộng thêm vào tổng (cộng nữa là tính trùng) và cũng không trừ ra.
_CONSUMPTION_COLS: list[tuple[str, str, str, str]] = [
    ("signed_lt_tonnes", "Tổng SL đã ký HĐ dài hạn", "số liệu năm", "tấn"),
    ("lt_export", "HĐ dài hạn — XK/UTXK", "cộng dồn", "tấn"),
    ("lt_domestic", "HĐ dài hạn — Tiêu thụ trong nước", "cộng dồn", "tấn"),
    ("spot_export", "HĐ chuyến — XK/UTXK", "cộng dồn", "tấn"),
    ("spot_domestic", "HĐ chuyến — Tiêu thụ trong nước", "cộng dồn", "tấn"),
    ("total_consumption", "Tổng tiêu thụ", "= 4 cột trên (cũ) + sản lượng hợp đồng (mới)", "tấn"),
    ("export_total", "Tổng XK/UTXK", "= dài hạn + chuyến + hợp đồng", "tấn"),
    ("domestic_total", "Tổng Tiêu thụ trong nước", "= dài hạn + chuyến + hợp đồng", "tấn"),
    ("internal_total", "Tiêu thụ nội bộ", "cộng dồn (nguồn hợp đồng)", "tấn"),
    ("revenue_ty", "Doanh thu cao su", "cộng dồn", "tỷ đồng"),
    ("avg_sell_price", "Giá bán bình quân", "= doanh thu / tiêu thụ", "triệu đ/tấn"),
    # Chi phí trên dòng bán — chỉ có ở nguồn hợp đồng (mảng sales/sales_own cũ không ghi chi phí).
    ("cost_lines", "Tổng chi phí trên dòng bán", "cộng dồn (nguồn hợp đồng)", "triệu đồng"),
    ("cost_export", "Chi phí — XK/UTXK", "cộng dồn (nguồn hợp đồng)", "triệu đồng"),
    ("cost_domestic", "Chi phí — Tiêu thụ trong nước", "cộng dồn (nguồn hợp đồng)", "triệu đồng"),
    ("cost_internal", "Chi phí — Tiêu thụ nội bộ", "cộng dồn (nguồn hợp đồng)", "triệu đồng"),
    # Ngày của ảnh chụp tồn kho — có thể sớm hơn ngày cuối kỳ nếu đơn vị chưa cập nhật tồn.
    ("stock_as_of", "Ngày lấy số tồn kho", "thời điểm", ""),
    ("stock_not_warehoused", "Tồn kho thành phẩm chế biến chưa nhập kho", "thời điểm", "tấn"),
    ("stock_warehoused", "Tồn kho thành phẩm đã nhập kho", "thời điểm", "tấn"),
    ("stock_finished", "Tổng tồn kho thành phẩm", "= chưa nhập kho + đã nhập kho", "tấn"),
    ("stock_finished_hd", "Số lượng đã ký hợp đồng chưa giao",
     "thời điểm — phần NẰM TRONG tồn kho thành phẩm đã có hợp đồng", "tấn"),
]
_TAIL_COLS: list[tuple[str, str, str, str]] = [
    ("stock_material", "Tồn kho nguyên liệu chưa sản xuất (quy khô)", "thời điểm", "tấn"),
    ("carry_lt_tonnes", "HĐ dài hạn năm trước chuyển sang", "số liệu năm", "tấn"),
    ("carry_spot_tonnes", "HĐ chuyến năm trước chuyển sang", "số liệu năm", "tấn"),
]

_TITLE = {
    "purchase": ("BÁO CÁO THU MUA", "Về công tác Thu mua và Tiêu thụ mủ nguyên liệu"),
    "consumption": ("BÁO CÁO TIÊU THỤ - TỒN KHO", "Về công tác Tiêu thụ - Tồn kho mủ cao su"),
}
# Chỉ tiêu KHÔNG được cộng ở dòng Tổng cộng (giá / tỷ lệ → tính lại hoặc bỏ trống).
_NO_SUM = {"price_latex_avg", "price_cup_avg", "price_cup_raw_avg",
           "price_rss_pressed_avg", "pct_plan", "avg_sell_price"}


def _columns(kind: str, grades: list[str]) -> tuple[list[tuple[str, str, str, str]], list[str]]:
    """Cột dữ liệu + đơn vị tính. Biểu tiêu thụ chèn tồn kho theo chủng loại vào giữa."""
    if kind == "purchase":
        cols = list(_PURCHASE_COLS)
    else:
        cols = list(_CONSUMPTION_COLS)
        # Tồn kho thành phẩm tách theo chủng loại (khối 1 + khối 2 gộp lại theo từng loại).
        cols += [(f"grade::{g}", g, "thời điểm", "tấn") for g in grades]
        cols += _TAIL_COLS
    return cols, [c[3] for c in cols]      # đơn vị suy thẳng từ cột → không thể lệch


def _value(row: dict, key: str) -> Any:
    if key.startswith("grade::"):
        return (row.get("stock_by_grade") or {}).get(key[7:])
    v = row.get(key)
    if key == "stock_as_of" and v:      # ngày hiển thị chuẩn VN (dữ liệu nội bộ giữ ISO)
        y, m, d = str(v)[:10].split("-")
        return f"{d}/{m}/{y}"
    return v


def build_period_xlsx(report: dict) -> bytes:
    """Dựng file .xlsx từ kết quả `period_report()` → bytes để trả về client."""
    kind = report["kind"]
    grades: list[str] = report.get("grades") or []
    cols, units = _columns(kind, grades)
    rows: list[dict] = report.get("rows") or []

    wb = Workbook()
    ws = wb.active
    ws.title = "Thu mua" if kind == "purchase" else "Tiêu thụ - Tồn kho"

    title, sub = _TITLE[kind]
    ncol = 2 + len(cols)
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=sub).font = Font(italic=True)
    ws.cell(row=3, column=1,
            value=f"Kỳ báo cáo: {report['date_from']} → {report['date_to']}").font = Font(bold=True)
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=min(ncol, 8))

    head, unit_row, formula_row = 5, 6, 7
    headers = ["Khu vực", "Đơn vị"] + [c[1] for c in cols]
    all_units = ["", ""] + units
    all_formulas = ["", ""] + [c[2] for c in cols]
    for i, (h, u, f) in enumerate(zip(headers, all_units, all_formulas), start=1):
        for r, v, bold in ((head, h, True), (unit_row, u, False), (formula_row, f, False)):
            cell = ws.cell(row=r, column=i, value=v)
            cell.font = Font(bold=bold, size=10 if r == head else 8,
                             italic=(r == formula_row), color="000000" if r != formula_row else "666666")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.fill = _HEAD_FILL
            cell.border = _BORDER

    r = formula_row + 1
    for row in rows:
        ws.cell(row=r, column=1, value=row.get("region") or "").border = _BORDER
        ws.cell(row=r, column=2, value=row.get("company")).border = _BORDER
        for i, (key, *_) in enumerate(cols, start=3):
            v = _value(row, key)
            c = ws.cell(row=r, column=i, value=v)
            if not isinstance(v, str):       # cột ngày là chữ → giữ nguyên, không ép định dạng số
                c.number_format = _NUM
            c.border = _BORDER
        r += 1

    # Dòng Tổng cộng: chỉ cộng chỉ tiêu dòng chảy/tồn kho; giá & % bỏ trống (không cộng được).
    ws.cell(row=r, column=2, value="Tổng cộng").font = Font(bold=True)
    ws.cell(row=r, column=1).fill = _TOTAL_FILL
    ws.cell(row=r, column=2).fill = _TOTAL_FILL
    for i, (key, *_) in enumerate(cols, start=3):
        total = None
        if key not in _NO_SUM:
            vals = [v for v in (_value(x, key) for x in rows) if isinstance(v, (int, float))]
            total = sum(vals) if vals else None
        c = ws.cell(row=r, column=i, value=total)
        c.number_format = _NUM
        c.font = Font(bold=True)
        c.fill = _TOTAL_FILL
        c.border = _BORDER

    ws.freeze_panes = ws.cell(row=formula_row + 1, column=3)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 24
    for i in range(3, ncol + 1):
        ws.column_dimensions[get_column_letter(i)].width = 15

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
