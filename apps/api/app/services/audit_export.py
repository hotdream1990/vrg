"""Xuất Nhật ký hoạt động ra Excel — để kèm biên bản/đối chiếu khi cần chứng minh.

Cột "Nội dung thay đổi" gói gọn diff trước→sau thành một chuỗi đọc được, giống cột trên màn hình.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_HEADERS = [
    ("Thời điểm", 20), ("Người thao tác", 22), ("Vai trò", 16), ("Đăng nhập hộ", 16),
    ("Nhóm số liệu", 24), ("Thao tác", 12), ("Bản ghi", 34), ("Ngày số liệu", 14),
    ("Đơn vị", 26), ("Nội dung thay đổi", 70), ("IP", 15), ("Ghi chú", 28),
]
_ROLE_LABEL = {"admin": "Quản trị viên", "editor": "Chuyên viên nhập liệu",
               "viewer": "Người xem", "member": "Đơn vị thành viên"}


def _flatten(value: Any, prefix: str = "", out: dict[str, Any] | None = None) -> dict[str, Any]:
    """Object/mảng lồng nhau → {'đường.dẫn': giá trị} để so từng ô."""
    out = {} if out is None else out
    if isinstance(value, list):
        for i, v in enumerate(value):
            _flatten(v, f"{prefix}[{i}]", out)
    elif isinstance(value, dict):
        for k, v in value.items():
            _flatten(v, f"{prefix}.{k}" if prefix else k, out)
    elif prefix:
        out[prefix] = value
    return out


def _fmt(v: Any) -> str:
    if v is None or v == "":
        return "(trống)"
    if isinstance(v, bool):
        return "Có" if v else "Không"
    return str(v)


def _changes(row: dict[str, Any]) -> str:
    """Diff trước→sau dạng chuỗi; thêm mới/xoá chỉ liệt kê một phía (không có mũi tên)."""
    before, after = _flatten(row.get("before")), _flatten(row.get("after"))
    action = row.get("action")
    parts = []
    for path in sorted(set(before) | set(after)):
        old, new = before.get(path), after.get(path)
        if old == new or (old in (None, "") and new in (None, "")):
            continue
        if action == "create":
            parts.append(f"{path}: {_fmt(new)}")
        elif action == "delete":
            parts.append(f"{path}: {_fmt(old)}")
        else:
            parts.append(f"{path}: {_fmt(old)} → {_fmt(new)}")
    return " · ".join(parts)


def build_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """Danh sách dòng nhật ký (đã lọc) → file Excel 1 sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Nhật ký hoạt động"
    head_fill = PatternFill("solid", fgColor="0A9E48")
    for col, (title, width) in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"

    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=str(r.get("at", ""))[:19].replace("T", " "))
        ws.cell(row=i, column=2, value=r.get("actor", ""))
        ws.cell(row=i, column=3, value=_ROLE_LABEL.get(r.get("actor_role", ""), ""))
        ws.cell(row=i, column=4, value=r.get("on_behalf", ""))
        ws.cell(row=i, column=5, value=r.get("entity_label", ""))
        ws.cell(row=i, column=6, value=r.get("action_label", ""))
        ws.cell(row=i, column=7, value=r.get("entity_key", ""))
        ws.cell(row=i, column=8, value=r.get("as_of") or "")
        ws.cell(row=i, column=9, value=r.get("company", ""))
        ws.cell(row=i, column=10, value=_changes(r)).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=11, value=r.get("ip", ""))
        ws.cell(row=i, column=12, value=r.get("note", ""))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
