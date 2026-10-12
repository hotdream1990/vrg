"""Kiểu dáng file Excel "Phiếu báo giá mủ thị trường" — theo mẫu Ban TTKD (Times New Roman, xanh lá đậm).

Chỉ chứa hằng số + hàm ghi ô; nội dung từng mục nằm ở `market_quote_excel*.py`.
Quy ước màu của mẫu: chữ XANH DƯƠNG = số liệu nhập · chữ ĐEN = số tự tính.
Chuỗi do người dùng gõ luôn qua `safe_cell` (chống "công thức Excel" nhúng trong tên khách hàng, ghi chú…).
"""

from __future__ import annotations

import math
from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.services.xlsx_text import safe_cell

FONT = "Times New Roman"
GREEN = "1F5C3A"
BAR_FILL = "C6E0B4"
GROUP_FILL = "E2EFDA"
INPUT_BLUE = "0000FF"
MUTED = "595959"

FMT_INT = r"#,##0;\(#,##0\);\-"
FMT_PCT = r"0.0%;\(0.0%\);\-"
FMT_DEC = r"#,##0.0;\(#,##0.0\);\-"
FMT_DATE = "dd/mm/yyyy"

_GREY = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_GREY, right=_GREY, top=_GREY, bottom=_GREY)
_WHITE = PatternFill("solid", fgColor="FFFFFF")

#: Độ rộng cột A..K (đơn vị ký tự) — rộng hơn mẫu ở A/B/F để chứa nhãn tỷ giá, tên đơn vị tư nhân, khách hàng, công ty thành viên.
COL_WIDTHS = {"A": 30, "B": 26, "C": 17, "D": 17.4, "E": 18, "F": 26, "G": 16, "H": 13.4, "I": 9, "J": 16, "K": 19.9}
LAST_COL = "K"


def fmt_for(value: Any, fmt: str) -> str:
    """Số nguyên giữ định dạng gốc; số có phần lẻ (vd 2.210,5 USD/tấn) hiện 1 số lẻ — khỏi bị làm tròn khi xem."""
    if fmt == FMT_INT and isinstance(value, float) and not value.is_integer():
        return FMT_DEC
    return fmt


def font(size: float = 11, *, bold: bool = False, color: str | None = None, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def put(ws: Worksheet, ref: str, value: Any, *, f: Font | None = None, fmt: str | None = None,
        fill: PatternFill | None = None, border: bool = False, h: str | None = None,
        wrap: bool = False, v: str | None = "center") -> None:
    """Ghi 1 ô: giá trị (chuỗi đã qua safe_cell) + font/định dạng/viền."""
    c = ws[ref]
    c.value = safe_cell(value)
    c.font = f or font()
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def bar(ws: Worksheet, row: int, text: str, last_col: str = LAST_COL) -> None:
    """Thanh tiêu đề mục (nền xanh nhạt, chữ xanh đậm) trải hết `A:last_col`."""
    ws.merge_cells(f"A{row}:{last_col}{row}")
    put(ws, f"A{row}", text, f=font(12, bold=True, color=GREEN), fill=PatternFill("solid", fgColor=BAR_FILL))
    ws.row_dimensions[row].height = 19.5


def header(ws: Worksheet, row: int, labels: dict[str, str]) -> None:
    """Dòng tiêu đề cột (nền xanh đậm, chữ trắng, in hoa)."""
    for col, label in labels.items():
        put(ws, f"{col}{row}", label, f=font(9, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=GREEN),
            border=True, h="center", wrap=True)
    ws.row_dimensions[row].height = 30


def label_cell(ws: Worksheet, ref: str, text: str) -> None:
    """Ô nhãn hàng (chữ đậm, viền; tên dài tự xuống dòng)."""
    put(ws, ref, text, f=font(bold=True), border=True, wrap=True)


def input_cell(ws: Worksheet, ref: str, value: Any, *, fmt: str | None = None, h: str | None = None,
               wrap: bool = False) -> None:
    """Ô số liệu nhập (chữ xanh dương, nền trắng, viền xám). Số rỗng → ô trống nhưng vẫn có viền."""
    put(ws, ref, value, f=font(color=INPUT_BLUE), fmt=fmt and fmt_for(value, fmt), fill=_WHITE,
        border=True, h=h, wrap=wrap)


def calc_cell(ws: Worksheet, ref: str, value: Any, *, fmt: str = FMT_INT) -> None:
    """Ô tự tính (chữ đen)."""
    put(ws, ref, value, fmt=fmt_for(value, fmt), border=True)


def note_row(ws: Worksheet, row: int, label: str, text: str, last_col: str = "E") -> None:
    """Dòng "Ghi chú:" — nhãn nhỏ đậm + nội dung gộp ô, tự tính chiều cao theo số dòng chữ."""
    put(ws, f"A{row}", label, f=font(9, bold=True), v="top")
    ws.merge_cells(f"B{row}:{last_col}{row}")
    put(ws, f"B{row}", text or "", wrap=True, v="top")
    width = sum(COL_WIDTHS[chr(c)] for c in range(ord("B"), ord(last_col) + 1))
    lines = sum(max(1, math.ceil(len(part) / max(width, 1))) for part in (text or "").split("\n"))
    ws.row_dimensions[row].height = max(15, 15 * lines)


def caption(ws: Worksheet, row: int, text: str, last_col: str = LAST_COL) -> None:
    """Dòng chú thích nhỏ, xám, nghiêng (gộp hết bề ngang)."""
    ws.merge_cells(f"A{row}:{last_col}{row}")
    put(ws, f"A{row}", text, f=font(9, color=MUTED, italic=True), wrap=True)
