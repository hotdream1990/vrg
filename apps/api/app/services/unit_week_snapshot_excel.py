"""Xuất Excel bản lưu số liệu tuần — 1 file, 2 sheet: Thu mua · Tiêu thụ - Tồn kho.

Mỗi sheet dựng bằng CHÍNH `unit_period_excel.build_period_xlsx` từ biểu đã lưu (không tính lại), nên
cột/đơn vị/dòng Tổng cộng y hệt file "Xuất Excel" của màn Báo cáo tổng hợp. Thêm 1 dòng ghi rõ
đây là bản lưu, chụp lúc nào — người nhận file không nhầm với số đang sống.
"""

from __future__ import annotations

import io
from copy import copy
from datetime import datetime
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app.services import unit_period_excel

_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
_NOTE_ROW = 4                      # dòng trống giữa "Kỳ báo cáo" và tiêu đề cột của mẫu


def stamp(iso: str | None) -> str:
    """ISO datetime → 'HH:MM ngày DD/MM/YYYY' theo giờ Việt Nam."""
    if not iso:
        return "—"
    t = datetime.fromisoformat(iso).astimezone(_TZ)
    return f"{t:%H:%M} ngày {t:%d/%m/%Y}"


def _copy_sheet(src: Worksheet, dst: Worksheet) -> None:
    """Chép giá trị + định dạng giữa 2 workbook (openpyxl chỉ tự chép trong cùng workbook)."""
    for row in src.iter_rows():
        for c in row:
            d = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                d.font, d.fill, d.border = copy(c.font), copy(c.fill), copy(c.border)
                d.alignment, d.number_format = copy(c.alignment), c.number_format
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for key, dim in src.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
    dst.freeze_panes = src.freeze_panes


def build(snap: dict[str, Any]) -> bytes:
    """Bản lưu (dict từ `unit_week_snapshot.detail`) → bytes .xlsx."""
    note = (f"BẢN LƯU {snap['label'].upper()} — chụp lúc {stamp(snap.get('taken_at'))} "
            f"(hạn nhập ngày Chủ nhật: {stamp(snap.get('deadline_at'))}). "
            "Sửa số liệu sau thời điểm chụp không làm đổi bản lưu này.")
    out = Workbook()
    out.remove(out.active)
    for kind in ("purchase", "consumption"):
        src = load_workbook(io.BytesIO(unit_period_excel.build_period_xlsx(snap[kind]))).active
        ws = out.create_sheet(src.title)
        _copy_sheet(src, ws)
        cell = ws.cell(row=_NOTE_ROW, column=1, value=note)
        cell.font = Font(bold=True, italic=True, color="9C2A00")
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def file_name(snap: dict[str, Any]) -> str:
    return f"snapshot-so-lieu-tuan-{snap['week_no']}-{snap['year']}.xlsx"
