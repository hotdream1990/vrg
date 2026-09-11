"""Xuất Excel cho màn "Cảnh báo bất thường" — một sheet mỗi nhóm luật + sheet "Tổng quan".

Dựng file trực tiếp bằng openpyxl (không dùng lại `unit_analytics_excel`): khuôn cột ở đây là
`columns: [{key, label}]` do `anomaly_rules.scan()` trả về — khác kiểu `Col` (khoá, nhãn, đơn vị
tính) của các bảng Thống kê đơn vị, nên tách riêng cho khỏi gượng ép hai khuôn vào một hàm.
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_HEAD_FILL = PatternFill("solid", fgColor="D9E7D5")
_SEV_FILL = {
    "high": PatternFill("solid", fgColor="F8D7DA"),
    "medium": PatternFill("solid", fgColor="FFF3CD"),
    "low": PatternFill("solid", fgColor="D1ECF1"),
}
_SEV_LABEL = {"high": "Nghiêm trọng", "medium": "Trung bình", "low": "Nhẹ"}
_THIN = Side(style="thin", color="9AA5A0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_NUM = "#,##0.00"


#: Ký tự Excel CẤM trong tên sheet — gặp là openpyxl ném lỗi và hỏng cả file xuất.
#: Nhãn nhóm tiếng Việt có dấu "/" (vd "Chưa nộp / thiếu một phần") nên phải lọc, không chỉ cắt độ dài.
_BAD_SHEET_CHARS = str.maketrans({c: "-" for c in "\\/*?:[]"})


def _sheet_name(label: str, used: set[str]) -> str:
    """Excel giới hạn 31 ký tự, cấm một số ký tự, không được trùng tên."""
    base = (label or "Nhóm").translate(_BAD_SHEET_CHARS).strip()[:31] or "Nhóm"
    name, i = base, 2
    while name in used:
        suffix = f" ({i})"
        name = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def _title_block(ws: Worksheet, title: str, period: str, note: str, ncols: int) -> None:
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Kỳ rà soát: {period}").font = Font(bold=True)
    if note:
        ws.cell(row=3, column=1, value=note).font = Font(italic=True, size=9, color="666666")
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, min(ncols, 8)))


def _group_sheet(ws: Worksheet, group: dict[str, Any], period: str) -> None:
    """Một nhóm cảnh báo = một sheet: cột động theo `group["columns"]`, dòng theo `group["rows"]`."""
    cols: list[dict[str, str]] = group.get("columns") or []
    rows: list[dict[str, Any]] = group.get("rows") or []
    _title_block(ws, str(group.get("label") or group.get("key") or ""), period,
                 str(group.get("desc") or ""), len(cols))

    head = 5
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=head, column=i, value=c.get("label", c.get("key")))
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER

    r = head + 1
    for row in rows:
        for i, c in enumerate(cols, start=1):
            value = row.get(c.get("key"))
            cell = ws.cell(row=r, column=i, value=value)
            cell.border = _BORDER
            if isinstance(value, (int, float)):
                cell.number_format = _NUM
        r += 1

    if rows and cols:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(cols))}{r - 1}"
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20


def _overview_sheet(ws: Worksheet, result: dict[str, Any]) -> None:
    """Sheet đầu tiên: số cảnh báo theo mức độ + danh sách nhóm để định vị nhanh trước khi lật sheet."""
    period = f"{result.get('date_from', '')} → {result.get('date_to', '')}"
    _title_block(ws, "CẢNH BÁO BẤT THƯỜNG — TỔNG QUAN", period, "", 4)

    summary = result.get("summary") or {}
    row = 5
    for label, key in (("Tổng số cảnh báo", "total"), ("Nghiêm trọng", "high"),
                       ("Trung bình", "medium"), ("Nhẹ", "low"),
                       ("Số đơn vị liên quan", "units")):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=summary.get(key, 0))
        row += 1

    row += 1
    headers = ["Nhóm cảnh báo", "Mức độ", "Số dòng", "Số đơn vị"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i, value=h)
        cell.font = Font(bold=True)
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
    row += 1

    for group in result.get("groups") or []:
        sev = group.get("severity", "low")
        ws.cell(row=row, column=1, value=group.get("label")).border = _BORDER
        sev_cell = ws.cell(row=row, column=2, value=_SEV_LABEL.get(sev, sev))
        sev_cell.border = _BORDER
        sev_cell.fill = _SEV_FILL.get(sev, _HEAD_FILL)
        ws.cell(row=row, column=3, value=group.get("count", 0)).border = _BORDER
        ws.cell(row=row, column=4, value=group.get("units", 0)).border = _BORDER
        row += 1

    for i, w in enumerate((36, 16, 12, 12), start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_xlsx(result: dict[str, Any]) -> bytes:
    """Dựng .xlsx từ kết quả `anomaly_rules.scan()`: sheet "Tổng quan" + một sheet mỗi nhóm."""
    wb = Workbook()
    _overview_sheet(wb.active, result)
    wb.active.title = "Tổng quan"
    used = {"Tổng quan"}
    period = f"{result.get('date_from', '')} → {result.get('date_to', '')}"
    for group in result.get("groups") or []:
        name = _sheet_name(str(group.get("label") or group.get("key") or "Nhóm"), used)
        _group_sheet(wb.create_sheet(name), group, period)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
