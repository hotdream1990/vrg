"""Phần Mục 5 (đề xuất mua từ khách hàng), Mục 6 (giá mủ tư nhân) và Bảng đối chiếu của file Excel báo giá."""

from __future__ import annotations

import math
from typing import Any

from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.core.market_meta import MARKET_QUOTE_GRADES
from app.services import market_quote_excel_style as st
from app.services.private_price_benchmark import DEFAULT_PROCESSING_COST, svr3l_cost

_UP, _DOWN = "FF008000", "FFC00000"   # xanh: cao hơn · đỏ: thấp hơn (như mẫu)


def _blank(it: dict[str, Any]) -> bool:
    return not (it.get("grade") or it.get("customer") or it.get("source")) and all(
        it.get(k) is None for k in ("qty", "price_vnd", "price_usd"))


def _half_up(x: float) -> int:
    """Làm tròn nửa LÊN (round() của Python làm tròn về số chẵn — 17,5 → 18 nhưng 16,5 → 16, hạ giá oan)."""
    return math.floor(x + 0.5)


def proposal_price_by_grade(items: list[dict[str, Any]]) -> dict[str, float]:
    """Đơn giá đề xuất (đồng/tấn) THEO CHỦNG LOẠI cho bảng đối chiếu.

    Một chủng loại có thể có nhiều dòng đề xuất → bình quân gia quyền theo số lượng (dòng không ghi số
    lượng bị bỏ qua); không dòng nào có số lượng thì bình quân thường. Một dòng duy nhất = chính giá đó.
    """
    by_grade: dict[str, list[tuple[float | None, float]]] = {}
    for it in items:
        if it.get("grade") and it.get("price_vnd") is not None:
            by_grade.setdefault(it["grade"], []).append((it.get("qty"), it["price_vnd"]))
    out: dict[str, float] = {}
    for grade, rows in by_grade.items():
        weighted = [(q, p) for q, p in rows if q]
        if weighted:
            out[grade] = _half_up(sum(q * p for q, p in weighted) / sum(q for q, _ in weighted))
        else:
            out[grade] = _half_up(sum(p for _, p in rows) / len(rows))
    return out


def proposal_section(ws: Worksheet, row: int, prop: dict[str, Any]) -> int:
    """Mục 5 — mỗi dòng đề xuất một hàng (cùng chủng loại có thể nhiều dòng). Trả về dòng trống kế tiếp."""
    st.bar(ws, row, "MỤC 5. ĐỀ XUẤT MUA TỪ KHÁCH HÀNG — HÀNG VRG   (số lượng: tấn · đơn giá: VNĐ/tấn, USD/tấn)", last_col="G")
    st.header(ws, row + 1, {"A": "CHỦNG LOẠI", "B": "KHÁCH HÀNG", "C": "SỐ LƯỢNG (TẤN)", "D": "ĐƠN GIÁ (ĐỒNG/TẤN)",
                            "E": "ĐƠN GIÁ (USD/TẤN)", "F": "NGUỒN", "G": "THÀNH TIỀN (ĐỒNG)"})
    r = row + 2
    items = [it for it in prop.get("items") or [] if not _blank(it)]
    for it in items:
        st.label_cell(ws, f"A{r}", it.get("grade") or "")
        st.input_cell(ws, f"B{r}", it.get("customer") or None, wrap=True)
        st.input_cell(ws, f"C{r}", it.get("qty"), fmt=st.FMT_INT)
        st.input_cell(ws, f"D{r}", it.get("price_vnd"), fmt=st.FMT_INT)
        st.input_cell(ws, f"E{r}", it.get("price_usd"), fmt=st.FMT_INT)
        st.input_cell(ws, f"F{r}", it.get("source") or None, wrap=True)
        qty, price = it.get("qty"), it.get("price_vnd")
        st.calc_cell(ws, f"G{r}", _half_up(qty * price) if qty is not None and price is not None else None)
        r += 1
    if not items:
        ws.merge_cells(f"A{r}:G{r}")
        st.put(ws, f"A{r}", "Chưa có đề xuất nào trong phiếu.", f=st.font(color=st.MUTED, italic=True), border=True, h="center")
        r += 1
    st.note_row(ws, r, "Ghi chú:", prop.get("note") or "", last_col="G")
    return r + 2


def private_section(ws: Worksheet, row: int, prices: dict[str, Any], cost: float | None,
                    catalog: list[str]) -> int:
    """Mục 6 — giá mủ tư nhân + giá thành SVR 3L theo giá min / max. Chỉ liệt kê đơn vị CÓ giá trong phiếu."""
    cost_used = DEFAULT_PROCESSING_COST if cost is None else cost
    st.bar(ws, row, "MỤC 6. GIÁ MỦ TƯ NHÂN   (đồng/độ TSC · nhập 1 giá, hoặc Giá + Giá max nếu là khoảng giá)", last_col="E")
    ws.merge_cells(f"A{row + 1}:B{row + 1}")
    st.label_cell(ws, f"A{row + 1}", "Chi phí gia công SVR 3L (đ/tấn)")
    ws[f"B{row + 1}"].border = st.BORDER
    st.input_cell(ws, f"C{row + 1}", cost_used, fmt=st.FMT_INT)
    st.put(ws, f"D{row + 1}", "Giá thành SVR 3L = Giá mủ × 1,08 × 100.000 + Chi phí gia công",
           f=st.font(9, color=st.MUTED))
    st.header(ws, row + 2, {"A": "ĐƠN VỊ TƯ NHÂN", "B": "GIÁ", "C": "GIÁ MAX", "D": "GIÁ THÀNH SVR 3L\n(THEO GIÁ MIN)",
                            "E": "GIÁ THÀNH SVR 3L\n(THEO GIÁ MAX)"})
    ws.row_dimensions[row + 2].height = 39
    names = [n for n in catalog if n in prices] + [n for n in prices if n not in catalog]
    r = row + 3
    for name in names:
        lo, hi = (prices[name] or {}).get("price"), (prices[name] or {}).get("price_max")
        if lo is None:
            continue
        st.label_cell(ws, f"A{r}", name)
        st.input_cell(ws, f"B{r}", lo, fmt=st.FMT_INT)
        st.input_cell(ws, f"C{r}", hi, fmt=st.FMT_INT)
        st.calc_cell(ws, f"D{r}", svr3l_cost(lo, cost_used))
        st.calc_cell(ws, f"E{r}", svr3l_cost(hi, cost_used) if hi is not None else None)
        r += 1
    if r == row + 3:
        ws.merge_cells(f"A{r}:E{r}")
        st.put(ws, f"A{r}", "Chưa có đơn vị tư nhân nào báo giá.", f=st.font(color=st.MUTED, italic=True),
               border=True, h="center")
        r += 1
    return r + 1


def _diff(base: float | None, other: float | None) -> tuple[float | None, float | None]:
    """(chênh lệch, % so với cột gốc) = other − base; thiếu số liệu hoặc gốc = 0 thì để trống."""
    if base is None or other is None:
        return None, None
    d = other - base
    return d, (d / base if base else None)


def comparison_section(ws: Worksheet, row: int, q: dict[str, Any]) -> int:
    """Bảng đối chiếu: tư nhân vs VRG (nội địa · xuất khẩu) và đề xuất khách hàng vs giá tư nhân."""
    def price(key: str, grade: str) -> float | None:
        return ((q.get(key) or {}).get("prices") or {}).get(grade)

    proposals = proposal_price_by_grade((q.get("customer_proposal") or {}).get("items") or [])
    st.bar(ws, row, "BẢNG ĐỐI CHIẾU SO SÁNH   (Chênh lệch = cột so sánh − cột gốc · Xanh: cao hơn · Đỏ: thấp hơn)")
    group = Font(name=st.FONT, size=9, bold=True, color=st.GREEN)
    for rng, text in ((f"B{row + 1}:E{row + 1}", "NỘI ĐỊA: TƯ NHÂN  VỚI VRG"), (f"F{row + 1}:I{row + 1}", "XUẤT KHẨU: VRG-TN (USD/T)"),
                      (f"J{row + 1}:K{row + 1}", "ĐỀ XUẤT KHÁCH HÀNG MUA HÀNG VRG")):
        ws.merge_cells(rng)
        c = ws[rng.split(":")[0]]
        c.value, c.font, c.fill = text, group, PatternFill("solid", fgColor=st.GROUP_FILL)
        c.alignment = Alignment(horizontal="center")
    st.header(ws, row + 2, {
        "A": "CHỦNG LOẠI", "B": "TƯ NHÂN\n(đ/tấn)", "C": "VRG\n(đ/tấn)", "D": "CHÊNH LỆCH\nVRG-TN", "E": "%",
        "F": "TƯ NHÂN\n(USD/tấn)", "G": "VRG\n(USD/tấn)", "H": "CHÊNH LỆCH\nVRG-TN", "I": "%",
        "J": "ĐỀ XUẤT\n(đ/tấn)", "K": "CHÊNH LỆCH\nGIÁ BID VÀ ĐỀ XUẤT"})
    first = row + 3
    for i, g in enumerate(MARKET_QUOTE_GRADES):
        r = first + i
        tn_vnd, vrg_vnd = price("domestic_private", g), price("domestic_vrg", g)
        tn_usd, vrg_usd = price("domestic_export", g), price("export_vrg", g)
        d1, p1 = _diff(tn_vnd, vrg_vnd)
        d2, p2 = _diff(tn_usd, vrg_usd)
        offer = proposals.get(g)
        st.label_cell(ws, f"A{r}", g)
        for col, v in (("B", tn_vnd), ("C", vrg_vnd), ("D", d1), ("F", tn_usd), ("G", vrg_usd), ("H", d2),
                       ("J", offer), ("K", _diff(tn_vnd, offer)[0])):
            st.calc_cell(ws, f"{col}{r}", v)
        st.calc_cell(ws, f"E{r}", p1, fmt=st.FMT_PCT)
        st.calc_cell(ws, f"I{r}", p2, fmt=st.FMT_PCT)
    last = first + len(MARKET_QUOTE_GRADES) - 1
    for cols in ("D", "E", "H", "I", "K"):
        rng = f"{cols}{first}:{cols}{last}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=_UP)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=_DOWN)))
    st.caption(ws, last + 1, "Đề xuất = đơn giá đồng/tấn ở Mục 5; chủng loại có nhiều dòng đề xuất lấy bình quân gia quyền "
                             "theo số lượng. Chênh lệch giá bid và đề xuất = Đề xuất − giá tư nhân (nội địa).")
    return last + 2
