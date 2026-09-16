"""Xuất Lịch sử truy cập ra Excel — 2 sheet: tổng hợp theo tài khoản + chi tiết từng lượt.

Dùng khi cần gửi kèm báo cáo "đơn vị nào thực sự dùng hệ thống" cho lãnh đạo.
"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.roles import role_label
from app.services.xlsx_text import safe_cell

_SUMMARY_HEADERS = [
    ("Tài khoản", 26), ("Vai trò", 24), ("Đơn vị", 30), ("Đăng nhập gần nhất", 20),
    ("Hoạt động gần nhất", 20), ("Số lần đăng nhập", 16), ("Đăng nhập sai", 14),
    ("Lượt xem trang", 15), ("Số ngày hoạt động", 17), ("Số trang khác nhau", 18),
    ("Trang hay vào nhất", 34),
]
_DETAIL_HEADERS = [
    ("Thời điểm", 20), ("Tài khoản", 26), ("Vai trò", 24), ("Đơn vị", 30),
    ("Sự kiện", 18), ("Trang", 34), ("Đường dẫn", 30), ("Đăng nhập hộ", 16), ("IP", 16),
]

_HEAD_FILL = PatternFill("solid", fgColor="0A9E48")


def _stamp(value: Any) -> str:
    """'2026-09-16T08:12:03+07:00' → '2026-09-16 08:12:03' (rỗng nếu không có)."""
    return str(value or "")[:19].replace("T", " ")


def _write_header(ws, headers: list[tuple[str, int]]) -> None:  # noqa: ANN001 - worksheet
    for col, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"


def build_xlsx(summary: list[dict[str, Any]], detail: list[dict[str, Any]]) -> bytes:
    """Tổng hợp + chi tiết (đều đã lọc sẵn) → file Excel 2 sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Theo tài khoản"
    _write_header(ws, _SUMMARY_HEADERS)
    for i, r in enumerate(summary, start=2):
        top = r.get("top_page", "")
        views = r.get("top_page_views", 0)
        for col, value in enumerate([
            r.get("username", ""), role_label(r.get("role", "")),
            r.get("company", ""), _stamp(r.get("last_login")), _stamp(r.get("last_seen")),
            r.get("logins", 0), r.get("failed_logins", 0), r.get("page_views", 0),
            r.get("active_days", 0), r.get("distinct_pages", 0),
            f"{top} ({views} lượt)" if top else "",
        ], start=1):
            ws.cell(row=i, column=col, value=safe_cell(value))

    ws2 = wb.create_sheet("Chi tiết lượt truy cập")
    _write_header(ws2, _DETAIL_HEADERS)
    for i, r in enumerate(detail, start=2):
        for col, value in enumerate([
            _stamp(r.get("at")), r.get("username", ""),
            role_label(r.get("role", "")), r.get("company", ""),
            r.get("event_label", ""), r.get("label", ""), r.get("path", ""),
            r.get("on_behalf", ""), r.get("ip", ""),
        ], start=1):
            ws2.cell(row=i, column=col, value=safe_cell(value))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
