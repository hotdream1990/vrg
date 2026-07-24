"""Nhập/xuất số liệu báo cáo đơn vị bằng Excel — TẢI MẪU · ĐỌC FILE (xem trước) · GHI.

Mỗi loại biểu khai báo cột MỘT chỗ (`SPECS`) rồi dùng chung cho cả sinh file mẫu lẫn đọc file
người dùng nộp → mẫu và bộ đọc không bao giờ lệch nhau.

4 loại (`kind`):
  purchase   — Thu mua: phần mủ nước/mủ chén 1 dòng / (đơn vị, ngày); thu mua THÀNH PHẨM nhiều
               dòng (mỗi chủng loại 1 dòng) → gom thành mảng `finished`
  sales      — Tiêu thụ: NHIỀU dòng / (đơn vị, ngày) → tách theo cột "Nguồn mủ" thành 2 mảng
               `sales` (mủ thu mua) và `sales_own` (mủ khai thác)
  stock      — Tồn kho: NHIỀU dòng / (đơn vị, ngày) → gom thành 3 khối tồn kho (chưa nhập kho ·
               đã nhập kho · đã ký HĐ) + ô nguyên liệu chưa sản xuất
  plan       — Kế hoạch năm: 1 dòng / (đơn vị, năm)

Ghi có MERGE: nhập Tiêu thụ không xoá Tồn kho của cùng bản ghi ngày đó và ngược lại.
"""

from __future__ import annotations

import io
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.market_meta import UNIT_STOCK_GRADES
from app.core.paths import bulletin_dir
from app.services import member_unit_repo, price_repo, unit_daily_repo
from app.services.unit_daily_fields import FINISHED_TABLE, SALE_TABLES

_BULLETIN = bulletin_dir()
if str(_BULLETIN) not in sys.path:
    sys.path.insert(0, str(_BULLETIN))

from bulletin.convert import r0  # noqa: E402 - 1 nguồn làm tròn nửa-lên dùng chung

TY = 1_000_000_000

CONTRACTS = {"Dài hạn": "long_term", "Chuyến": "spot"}
CHANNELS = {"XK / UTXK": "export", "Nội tiêu": "domestic"}
# Nguồn mủ của dòng tiêu thụ → ghi vào bảng nào (khớp `SALE_TABLES` ở unit_daily_fields).
# File cũ không có cột này: dòng trống mặc định là mủ thu mua (giữ nguyên cách hiểu trước đây).
SALE_SOURCES = {"Mủ thu mua": "sales", "Mủ khai thác": "sales_own"}
# 3 khối tồn kho nhập theo dòng (khối 4 "nguyên liệu chưa sản xuất" là 1 ô riêng, không theo dòng).
CUP_BASES = {"Độ TSC": "tsc", "Độ DRC": "drc"}
CCYS = {"VND": "VND", "USD": "USD"}
STOCK_GROUPS = {
    "Chế biến chưa nhập kho": "not_warehoused",
    "Đã nhập kho": "warehoused",
    "Đã ký HĐ": "signed_undelivered",
}
GRADES = list(UNIT_STOCK_GRADES)


@dataclass
class Col:
    key: str
    title: str
    unit: str = ""
    required: bool = False
    type: str = "num"                       # num | text | date | year | enum
    choices: dict[str, str] | None = None   # enum: {nhãn hiển thị: giá trị lưu}
    width: int = 16


@dataclass
class Spec:
    title: str
    sheet: str
    note: str
    cols: list[Col] = field(default_factory=list)


_UNIT_COL = Col("company", "Đơn vị", required=True, type="text", width=26)
_DATE_COL = Col("as_of", "Ngày", "dd/mm/yyyy", required=True, type="date")

SPECS: dict[str, Spec] = {
    "purchase": Spec(
        "BIỂU NHẬP — THU MUA", "Thu mua",
        "Mủ nước / mủ chén: mỗi đơn vị 1 dòng / 1 ngày. Đơn giá ghi vào kho 'Giá mủ nguyên liệu'. "
        "Mủ nước luôn theo độ TSC; mủ chén theo cột 'Đơn giá mủ chén tính theo'. "
        "THU MUA THÀNH PHẨM tính theo CHỦNG LOẠI: mua mấy chủng loại thì thêm bấy nhiêu dòng cho "
        "cùng (đơn vị, ngày) — các cột mủ nước/mủ chén chỉ điền ở dòng đầu, dòng sau để trống. "
        "File có dòng thành phẩm sẽ GHI ĐÈ toàn bộ phần thành phẩm của ngày đó; không có dòng nào "
        "thì phần thành phẩm đã nhập trên web được giữ nguyên.",
        [_UNIT_COL, _DATE_COL,
         Col("latex_wet", "SL thu mua mủ nước", "tấn"),
         Col("coagulum", "SL thu mua mủ chén", "tấn"),
         Col("price_latex", "Đơn giá mủ nước", "đồng/độ TSC", width=18),
         Col("price_cup", "Đơn giá mủ chén", "đồng/độ", width=18),
         Col("cup_basis", "Đơn giá mủ chén tính theo", "mặc định Độ TSC", type="enum",
             choices=CUP_BASES, width=22),
         Col("finished_grade", "Chủng loại thành phẩm", "chỉ dòng thu mua thành phẩm",
             type="enum", choices={g: g for g in GRADES}, width=22),
         Col("finished_qty", "SL thu mua thành phẩm", "tấn", width=20),
         Col("finished_price", "Đơn giá thành phẩm",
             "triệu đ/tấn khi VND · USD/tấn khi USD", width=26),
         Col("finished_ccy", "Đơn giá thành phẩm bằng", type="enum",
             choices=CCYS, width=20)]),
    "sales": Spec(
        "BIỂU NHẬP — TIÊU THỤ", "Tiêu thụ",
        "Mỗi dòng = 1 hợp đồng bán. Cùng (đơn vị, ngày) có thể nhiều dòng — hệ thống tự gộp. "
        "Cột 'Nguồn mủ' tách mủ thu mua / mủ khai thác (lưu riêng, tổng vẫn cộng chung). "
        "File này GHI ĐÈ TOÀN BỘ phần Tiêu thụ của ngày đó (cả 2 nguồn mủ) — dòng nào không có "
        "trong file sẽ bị xoá; phần Tồn kho giữ nguyên. "
        "Các file đính kèm (bộ Hợp đồng · phiếu xuất kho · hoá đơn) tải lên trên web — "
        "Excel không mang file được.",
        [_UNIT_COL, _DATE_COL,
         Col("code", "Mã HĐ/PL", "số hợp đồng / phụ lục", type="text", width=20),
         Col("source", "Nguồn mủ", "để trống = mủ thu mua", type="enum",
             choices=SALE_SOURCES, width=18),
         Col("contract", "Loại HĐ", required=True, type="enum", choices=CONTRACTS),
         Col("channel", "Hình thức", required=True, type="enum", choices=CHANNELS, width=18),
         Col("grade", "Loại mủ", required=True, type="enum",
             choices={g: g for g in GRADES}, width=20),
         Col("qty", "Số lượng", "tấn"),
         Col("price", "Giá bán", "triệu đ/tấn khi VND · USD/tấn khi USD", width=26),
         Col("sales_ccy", "Giá bán bằng", type="enum", choices=CCYS, width=14),
         Col("warehouse_date", "Ngày xuất kho", "dd/mm/yyyy", type="date", width=18),
         Col("invoice_date", "Ngày xuất hoá đơn", "dd/mm/yyyy", type="date", width=18)]),
    "stock": Spec(
        "BIỂU NHẬP — TỒN KHO", "Tồn kho",
        "Mỗi dòng = 1 dòng tồn kho (số THỜI ĐIỂM cuối ngày, không cộng dồn). "
        "Riêng nhóm 'Đã ký HĐ' là HỢP ĐỒNG có vòng đời: nhập MỘT LẦN, hệ thống tự giữ ở nhóm này "
        "từ cột 'Ngày' (= ngày bắt đầu tồn kho) đến HẾT NGÀY TRƯỚC 'Ngày giao'; chưa giao thì để "
        "trống 'Ngày giao'. KHÔNG nhập lại hợp đồng đó cho các ngày sau. Nhóm này là phần NẰM "
        "TRONG tồn kho thành phẩm đã có hợp đồng nhưng chưa giao — KHÔNG cộng thêm vào tồn kho "
        "(cộng nữa là tính trùng) và cũng không trừ ra, nên không vượt quá tổng tồn kho. "
        "File HĐ scan đính kèm trên web.",
        [_UNIT_COL, _DATE_COL,
         Col("group", "Nhóm", required=True, type="enum", choices=STOCK_GROUPS, width=20),
         Col("grade", "Chủng loại", required=True, type="enum",
             choices={g: g for g in GRADES}, width=20),
         Col("code", "Mã HĐ/PL", "chỉ nhóm đã ký HĐ", type="text", width=20),
         Col("qty", "Số lượng", "tấn"),
         Col("price", "Đơn giá", "chỉ nhóm đã ký HĐ", width=18),
         Col("stock_ccy", "Đơn giá bằng", type="enum", choices=CCYS, width=14),
         Col("delivery_date", "Lịch giao", "dd/mm/yyyy · dự kiến", type="date", width=18),
         Col("delivered_date", "Ngày giao", "dd/mm/yyyy · để trống nếu chưa giao",
             type="date", width=22),
         Col("stock_material", "Tồn kho nguyên liệu chưa sản xuất (quy khô)",
             "tấn", width=34)]),
    "plan": Spec(
        "BIỂU NHẬP — KẾ HOẠCH NĂM", "Kế hoạch năm",
        "Mỗi dòng = 1 đơn vị / 1 năm. Số liệu nhập 1 lần, cập nhật khi có thay đổi.",
        [_UNIT_COL, Col("year", "Năm", required=True, type="year", width=10),
         Col("plan_tonnes", "Kế hoạch thu mua", "tấn", width=20),
         Col("signed_lt_tonnes", "HĐ dài hạn đã ký", "tấn", width=20),
         Col("carry_lt_tonnes", "HĐ dài hạn năm trước chuyển sang", "tấn", width=28),
         Col("carry_spot_tonnes", "HĐ chuyến năm trước chuyển sang", "tấn", width=28)]),
}

_HEAD_FILL = PatternFill("solid", fgColor="D9E7D5")
_REQ_FILL = PatternFill("solid", fgColor="FCE9E7")

_HINT_MAX = 60      # danh sách dài hơn thì chỉ trỏ sang sheet "Danh mục" cho khỏi vỡ dòng


def _hint(c: Col) -> str:
    """Dòng gợi ý dưới tiêu đề.

    Cột chọn phải GHI RÕ giá trị hợp lệ: dropdown chỉ hiện khi bấm vào ô, và nhiều trình xem
    (Quick Look, Numbers, preview trên web) không hiển thị dropdown — người dùng sẽ không biết
    điền gì nếu chỉ dựa vào data validation.
    """
    if c.type != "enum" or not c.choices:
        return c.unit
    labels = " · ".join(c.choices)
    body = labels if len(labels) <= _HINT_MAX else "xem sheet 'Danh mục'"
    return f"{c.unit} · chọn: {body}" if c.unit else f"chọn: {body}"


def template_columns(kind: str, allowed_units: list[str] | None) -> tuple[list[Col], str | None]:
    """Cột của file mẫu + đơn vị mặc định.

    Tài khoản chỉ quản 1 đơn vị → BỎ cột 'Đơn vị' (khỏi gõ tên mình), khi nhập tự gán đơn vị đó.
    """
    spec = SPECS[kind]
    only = allowed_units[0] if allowed_units and len(allowed_units) == 1 else None
    cols = [c for c in spec.cols if not (only and c.key == "company")]
    return cols, only


def build_template(kind: str, allowed_units: list[str] | None = None) -> bytes:
    """Sinh file Excel mẫu: header + đơn vị tính + ô chọn sẵn (dropdown) + sheet Danh mục."""
    spec = SPECS[kind]
    units = allowed_units if allowed_units else member_unit_repo.active_names()
    cols, only = template_columns(kind, allowed_units)
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet

    ws.cell(row=1, column=1, value=spec.title).font = Font(bold=True, size=14)
    note = spec.note + (f"  ·  Áp dụng cho đơn vị: {only}" if only else "")
    ws.cell(row=2, column=1, value=note).font = Font(italic=True, color="666666")
    ws.cell(row=3, column=1, value="Cột tô đỏ là BẮT BUỘC. Không đổi tên/thứ tự cột.").font = Font(
        italic=True, size=9, color="B03A2E")

    head = 5
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=head, column=i, value=c.title)
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.fill = _REQ_FILL if c.required else _HEAD_FILL
        hint = "chọn: xem sheet 'Danh mục'" if (c.key == "company" and units) else _hint(c)
        u = ws.cell(row=head + 1, column=i, value=hint)
        u.font = Font(size=8, italic=True, color="666666")
        u.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = c.width
    ws.row_dimensions[head + 1].height = 30     # đủ chỗ cho gợi ý xuống dòng

    # Danh mục tham chiếu: đơn vị + MỌI cột chọn, để người dùng đọc được giá trị hợp lệ
    # mà không phải bấm vào từng ô.
    ref = wb.create_sheet("Danh mục")
    ranges: dict[str, str] = {}                 # key cột -> vùng tham chiếu cho dropdown
    ref.cell(row=1, column=1, value="Đơn vị").font = Font(bold=True)
    for r, u in enumerate(units, start=2):
        ref.cell(row=r, column=1, value=u)
    ref.column_dimensions["A"].width = 28
    if units:
        ranges["company"] = f"'Danh mục'!$A$2:$A${len(units) + 1}"

    enums = [c for c in cols if c.type == "enum" and c.choices]
    for j, c in enumerate(enums, start=2):
        letter = get_column_letter(j)
        ref.cell(row=1, column=j, value=c.title).font = Font(bold=True)
        for r, label in enumerate(c.choices or {}, start=2):
            ref.cell(row=r, column=j, value=label)
        ref.column_dimensions[letter].width = min(
            max(len(c.title), *(len(x) for x in c.choices or {})) + 2, 30)
        ranges[c.key] = f"'Danh mục'!${letter}$2:${letter}${len(c.choices or {}) + 1}"

    first, last = head + 2, head + 501          # 500 dòng cho người dùng nhập
    for i, c in enumerate(cols, start=1):
        rng = ranges.get(c.key)
        if not rng:
            continue
        dv = DataValidation(type="list", formula1=rng, allow_blank=True)
        ws.add_data_validation(dv)
        letter = get_column_letter(i)
        dv.add(f"{letter}{first}:{letter}{last}")

    ws.freeze_panes = ws.cell(row=first, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _as_date(v: Any) -> str | None:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _as_num(v: Any) -> float | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    # Người dùng hay gõ kiểu vi-VN: "1.234,5"
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_upload(kind: str, data: bytes,
                 allowed_units: list[str] | None = None) -> dict[str, Any]:
    """Đọc file người dùng nộp → danh sách dòng đã chuẩn hoá + lỗi từng dòng (để XEM TRƯỚC)."""
    spec = SPECS[kind]
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:                                  # noqa: BLE001
        raise ValueError("Không đọc được file — hãy dùng đúng file mẫu (.xlsx).") from exc
    ws = wb[spec.sheet] if spec.sheet in wb.sheetnames else wb.worksheets[0]

    units = set(member_unit_repo.active_names())
    allowed = set(allowed_units) if allowed_units is not None else None
    default_company = allowed_units[0] if allowed_units and len(allowed_units) == 1 else None

    # Dò cột theo TIÊU ĐỀ (không theo vị trí cố định) → chịu được khi thiếu/đổi chỗ cột.
    head = 5
    titles = {}
    for i in range(1, ws.max_column + 1):
        t = str(ws.cell(row=head, column=i).value or "").strip()
        if t:
            titles.setdefault(t, i)
    idx: dict[str, int | None] = {c.key: titles.get(c.title) for c in spec.cols}
    missing = [c.title for c in spec.cols
               if c.required and idx[c.key] is None and not (c.key == "company" and default_company)]
    if missing:
        raise ValueError("File thiếu cột bắt buộc: " + ", ".join(missing) + ". Hãy dùng đúng file mẫu.")
    rows: list[dict[str, Any]] = []
    for r in range(head + 2, ws.max_row + 1):
        raw = [ws.cell(row=r, column=idx[c.key]).value if idx[c.key] else None for c in spec.cols]
        if all(v is None or str(v).strip() == "" for v in raw):
            continue
        rec: dict[str, Any] = {"_row": r, "_errors": []}
        for c, v in zip(spec.cols, raw):
            if c.key == "company" and idx["company"] is None and default_company:
                rec["company"] = default_company     # mẫu không có cột Đơn vị → gán đơn vị của tài khoản
                continue
            if c.type == "date":
                val = _as_date(v)
                if val is None and (c.required or (v not in (None, ""))):
                    rec["_errors"].append(f"{c.title}: ngày không hợp lệ")
            elif c.type == "year":
                n = _as_num(v)
                val = int(n) if n else None
                if val is None or not (2020 <= val <= 2100):
                    val = None
                    rec["_errors"].append(f"{c.title}: năm không hợp lệ")
            elif c.type == "enum":
                label = str(v or "").strip()
                val = (c.choices or {}).get(label)
                if val is None and label:
                    rec["_errors"].append(f"{c.title}: '{label}' không hợp lệ")
                elif val is None and c.required:
                    rec["_errors"].append(f"{c.title}: bắt buộc")
            elif c.type == "text":
                val = str(v).strip() if v not in (None, "") else None
            else:
                val = _as_num(v)
                if val is None and v not in (None, ""):
                    rec["_errors"].append(f"{c.title}: không phải số")
            rec[c.key] = val
            if c.required and rec.get(c.key) in (None, "") and not rec["_errors"]:
                rec["_errors"].append(f"{c.title}: bắt buộc")

        if rec.get("company") and rec["company"] not in units:
            rec["_errors"].append(f"Đơn vị '{rec['company']}' không có trong danh mục")
        elif allowed is not None and rec.get("company") and rec["company"] not in allowed:
            rec["_errors"].append(f"Đơn vị '{rec['company']}' không thuộc quyền của tài khoản")
        rows.append(rec)

    _mark_actions(kind, rows)
    ok = sum(1 for r in rows if not r["_errors"])
    return {
        "kind": kind, "rows": rows,
        # Nhãn cột tiếng Việt để bảng xem trước hiển thị dễ đọc (không phải khoá thô).
        "columns": [{"key": c.key, "title": c.title, "unit": c.unit} for c in spec.cols
                    if not (c.key == "company" and idx["company"] is None and default_company)],
        "summary": {"total": len(rows), "ok": ok, "error": len(rows) - ok},
    }


def _mark_actions(kind: str, rows: list[dict]) -> None:
    """Đánh dấu mỗi dòng sẽ TẠO MỚI hay GHI ĐÈ số đã có (để người dùng biết trước khi xác nhận)."""
    if kind == "plan":
        for r in rows:
            if r.get("company") and r.get("year"):
                cur = unit_daily_repo.year_plan(int(r["year"]), [r["company"]])
                r["_action"] = "update" if cur.get(r["company"]) else "create"
        return
    daily_kind = "purchase" if kind == "purchase" else "consumption"
    seen: dict[tuple[str, str], bool] = {}
    for r in rows:
        c, d = r.get("company"), r.get("as_of")
        if not c or not d:
            continue
        key = (c, d)
        if key not in seen:
            seen[key] = unit_daily_repo.has_entry(daily_kind, d, c)
        r["_action"] = "update" if seen[key] else "create"


def _line_revenue_vnd(qty, price, ccy: str, fx: float | None) -> float | None:
    """Doanh thu 1 dòng bán → BASE = đồng. Cùng công thức với web (`lineRevenueVnd`) để mở phiếu
    ra lưu lại KHÔNG đổi số. VND: giá là triệu đ/tấn. USD: cần tỷ giá, thiếu → None (không đoán)."""
    q, p = _as_num(qty), _as_num(price)
    if q is None or p is None:
        return 0.0                       # dòng chưa điền số → không đóng góp doanh thu
    if ccy != "USD":
        return q * p * 1_000_000
    return None if fx is None else q * p * fx


def _upsert_contract(r: dict, company: str, start_date: str, ccy: str | None,
                     username: str | None) -> None:
    """1 dòng Excel nhóm 'Đã ký HĐ' → thêm/cập nhật hợp đồng (bảng `unit_stock_contract`).

    Khớp lại hợp đồng cũ theo (đơn vị, mã HĐ/PL, chủng loại, ngày bắt đầu) để nhập lại cùng file
    KHÔNG sinh bản sao — nhập lại là SỬA, đúng như cách các biểu khác ghi đè theo (đơn vị, ngày).
    """
    from app.services import unit_stock_contract_repo

    existing = unit_stock_contract_repo.list_contracts(companies=[company])
    match = next((c for c in existing
                  if c["start_date"] == start_date and c["grade"] == r["grade"]
                  and (c["code"] or "") == (r.get("code") or "")), None)
    unit_stock_contract_repo.save({
        "id": match["id"] if match else None,
        "code": r.get("code"), "grade": r["grade"], "qty": r.get("qty"), "price": r.get("price"),
        "ccy": ccy or "VND", "fx": match.get("fx") if match else None,
        "start_date": start_date, "delivery_date": r.get("delivery_date"),
        "delivered_date": r.get("delivered_date"),
        "file": match.get("file") if match else None,        # file HĐ scan chỉ đính kèm trên web
        "filename": match.get("filename") if match else None,
    }, company, username)


def _sale_line(r: dict, ccy: str, fx: float | None) -> dict:
    """1 dòng Excel → 1 dòng bán để lưu.

    Gán loại tiền + tỷ giá xuống TỪNG DÒNG (Excel chỉ hỏi 1 lần cho cả file) để số đã lưu khớp
    đúng cách web tính lại doanh thu — mở phiếu ra lưu lại không bị đổi số.
    """
    return {
        "code": r.get("code"),
        "contract": r["contract"], "channel": r["channel"], "grade": r["grade"],
        "qty": r.get("qty"), "price": r.get("price"), "ccy": ccy, "fx": fx,
        "warehouse_date": r.get("warehouse_date"), "invoice_date": r.get("invoice_date"),
    }


def commit_rows(kind: str, rows: list[dict], username: str | None,
                allowed_units: list[str] | None = None) -> dict[str, Any]:
    """Ghi các dòng HỢP LỆ vào hệ thống (bỏ qua dòng có lỗi). Trả số bản ghi đã ghi."""
    # Client gửi lại danh sách dòng nên phải KIỂM LẠI ở đây (không tin bước xem trước):
    # đơn vị phải có thật và thuộc quyền tài khoản.
    valid_units = set(member_unit_repo.active_names())
    if allowed_units is not None:
        valid_units &= set(allowed_units)
    good = [r for r in rows
            if not (r.get("_errors") or []) and r.get("company") in valid_units]
    if kind == "plan":
        n = 0
        for r in good:
            unit_daily_repo.set_year_plan(int(r["year"]), r["company"], r.get("plan_tonnes"),
                                          r.get("signed_lt_tonnes"), r.get("carry_lt_tonnes"),
                                          r.get("carry_spot_tonnes"), username)
            n += 1
        return {"saved": n, "skipped": len(rows) - len(good), "warnings": []}

    currencies = member_unit_repo.currency_by_name()
    warnings: list[str] = []

    # Gom theo (đơn vị, ngày)
    grouped: dict[tuple[str, str], list[dict]] = {}
    for r in good:
        if r.get("company") and r.get("as_of"):
            grouped.setdefault((r["company"], r["as_of"]), []).append(r)

    saved = 0
    for (company, as_of), items in grouped.items():
        if kind == "purchase":
            # Mủ nước/mủ chén là số CỦA NGÀY (1 giá trị), thành phẩm là NHIỀU DÒNG theo chủng loại
            # → lấy ô đầu tiên có số cho phần theo ngày, gom mọi dòng có chủng loại cho thành phẩm.
            first = lambda k: next((r.get(k) for r in items if r.get(k) is not None), None)  # noqa: E731
            it = {k: first(k) for k in ("latex_wet", "coagulum", "price_latex", "price_cup")}
            # MERGE: giữ các ô đã nhập trên web mà file không có (cờ không thu mua, đơn giá nội tệ…).
            cur = unit_daily_repo.entries_on("purchase", as_of).get(company) or {}
            fields = dict(cur.get("fields") or {})
            fields.update({k: v for k, v in it.items()
                           if v is not None and k in ("latex_wet", "coagulum")})
            # Cách tính độ của mủ chén phải gán TRƯỚC khi ghi, không thì không được lưu.
            basis = first("cup_basis") or fields.get("cup_basis")
            if basis:
                fields["cup_basis"] = basis
            # Thành phẩm: tỷ giá không có trong file → lấy lại tỷ giá đã nhập trên web (nếu có).
            old_fx = next((r.get("fx") for r in (fields.get(FINISHED_TABLE) or [])
                           if isinstance(r, dict) and r.get("fx")), None)
            finished = [{"grade": r["finished_grade"], "qty": r.get("finished_qty"),
                         "price": r.get("finished_price"),
                         "ccy": (ccy := r.get("finished_ccy") or "VND"),
                         "fx": old_fx if ccy == "USD" else None}
                        for r in items if r.get("finished_grade")]
            if finished:
                fields[FINISHED_TABLE] = finished
                if old_fx is None and any(r["ccy"] == "USD" for r in finished):
                    warnings.append(
                        f"{company} {as_of}: có dòng thành phẩm bằng USD chưa có tỷ giá "
                        f"(nhập tỷ giá ở màn Thu mua rồi lưu lại).")
            unit_daily_repo.upsert("purchase", as_of, company, fields, username)
            for key, ptype in (("price_latex", "purchase"), ("price_cup", "purchase_cup")):
                if it.get(key) is not None:
                    price_repo.upsert_record({
                        "as_of": as_of, "source": "vrg", "grade": company, "contract": "",
                        "price_type": ptype, "price": float(it[key]),
                        "currency": "VND",
                        "unit": ("đồng/độ DRC" if ptype == "purchase_cup" and basis == "drc"
                                 else "đồng/độ TSC")})
        else:
            # MERGE: giữ nguyên phần còn lại của bản ghi ngày đó.
            cur = unit_daily_repo.entries_on("consumption", as_of).get(company) or {}
            fields = dict(cur.get("fields") or {})
            if kind == "sales":
                # "Giá bán bằng" là cột của TỪNG DÒNG → loại tiền theo từng dòng, y như form web
                # (một ngày có thể vừa bán USD vừa bán VNĐ). Dòng để trống thì lấy loại tiền của
                # bản ghi, cuối cùng mới suy từ đơn vị (nước ngoài → USD).
                default_ccy = (next((r.get("sales_ccy") for r in items if r.get("sales_ccy")), None)
                               or fields.get("sales_ccy")
                               or ("VND" if currencies.get(company, "VND") == "VND" else "USD"))
                fields["sales_ccy"] = default_ccy   # loại tiền mặc định cho dòng thêm mới trên web
                fx = _as_num(fields.get("fx_revenue"))
                # File Tiêu thụ mang CẢ 2 nguồn mủ (cột "Nguồn mủ") → ghi đè trọn phần tiêu thụ;
                # nguồn nào không có dòng nào trong file thì thành rỗng.
                lines: dict[str, list[dict]] = {t: [] for t in SALE_TABLES}
                revenue, missing_fx = 0.0, False
                for r in items:
                    rc = r.get("sales_ccy") or default_ccy
                    table = r.get("source") or "sales"
                    if table not in lines:
                        table = "sales"
                    lines[table].append(_sale_line(r, rc, fx if rc == "USD" else None))
                    got = _line_revenue_vnd(r.get("qty"), r.get("price"), rc, fx)
                    if got is None:
                        missing_fx = True      # dòng USD thiếu tỷ giá → không cộng, KHÔNG đoán bừa
                    else:
                        revenue += got
                for table in SALE_TABLES:
                    fields[table] = lines[table]
                fields["revenue"] = r0(revenue)   # nửa LÊN như mọi số tiền khác của dự án
                if missing_fx:
                    warnings.append(
                        f"{company} {as_of}: có dòng bán bằng USD chưa có tỷ giá nên chưa cộng "
                        f"vào doanh thu (nhập tỷ giá ở màn Báo cáo tiêu thụ rồi lưu lại).")
            else:
                # Tồn kho = số THỜI ĐIỂM, đơn vị TẤN (mẫu tuần mục 11–14).
                sccy = next((r.get("stock_ccy") for r in items if r.get("stock_ccy")), None)
                if sccy:
                    fields["stock_ccy"] = sccy
                pick = lambda g: [r for r in items if r.get("group") == g]  # noqa: E731
                fields["stock_not_warehoused"] = [
                    {"grade": r["grade"], "qty": r.get("qty")} for r in pick("not_warehoused")]
                fields["stock_warehoused"] = [
                    {"grade": r["grade"], "qty": r.get("qty")} for r in pick("warehoused")]
                # Nhóm "Đã ký HĐ" KHÔNG nằm trong payload ngày nữa — mỗi dòng là 1 HỢP ĐỒNG có
                # vòng đời riêng: cột "Ngày" = ngày bắt đầu tồn kho. Trùng (đơn vị, mã HĐ, chủng
                # loại, ngày bắt đầu) thì cập nhật lại chính hợp đồng đó, không tạo bản sao.
                for r in pick("signed_undelivered"):
                    try:
                        _upsert_contract(r, company, as_of, sccy, username)
                    except ValueError as exc:
                        warnings.append(f"{company} {as_of}: {exc}")
                mat = next((r.get("stock_material") for r in items
                            if r.get("stock_material") is not None), None)
                if mat is not None:
                    fields["stock_material"] = mat
            unit_daily_repo.upsert("consumption", as_of, company, fields, username)
        saved += 1
    return {"saved": saved, "skipped": len(rows) - len(good), "warnings": warnings}
