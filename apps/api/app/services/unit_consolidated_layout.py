"""KHUÔN của biểu TỔNG HỢP gửi Tập đoàn: hằng số cột · nhóm chủng loại · 5 dòng tiêu đề.

Phần dựng số liệu nằm ở `unit_consolidated_excel` (tách đôi cho mỗi file một việc).

Bám đúng file mẫu Ban TTKD đang dùng ("Tuần 37-Tổng hợp BC tiêu thụ VRG 17.09.2026"): tiêu đề 4
tầng, mỗi ĐƠN VỊ một dòng gom theo KHU VỰC (I…VI), dòng tổng của từng khu vực và dòng TẬP ĐOÀN
trên cùng, khối ghi chú ở cuối.

Số liệu lấy nguyên từ `unit_period_report.period_report("consumption", 01/01 → ngày chốt)` — cùng
một nguồn với màn Báo cáo theo kỳ, nên hai chỗ không thể lệch nhau. Các cột SUY RA (i, j, k, l, m,
và cột "chưa có hợp đồng") ghi bằng CÔNG THỨC Excel như mẫu: người dùng sửa một ô là tổng tự đổi.

Ba điểm phải biết khi đọc số:
- **Tồn kho thành phẩm (Q)** = khối 1 + khối 2 của biểu Tồn kho; **đã có HĐ (R)** = phần đã ký hợp
  đồng chưa giao, tính từ hợp đồng chứ không từ lời khai tồn kho. Hai nguồn khác nhau nên R có thể
  lớn hơn Q ở đơn vị chốt tồn kho trễ — lúc đó cột chủng loại bị kẹp về 0 (không in số âm).
- **Chi tiết chủng loại (S…AB)** = tồn kho của chủng loại đó TRỪ phần đã ký hợp đồng của chính nó.
- **KH Sản xuất + Thu mua (D/E/F)** để trống đúng như mẫu (chốt với chủ dự án 22/09/2026).
"""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

#: 10 nhóm chủng loại của biểu (cột S…AB) → các chủng loại trong hệ thống.
#: Chủng loại không nằm trong bảng này (kể cả mủ nguyên liệu khai nhầm vào tồn thành phẩm) dồn vào
#: "Ngoại lệ" — thà hiện ở cột ngoại lệ còn hơn rơi mất sản lượng khỏi biểu.
GRADE_BUCKETS: list[tuple[str, tuple[str, ...]]] = [
    ("CV50/\n60", ("SVR CV 50", "SVR CV60")),
    ("SVR10CV/\n20CV", ("SVR 10CV / 20CV",)),
    ("SVRL/\n3L& 3L Mix", ("SVR L", "SVR 3L", "SVR 3L Mix")),
    ("RSS", ("RSS 1", "RSS 3", "RSS 5")),
    ("SVR5/\n5S", ("SVR 5", "SVR 5S")),
    ("SVR10/\n20", ("SVR 10 / CSR 10", "SVR 20 / CSR 20")),
    ("Latex\n(quy khô)", ("LATEX",)),
    ("Ngoại\nlệ", ("Mủ ngoại lệ", "Chủng loại khác")),
    ("Skim", ("Skim Block",)),
    ("SVR10\nMIX", ("SVR 10 Mix",)),
]
_OTHER_BUCKET = 7                      # chỉ số cột "Ngoại lệ" — nơi dồn chủng loại lạ

#: Thứ tự khu vực trên biểu + số La Mã (đúng thứ tự file mẫu).
REGION_ORDER: list[str] = ["Đông Nam Bộ", "Tây Nguyên", "Duyên Hải Miền Trung",
                           "Miền Núi phía Bắc", "Campuchia", "Lào"]
_ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
_OTHER_REGION = "Khu vực khác"

NOTES: list[str] = [
    "Số liệu báo cáo chốt vào thứ 5 hàng tuần, thời gian gửi báo cáo vào thứ 6 hàng tuần.",
    "Tồn kho thành phẩm là: số lượng bao gồm chưa có HĐ và đã có HĐ (đã có HĐ là gồm SL hàng đã ký "
    "HĐ nhưng chưa giao hàng, đối với HĐDH thì số lượng ký đã có xác định giá bán).",
    "Nguyên liệu: không phải sản phẩm tồn kho.",
    "Cột “Chưa có hợp đồng” theo chủng loại = tồn kho chủng loại đó trừ phần đã ký hợp đồng chưa "
    "giao; đơn vị nào chốt tồn kho trễ hơn ngày ký hợp đồng thì phần chênh được kẹp về 0.",
]

_NUM = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
_FONT = "Times New Roman"
_THIN = Side(style="thin", color="808080")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="D9E7D5")
_GROUP_FILL = PatternFill("solid", fgColor="CCFF99")
_TOTAL_FILL = PatternFill("solid", fgColor="FFE699")
_WIDTHS = {"A": 4, "B": 5, "C": 45, "D": 9, "E": 9, "F": 10, "G": 12, "H": 12, "I": 10, "J": 12,
           "K": 12, "L": 11, "M": 11, "N": 12, "O": 12, "P": 11, "Q": 10, "R": 11, "AC": 11,
           "AD": 12, "AE": 14, "AF": 14, "AG": 22}

# Cột lấy thẳng từ báo cáo kỳ: {chữ cái cột: khoá dữ liệu}. Các cột còn lại là công thức hoặc trống.
_VALUE_COLS: dict[str, str] = {
    "G": "signed_lt_tonnes", "H": "lt_export", "I": "lt_domestic",
    "J": "spot_export", "K": "spot_domestic",
    "Q": "stock_finished", "R": "stock_finished_hd", "AD": "stock_material",
    "AE": "carry_lt_tonnes", "AF": "carry_spot_tonnes",
}
_GRADE_COLS = [get_column_letter(19 + i) for i in range(len(GRADE_BUCKETS))]     # S…AB
#: Cột được CỘNG ở dòng khu vực / Tập đoàn (cột công thức tự tính lại, không cộng).
_SUM_COLS = ["G", "H", "I", "J", "K", "Q", "R", *_GRADE_COLS, "AD", "AE", "AF"]


def row_formulas(r: int) -> dict[str, str]:
    """Các cột suy ra của MỘT dòng bất kỳ (đơn vị · khu vực · Tập đoàn) — đúng công thức của mẫu."""
    return {"L": f"=H{r}+I{r}", "M": f"=J{r}+K{r}", "N": f"=H{r}+J{r}", "O": f"=I{r}+K{r}",
            "P": f"=SUM(H{r}:K{r})", "AC": f"=SUM(S{r}:AB{r})"}


def grade_cells(row: dict[str, Any]) -> list[float | None]:
    """Tồn kho CHƯA CÓ HỢP ĐỒNG theo 10 nhóm chủng loại của biểu.

    = tồn kho của chủng loại − phần đã ký hợp đồng chưa giao của chính chủng loại đó, kẹp ≥ 0
    (hai số đến từ hai nguồn khác nhau nên hiệu có thể âm khi đơn vị chốt tồn kho trễ).
    """
    stock = row.get("stock_by_grade") or {}
    signed = row.get("stock_hd_by_grade") or {}
    out = [0.0] * len(GRADE_BUCKETS)
    known = {g: i for i, (_, grades) in enumerate(GRADE_BUCKETS) for g in grades}
    for grade, qty in stock.items():
        left = (qty or 0.0) - (signed.get(grade) or 0.0)
        out[known.get(grade, _OTHER_BUCKET)] += max(0.0, left)
    return [round(v, 3) or None for v in out]


def write(ws, r: int, col: str, value: Any, *, bold=False, fill=None, fmt=_NUM) -> None:
    c = ws[f"{col}{r}"]
    # Làm tròn 3 số lẻ NGAY KHI GHI: số trong DB là float nên hay ra đuôi 3753.719999999998,
    # người nhận file thấy là tưởng số liệu bẩn (mẫu của Ban TTKD cũng để 2-3 số lẻ).
    c.value = round(value, 3) if isinstance(value, float) else value
    c.font = Font(name=_FONT, size=9, bold=bold)
    c.border = _BORDER
    c.alignment = Alignment(vertical="center", wrap_text=False)
    if isinstance(value, (int, float)) or (isinstance(value, str) and value.startswith("=")):
        c.number_format = fmt
    if fill:
        c.fill = fill


def header(ws, as_of: str, year: int) -> None:
    """5 dòng đầu: tiêu đề + 3 tầng tiêu đề cột + dòng chữ cái (a, b, c…) như mẫu."""
    d, m, y = as_of[8:10], as_of[5:7], as_of[:4]
    ws.merge_cells("A1:AG1")
    title = ws["A1"]
    title.value = (f"TỔNG HỢP BÁO CÁO\n SẢN XUẤT - TỒN KHO - TIÊU THỤ \nNĂM {year}\n"
                   f"Báo cáo lũy kế đến ngày {d}/{m}/{y}")
    title.font = Font(name=_FONT, size=11, bold=True)
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    tall = {"A": "TT", "B": "TT", "C": "Khu vực / Đơn vị",
            "G": f"Tổng \nSố Lượng đã ký theo CT HĐDH {year}\n(Số liệu lũy kế)",
            "Q": "Tồn kho thành phẩm (tấn)",
            "R": "Trong đó, thành phẩm đã có Hợp đồng (tấn)",
            "AC": "Số lượng chưa có hợp đồng",
            "AD": "Tồn kho Nguyên liệu chưa có HĐ (đối với các đơn vị chưa có nhà máy chế biến)",
            "AE": f"Số lượng\n tiêu thụ HĐDH {year - 1} chuyển sang {year}",
            "AF": f"Số lượng \ntiêu thụ \nHĐ CHUYẾN \n{year - 1} chuyển sang {year}",
            "AG": "Ghi chú"}
    for col, text in tall.items():
        ws.merge_cells(f"{col}2:{col}4")
        write(ws, 2, col, text, bold=True, fill=_HEAD_FILL)

    ws.merge_cells("D2:F2")
    write(ws, 2, "D", f"KH Sản xuất + Thu mua\n{year}", bold=True, fill=_HEAD_FILL)
    ws.merge_cells("H2:P2")
    write(ws, 2, "H", "Tiêu thụ", bold=True, fill=_HEAD_FILL)
    ws.merge_cells("S2:AB2")
    write(ws, 2, "S", "Số lượng thành phẩm chưa có HĐ \n(Chi tiết từng chủng loại)\n(Tấn)",
           bold=True, fill=_HEAD_FILL)

    for col, text in (("D", "Khai\n thác"), ("E", "Thu\n mua "), ("F", "Cộng")):
        ws.merge_cells(f"{col}3:{col}4")
        write(ws, 3, col, text, bold=True, fill=_HEAD_FILL)
    ws.merge_cells("H3:I3")
    write(ws, 3, "H", "Tổng\n Số Lượng \nHĐ DÀI HẠN\nđã thực hiện \n(Lũy kế từ đầu năm)",
           bold=True, fill=_HEAD_FILL)
    ws.merge_cells("J3:K3")
    write(ws, 3, "J", "Tổng \nSố Lượng\n HĐ CHUYẾN\nđã thực hiện\n(Lũy kế từ đầu năm)",
           bold=True, fill=_HEAD_FILL)
    # L…P KHÔNG gộp dòng 3-4: dòng 3 là tên cột, dòng 4 là công thức của mẫu (i=(e+f)…).
    for col, text in (("L", "CỘNG \nSố lượng \nHĐ DÀI HẠN"), ("M", "CỘNG \nSố lượng \nHĐ CHUYẾN"),
                      ("N", "CỘNG \nSố lượng \n(XK UTXK) của DH & chuyến"),
                      ("O", "CỘNG \nSố lượng \n(Nội tiêu) của DH & chuyến"),
                      ("P", "CỘNG\nSố lượng\n HĐ DÀI HẠN và HĐ CHUYẾN \n(Tổng tiêu thụ)")):
        write(ws, 3, col, text, bold=True, fill=_HEAD_FILL)
    for i, (label, _) in enumerate(GRADE_BUCKETS):
        col = _GRADE_COLS[i]
        ws.merge_cells(f"{col}3:{col}4")
        write(ws, 3, col, label, bold=True, fill=_HEAD_FILL)
    for col, text in (("H", "XK UTXK"), ("I", "Nội tiêu"), ("J", "XK UTXK"), ("K", "Nội Tiêu"),
                      ("L", "i=(e +f)"), ("M", "j= (g + h)"), ("N", "k=(e + g)"),
                      ("O", "l=(f +h)"), ("P", "m=(e+f+g+h)")):
        write(ws, 4, col, text, bold=True, fill=_HEAD_FILL)

    letters = ["a", "b", "c", "c", "d", "d", "d", "e", "f", "g", "h", "i", "j", "k", "l", "m",
               "n", "o", "p", "q", "r", "s", "t", "u", "v", "w", "y", "z",
               "(1)=(p+q+r+s+t+u+v+w+y+z)"]
    for i, text in enumerate(letters, start=1):
        write(ws, 5, get_column_letter(i), text, bold=False, fill=_HEAD_FILL)
    for cell in (2, 3, 4, 5):
        for i in range(1, 34):
            c = ws.cell(row=cell, column=i)
            c.border = _BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if not c.font.name:
                c.font = Font(name=_FONT, size=9, bold=True)
    ws.row_dimensions[1].height = 62
    ws.row_dimensions[3].height = 79
