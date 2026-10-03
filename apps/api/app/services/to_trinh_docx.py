"""Xuất TỜ TRÌNH giá sàn ra Word (.docx) theo mẫu mới — cùng khung chữ với bản HTML/PDF
([to_trinh_view]) nên hai bản luôn khớp. Word để Ban TTKD chỉnh câu chữ lần cuối, trình ký.

Định dạng theo mẫu 54/TTr-TTKD: A4, Times New Roman 13, lề trái 3 cm, đoạn thụt đầu dòng 1 cm, căn đều.
"""
from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

FONT = "Times New Roman"
HEAD_FILL = "EAF1DD"
_C, _L, _R, _J = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT,
                  WD_ALIGN_PARAGRAPH.JUSTIFY)


def _run(p, text: str, *, bold=False, italic=False, size: float | None = None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    return r


def _p(doc_or_cell, text: str = "", *, align=_J, indent=True, bold=False, italic=False, size=None,
       space_after: float = 3):
    p = doc_or_cell.add_paragraph()
    p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before, fmt.space_after = Pt(0), Pt(space_after)
    if indent:
        fmt.first_line_indent = Cm(1)
    if text:
        _run(p, text, bold=bold, italic=italic, size=size)
    return p


def _para(doc, para: dict[str, str], lead_bold_italic: bool = True) -> None:
    if not (para.get("lead") or para.get("text")):
        return
    p = _p(doc)
    if para.get("lead"):
        _run(p, para["lead"] + " ", bold=True, italic=lead_bold_italic)
    _run(p, para.get("text", ""))


def _shade(cell, fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _cell(cell, text: str, *, align=_C, bold=False, size: float = 12, fill: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate((text or "").split("\n")):
        if i:
            _run(p, "\n", size=size)
        _run(p, line, bold=bold, size=size)
    if fill:
        _shade(cell, fill)


def _widths(t, cms: list[float]) -> None:
    """Độ rộng cột cố định (Word tự co cột theo chữ → 'FUTURES' vỡ dòng)."""
    t.autofit = False
    for col, w in zip(t.columns, cms):
        col.width = Cm(w)                     # w:gridCol — LibreOffice đọc chỗ này
    for row in t.rows:
        for cell, w in zip(row.cells, cms):
            cell.width = Cm(w)                # w:tcW — Word đọc chỗ này


def _grid(doc, rows: int, cols: int, cms: list[float]):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _widths(t, cms)
    return t


def _head(t, row: int, col: int, text: str, *, down: int = 0, right: int = 0) -> None:
    cell = t.cell(row, col)
    if down or right:
        cell = cell.merge(t.cell(row + down, col + right))
    _cell(cell, text, bold=True, fill=HEAD_FILL)


def _settlement(doc, v: dict) -> None:
    t = _grid(doc, 2 + len(v["sett"]), 8, [1.3, 2.6, 2.0, 1.9, 2.3, 2.3, 1.8, 1.5])
    for c, txt in enumerate(["STT", "SÀN\nFUTURES", "Chủng loại", "Đơn vị\ntính", f"Giá\n({v['d2']})",
                             f"Giá\n({v['d1']})"]):
        _head(t, 0, c, txt, down=1)
    _head(t, 0, 6, "Thay đổi", right=1)
    _head(t, 1, 6, "USD/T")
    _head(t, 1, 7, "%")
    for i, r in enumerate(v["sett"], start=2):
        if r["span"]:
            _cell(t.cell(i, 0).merge(t.cell(i + r["span"] - 1, 0)), r["stt"])
            _cell(t.cell(i, 1).merge(t.cell(i + r["span"] - 1, 1)), r["san"])
        prev = r["prev"] + (f" {r['prev_tag']}" if r["prev_tag"] else "")
        curr = r["curr"] + (f" {r['curr_tag']}" if r["curr_tag"] else "")
        for c, txt in enumerate([r["grade"], r["unit"], prev, curr, r["d"], r["pct"]], start=2):
            _cell(t.cell(i, c), txt)


def _physical(doc, v: dict) -> None:
    t = _grid(doc, 2 + len(v["phys"]), 5, [4.6, 2.9, 2.9, 2.7, 2.6])
    for c, txt in enumerate(["SÀN\nPHYSICAL", f"Giá\n({v['d2']})", f"Giá\n({v['d1']})"]):
        _head(t, 0, c, txt, down=1)
    _head(t, 0, 3, "Thay đổi", right=1)
    _head(t, 1, 3, "USD/T")
    _head(t, 1, 4, "%")
    for i, r in enumerate(v["phys"], start=2):
        _cell(t.cell(i, 0), r["grade"], align=_L)
        for c, txt in enumerate([r["prev"], r["curr"], r["d"], r["pct"]], start=1):
            _cell(t.cell(i, c), txt)


def _proposal(doc, v: dict) -> None:
    t = _grid(doc, 1 + len(v["prop"]), 5, [3.3, 3.0, 3.1, 3.2, 3.1])
    lan, prev = v["prop_lan"], v["prop_prev"]
    for c, txt in enumerate(["Chủng Loại", f"Giá dự kiến\n{lan}", f"(+/-)\n(usd/tấn)\nso với\n{prev}",
                             f"Giá dự kiến\n{lan}", f"(+/-)\n(đ/tấn)\nso với\n{prev}"]):
        _cell(t.cell(0, c), txt, bold=True, fill=HEAD_FILL)
    for i, r in enumerate(v["prop"], start=1):
        _cell(t.cell(i, 0), r["grade"], align=_L)
        for c, k in enumerate(["fob", "fob_d", "vnd", "vnd_d"], start=1):
            _cell(t.cell(i, c), r[k], align=_R)


def _header(doc, v: dict) -> None:
    t = doc.add_table(rows=2, cols=2)
    _widths(t, [6.4, 9.6])
    left, right = t.cell(0, 0), t.cell(0, 1)
    for cell, lines, rule in ((left, v["org"], "––––––––"), (right, v["nation"], "––––––––––––––––––––––––")):
        cell.text = ""
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.alignment, p.paragraph_format.space_after = _C, Pt(0)
            _run(p, line, bold=True, size=12.5)
        _p(cell, rule, align=_C, indent=False, space_after=0)
    _cell(t.cell(1, 0), v["dept"], bold=True, size=13)
    _p(t.cell(1, 0), v["so"], align=_C, indent=False, space_after=0)
    t.cell(1, 1).text = ""
    _run(t.cell(1, 1).paragraphs[0], v["place_date"], italic=True)
    t.cell(1, 1).paragraphs[0].alignment = _C


def _signers(doc, s: dict) -> None:
    """Hàng chức danh · hàng tên (chừa chỗ ký) riêng → tên hai bên luôn thẳng hàng dù chức danh dài."""
    t = doc.add_table(rows=4, cols=2)
    _widths(t, [8.0, 8.0])
    for col, (role, name) in enumerate([(s["left_role"], s["left_name"]), (s["right_role"], s["right_name"])]):
        _cell(t.cell(0, col), role, bold=True, size=12.5)
        _cell(t.cell(1, col), name, bold=True, size=13)
    _cell(t.cell(2, 0).merge(t.cell(2, 1)), s["approver_role"], bold=True, size=12.5)
    _cell(t.cell(3, 0).merge(t.cell(3, 1)), s["approver_name"], bold=True, size=13)
    for r, row in enumerate(t.rows):
        for cell in row.cells:
            fmt = cell.paragraphs[0].paragraph_format
            fmt.keep_with_next = r < 3        # cả khối ký nằm trọn một trang
            if r in (1, 3):
                fmt.space_before = Pt(72)


def build_docx(v: dict[str, Any]) -> bytes:
    """`v` = to_trinh_view.build(...) → bytes .docx."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2), Cm(1.8), Cm(3), Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(13)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    _header(doc, v)
    _p(doc, "TỜ TRÌNH", align=_C, indent=False, bold=True, size=14, space_after=0)
    _p(doc, v["subject"], align=_C, indent=False, space_after=0)
    _p(doc, "––––––––––––––––––––––––", align=_C, indent=False, space_after=0)
    _p(doc, v["to"], align=_C, indent=False)
    _p(doc, v["sec1"], bold=True)
    _settlement(doc, v)
    if v["futures_note"]:
        _p(doc, v["futures_note"], italic=True, size=10, indent=False)
    for para in v["futures"]:
        _para(doc, para)
    _p(doc, v["sec2"], bold=True)
    _physical(doc, v)
    if v["physical_note"]:
        _p(doc, v["physical_note"])
    for para in [*v["physical"], *v["outlook"]]:
        _para(doc, para)
    _para(doc, v["inventory"], lead_bold_italic=False)
    _p(doc, v["intro"])
    _proposal(doc, v)
    _p(doc, v["closing"]).paragraph_format.keep_with_next = True
    _signers(doc, v["signers"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
