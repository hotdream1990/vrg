"""Test Lịch sử truy cập: ghi đăng nhập · ghi lượt vào trang (có chống trùng) · tra cứu · tổng hợp."""

from __future__ import annotations

import pytest
from sqlalchemy import text

from app.core.access_meta import normalize_path, page_label
from app.core.db import db_healthy, session_scope
from app.services import access_repo, access_stats

pytestmark = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")

_USER = "__access_test__"
_LEADER = {"username": _USER, "role": "leader", "member_units": ["Cao su Test A", "Cao su Test B"]}


@pytest.fixture(autouse=True)
def _cleanup():
    """Dọn đúng dòng của test này (không đụng lịch sử thật)."""
    yield
    with session_scope() as db:
        db.execute(text("DELETE FROM access_log WHERE username = :u"), {"u": _USER})


def _rows(**filters) -> list[dict]:
    return access_repo.search(username=_USER, limit=50, **filters)["items"]


def test_normalize_path_gom_doan_dong():
    assert normalize_path("/ho-tro/12?tab=1") == "/ho-tro/:id"
    assert normalize_path("/ban-tin/xem/ban-tin-2026-09-16.pdf") == "/ban-tin/xem/:id"
    assert normalize_path("/bao-cao-ton-kho/") == "/bao-cao-ton-kho"
    assert normalize_path("/") == "/"
    assert page_label("/bao-cao-ton-kho") == "Tồn kho (theo ngày)"
    # Trang chưa khai vẫn tra cứu được (hiện đường dẫn thô), KHÔNG được ném lỗi.
    assert normalize_path("/mot-trang-la") == "/mot-trang-la"
    assert page_label("/mot-trang-la") == ""


def test_ghi_dang_nhap_thanh_cong_va_that_bai():
    access_repo.log_login(_USER, _LEADER)
    access_repo.log_login(_USER, ok=False)
    events = [r["event"] for r in _rows()]
    assert sorted(events) == ["login", "login_failed"]
    ok_row = next(r for r in _rows() if r["event"] == "login")
    assert ok_row["role"] == "leader"
    assert ok_row["company"] == "Cao su Test A, Cao su Test B"
    assert ok_row["event_label"] == "Đăng nhập"


def test_ghi_luot_vao_trang_kem_ten_tieng_viet():
    access_repo.log_page(_LEADER, "/bao-cao-ton-kho?ngay=2026-09-16")
    row = _rows()[0]
    assert row["path"] == "/bao-cao-ton-kho"
    assert row["label"] == "Tồn kho (theo ngày)"
    assert row["event"] == "page"


def test_chong_trung_lan_vao_lien_tiep_cung_trang():
    """React render lại / bấm F5 liên tục KHÔNG được sinh nhiều dòng."""
    for _ in range(3):
        access_repo.log_page(_LEADER, "/nhu-cau-thi-truong")
    access_repo.log_page(_LEADER, "/ke-hoach-nam")
    paths = [r["path"] for r in _rows()]
    assert paths.count("/nhu-cau-thi-truong") == 1
    assert paths.count("/ke-hoach-nam") == 1


def test_loc_theo_su_kien_va_vai_tro():
    access_repo.log_login(_USER, _LEADER)
    access_repo.log_page(_LEADER, "/ho-tro")
    assert all(r["event"] == "page" for r in _rows(event="page"))
    assert len(_rows(role="leader")) == 2
    assert _rows(role="member") == []


def test_tong_hop_theo_tai_khoan():
    access_repo.log_login(_USER, _LEADER)
    access_repo.log_page(_LEADER, "/ho-tro")
    access_repo.log_page(_LEADER, "/bao-cao-ton-kho")
    row = next(r for r in access_stats.summary(username=_USER) if r["username"] == _USER)
    assert row["logins"] == 1
    assert row["page_views"] == 2
    assert row["distinct_pages"] == 2
    assert row["last_login"] and row["last_seen"]
    assert row["top_page"] in {"Hỗ trợ & Thông báo", "Tồn kho (theo ngày)"}
    top = access_stats.top_pages(username=_USER)
    assert {t["path"] for t in top} == {"/ho-tro", "/bao-cao-ton-kho"}


def test_beacon_chi_nhan_duong_dan_that():
    """Chuỗi lạ (công thức Excel, chữ ký lạ) bị chặn ngay ở cửa, không vào được nhật ký."""
    import pydantic

    from app.routers.access_log import PageView

    assert PageView(path="/bao-cao-ton-kho").path == "/bao-cao-ton-kho"
    for junk in ['=HYPERLINK("http://xau.vn")', "khong-co-gach-cheo", "/<script>alert(1)</script>"]:
        with pytest.raises(pydantic.ValidationError):
            PageView(path=junk)


def test_xuat_excel_khong_de_lot_cong_thuc():
    """Tên tài khoản gõ ở màn đăng nhập là chuỗi NGƯỜI NGOÀI nhập được → phải vô hại trong Excel."""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.services import access_export

    attack = '=HYPERLINK("http://xau.vn","bấm đi")'
    summary = [{"username": attack, "role": "leader", "company": "", "last_seen": None,
                "last_login": None, "logins": 0, "failed_logins": 1, "page_views": 0,
                "active_days": 0, "distinct_pages": 0, "top_page": "", "top_page_views": 0}]
    detail = [{"at": "2026-09-16T08:00:00", "username": attack, "role": "", "company": "",
               "event_label": "Đăng nhập thất bại", "label": "Đăng nhập", "path": "/login",
               "on_behalf": "", "ip": "1.2.3.4"}]
    wb = load_workbook(BytesIO(access_export.build_xlsx(summary, detail)))
    for sheet in wb.worksheets:
        cell = sheet.cell(row=2, column=1 if sheet.title == "Theo tài khoản" else 2).value
        assert cell.startswith("'"), f"{sheet.title}: Excel vẫn hiểu ô là công thức"
        assert attack in cell  # vẫn đọc được nguyên văn, chỉ thôi bị chạy như công thức


def test_tong_hop_bo_qua_luot_dang_nhap_ho():
    """Quản trị 'đăng nhập hộ' thì lượt đó KHÔNG được tính là hoạt động của chủ tài khoản."""
    from app.core import request_ctx

    access_repo.log_page(_LEADER, "/ke-hoach-nam")
    request_ctx.set_request(_USER, "127.0.0.1", on_behalf="admin")
    try:
        access_repo.log_page(_LEADER, "/bao-cao-tieu-thu")
    finally:
        request_ctx.set_request("", "")

    assert len(_rows()) == 2                      # tab Chi tiết vẫn thấy đủ cả hai
    row = next(r for r in access_stats.summary(username=_USER) if r["username"] == _USER)
    assert row["page_views"] == 1                 # bảng tổng hợp chỉ tính lượt của chính họ
    assert [t["path"] for t in access_stats.top_pages(username=_USER)] == ["/ke-hoach-nam"]


def test_tong_hop_hien_ca_tai_khoan_chua_truy_cap():
    """Cấp tài khoản mà chưa dùng bao giờ vẫn phải có một dòng — đó là câu hỏi chính của quản trị."""
    from sqlalchemy import text as sql

    from app.core.security import hash_password
    from app.services import user_repo

    idle = "__idle_probe__@vrg.vn"
    with session_scope() as db:
        db.execute(sql("INSERT INTO app_user (username, password_hash, full_name, role, "
                       "permissions, member_units) VALUES (:u, :p, 'probe', 'leader', '[]'::jsonb, "
                       "'[\"Đơn vị Test\"]'::jsonb) ON CONFLICT (username) DO NOTHING"),
                   {"u": idle, "p": hash_password("x" * 12)})
    try:
        row = next(r for r in access_stats.summary(role="leader") if r["username"] == idle)
        assert row["last_seen"] is None and row["last_login"] is None
        assert row["page_views"] == 0 and row["logins"] == 0
        assert row["company"] == "Đơn vị Test"
        assert idle in access_repo.known_users()          # chọn được ở ô lọc "Tài khoản"
        # Đã khoá thì thôi liệt kê (chưa truy cập là đương nhiên).
        user_repo.update_user(idle, {"is_active": False})
        assert all(r["username"] != idle for r in access_stats.summary(role="leader"))
    finally:
        with session_scope() as db:
            db.execute(sql("DELETE FROM app_user WHERE username = :u"), {"u": idle})
            db.execute(sql("DELETE FROM audit_log WHERE entity = 'user' AND entity_key = :u"),
                       {"u": idle})
