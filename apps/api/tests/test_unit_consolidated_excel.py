"""Biểu TỔNG HỢP gửi Tập đoàn — khoá lại KHUÔN của mẫu Ban TTKD và cách gom chủng loại.

Test dựng file từ dữ liệu giả (không đụng DB) để khuôn biểu không lặng lẽ trôi: đổi cột, đổi thứ tự
khu vực hay đổi nhóm chủng loại là đỏ ngay.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from app.core.market_meta import UNIT_GRADES
from app.services import unit_consolidated_excel as mod
from app.services.unit_consolidated_layout import GRADE_BUCKETS, grade_cells

UNITS = [
    {"name": "Đơn vị A", "region": "Đông Nam Bộ", "merged_into": None, "merged_at": None},
    {"name": "Đơn vị B", "region": "Đông Nam Bộ", "merged_into": "Đơn vị A", "merged_at": "2026-08-21"},
    {"name": "Đơn vị C", "region": "Lào", "merged_into": None, "merged_at": None},
]
ROW_A = {
    "company": "Đơn vị A", "signed_lt_tonnes": 100.0,
    "lt_export": 10.0, "lt_domestic": 20.0, "spot_export": 30.0, "spot_domestic": 40.0,
    "stock_finished": 50.0, "stock_finished_hd": 8.0, "stock_material": 5.0,
    "carry_lt_tonnes": 1.0, "carry_spot_tonnes": 2.0,
    "stock_by_grade": {"SVR CV 50": 8.0, "SVR CV60": 12.0, "SVR 3L": 20.0, "Chủng loại khác": 10.0},
    "stock_hd_by_grade": {"SVR CV 50": 3.0, "SVR 3L": 5.0},
}


@pytest.fixture()
def book(monkeypatch):
    monkeypatch.setattr(mod.member_unit_repo, "list_units", lambda include_inactive=False: UNITS)
    data = mod.build("2026-09-17", {"rows": [ROW_A]})
    return openpyxl.load_workbook(io.BytesIO(data))


def test_khuon_bieu_dung_mau(book) -> None:
    """Tiêu đề · dòng TẬP ĐOÀN · dòng khu vực đúng chỗ, và tổng là CÔNG THỨC chứ không phải số chết."""
    ws = book.worksheets[0]
    assert "TỔNG HỢP BÁO CÁO" in ws["A1"].value and "17/09/2026" in ws["A1"].value
    assert ws["A6"].value == "TẬP ĐOÀN"
    assert ws["B7"].value == "I" and ws["C7"].value == "Đông Nam Bộ"
    assert ws["C10"].value == "Lào" and ws["B10"].value == "II"     # khu vực rỗng bị bỏ qua
    assert ws["P8"].value == "=SUM(H8:K8)" and ws["AC8"].value == "=SUM(S8:AB8)"
    assert ws["G7"].value == "=G8+G9"                                # khu vực cộng đúng dòng con
    assert ws["G6"].value == "=G7+G10"                               # Tập đoàn cộng các khu vực


def test_so_lieu_don_vi_va_don_vi_da_sap_nhap(book) -> None:
    """Đơn vị có số liệu ghi đúng cột; đơn vị đã sáp nhập để trống kèm ghi chú (như mẫu)."""
    ws = book.worksheets[0]
    assert (ws["C8"].value, ws["G8"].value, ws["H8"].value, ws["K8"].value) == ("Đơn vị A", 100.0, 10.0, 40.0)
    assert (ws["Q8"].value, ws["R8"].value, ws["AD8"].value) == (50.0, 8.0, 5.0)
    assert (ws["AE8"].value, ws["AF8"].value) == (1.0, 2.0)
    assert ws["C9"].value == "Đơn vị B" and ws["G9"].value is None
    assert ws["AG9"].value == "Đã hợp nhất số liệu vào Đơn vị A từ 21/08/2026"


def test_chung_loai_gom_dung_nhom_va_tru_phan_da_ky() -> None:
    """Cột chủng loại = tồn kho của chủng loại đó TRỪ phần đã ký; chủng loại lạ dồn vào "Ngoại lệ"."""
    cells = grade_cells(ROW_A)
    labels = [b[0] for b in GRADE_BUCKETS]
    got = dict(zip(labels, cells, strict=True))
    assert got[labels[0]] == 17.0        # CV50/60: (8-3) + 12
    assert got[labels[2]] == 15.0        # SVR L/3L: 20 - 5
    assert got[labels[7]] == 10.0        # "Chủng loại khác" → Ngoại lệ
    assert sum(c or 0 for c in cells) == 42.0 == ROW_A["stock_finished"] - ROW_A["stock_finished_hd"]


def test_khong_in_so_am_khi_da_ky_nhieu_hon_ton_kho() -> None:
    """Tồn kho chốt trễ hơn ngày ký hợp đồng ⇒ hiệu âm, phải kẹp về 0 chứ không in số âm."""
    row = {"stock_by_grade": {"SVR 3L": 5.0}, "stock_hd_by_grade": {"SVR 3L": 9.0}}
    assert all(c is None or c >= 0 for c in grade_cells(row))


def test_moi_chung_loai_deu_co_cho_trong_bieu() -> None:
    """Chủng loại mới thêm vào hệ thống mà quên xếp nhóm thì vẫn hiện ở "Ngoại lệ", không rơi mất."""
    row = {"stock_by_grade": {g: 1.0 for g in UNIT_GRADES}, "stock_hd_by_grade": {}}
    assert sum(c or 0 for c in grade_cells(row)) == len(UNIT_GRADES)
