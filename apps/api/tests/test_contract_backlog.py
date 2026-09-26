"""Test SẢN LƯỢNG CÒN PHẢI GIAO (`contract_backlog`) — HĐ chuyến · HĐ dài hạn theo cam kết HĐ mẹ.

Khoá những chỗ dễ sai nhất (plan 260926 Q1–Q4):
  1. KHÔNG đếm trùng: phụ lục của HĐ mẹ có cam kết chỉ góp ở cấp HĐ mẹ, không góp thêm vào ô chuyến
     / dài hạn ngoài HĐ mẹ.
  2. Còn phải giao = max(cam kết − đã giao, phụ lục đã ký chưa giao); HĐ mẹ hết hạn → 0 + thiếu hụt.
  3. HĐ mẹ không có cam kết → phụ lục tính như hợp đồng thường.
  4. Cùng gốc QUY KHÔ với cột tiêu thụ, lọc chủng loại ở mức dòng, phạm vi đơn vị.
Dữ liệu dựng thẳng bằng SQL: màn nhập hợp đồng có luật riêng (khách hàng, ngày không ở tương lai…)
không liên quan tới phép tính ở đây, và ngày cố định giữ test không đổi màu theo lịch.
"""

from __future__ import annotations

import io
import json

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.core.db import db_healthy, session_scope
from app.main import app
from app.services import contract_backlog, member_unit_merge, user_repo

pytestmark = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")

client = TestClient(app)
UNIT_A, UNIT_B = "_zz_bl_a", "_zz_bl_b"
UNITS = [UNIT_A, UNIT_B]
AS_OF = "2026-06-30"
SVR, LATEX = "SVR 10 / CSR 10", "LATEX"


def _ln(qty: float | None, grade: str = SVR, dry: float | None = None) -> dict:
    return {"grade": grade, "qty": qty, "qty_dry": dry, "price": 40.0, "ccy": "VND"}


def _master(company: str, code: str, lines: list[dict], sign: str | None = "2026-01-05",
            expiry: str | None = "2026-12-31", mtype: str = "long_term") -> int:
    with session_scope() as db:
        return db.execute(text(
            "INSERT INTO master_contract (company, code, master_type, sign_date, expiry_date, lines) "
            "VALUES (:c, :code, :t, :s, :e, CAST(:l AS jsonb)) RETURNING id"),
            {"c": company, "code": code, "t": mtype, "s": sign, "e": expiry,
             "l": json.dumps(lines)}).scalar()


def _contract(company: str, code: str, lines: list[dict], *, ctype: str | None = "spot",
              master_id: int | None = None, parent_id: int | None = None,
              sign: str = "2026-02-01", delivered_at: str | None = None,
              multi: bool = False) -> int:
    with session_scope() as db:
        return db.execute(text(
            "INSERT INTO sales_contract (company, parent_id, code, delivery_type, contract_type, "
            "  sign_date, lines, delivered, delivered_at, master_id) "
            "VALUES (:c, :p, :code, :dt, :t, :s, CAST(:l AS jsonb), :dd, :d, :m) RETURNING id"),
            {"c": company, "p": parent_id, "code": code, "dt": "multi" if multi else "single",
             "t": ctype, "s": sign, "l": json.dumps(lines), "dd": delivered_at is not None,
             "d": delivered_at, "m": master_id}).scalar()


def _wipe() -> None:
    with session_scope() as db:
        for tbl in ("sales_contract", "master_contract"):
            db.execute(text(f"DELETE FROM {tbl} WHERE company = ANY(:u)"), {"u": UNITS})


@pytest.fixture(autouse=True)
def _clean():
    _wipe()
    yield
    _wipe()


def _bl(company: str = UNIT_A, **kw) -> dict:
    return contract_backlog.backlog_on(AS_OF, UNITS, **kw).get(company) or {}


def test_chia_chuyen_dai_han_khong_dem_trung() -> None:
    m = _master(UNIT_A, "HDDH-1", [_ln(1000)])
    _contract(UNIT_A, "PL-1", [_ln(300)], ctype="long_term", master_id=m, delivered_at="2026-03-01")
    _contract(UNIT_A, "PL-2", [_ln(200)], ctype="long_term", master_id=m)    # ký, chưa giao
    _contract(UNIT_A, "CH-1", [_ln(100)], ctype="spot")
    _contract(UNIT_A, "DH-LE", [_ln(50)], ctype="long_term")                  # dài hạn ngoài HĐ mẹ
    _contract(UNIT_A, "KO-LOAI", [_ln(30)], ctype=None)

    b = _bl()
    assert (b["spot_undelivered"], b["lt_unlinked_undelivered"], b["unknown_undelivered"]) == (
        pytest.approx(100), pytest.approx(50), pytest.approx(30))
    assert b["master_committed"] == pytest.approx(1000)
    assert b["master_delivered"] == pytest.approx(300)
    # 700 = cam kết − đã giao; PL-2 (200) nằm TRONG 700 chứ không cộng thêm → không đếm trùng.
    assert b["master_remaining"] == pytest.approx(700)
    assert b["lt_remaining"] == pytest.approx(750)
    assert b["to_deliver"] == pytest.approx(880)
    assert b["master_pct"] == pytest.approx(30) and b["masters"] == 1
    item = b["items"][0]
    assert (item["id"], item["code"], item["expired"]) == (m, "HDDH-1", False)
    assert item["pct"] == pytest.approx(30) and item["remaining"] == pytest.approx(700)


def test_phu_luc_ky_vuot_cam_ket_thi_lay_phan_da_ky() -> None:
    """Cam kết 500, đã giao 300 (qua đợt giao), phụ lục còn 400 chưa giao → phải giao 400, không 200."""
    m = _master(UNIT_A, "HDNT-1", [_ln(500)], mtype="principle")
    pl = _contract(UNIT_A, "PL-X", [_ln(700)], ctype="long_term", master_id=m, multi=True)
    _contract(UNIT_A, "PL-X/1", [_ln(300)], ctype=None, parent_id=pl, delivered_at="2026-04-10")

    b = _bl()
    assert b["master_delivered"] == pytest.approx(300)       # đợt giao của phụ lục được đếm
    assert b["master_remaining"] == pytest.approx(400)
    assert b["lt_unlinked_undelivered"] == 0 and b["to_deliver"] == pytest.approx(400)
    assert b["items"][0]["master_type"] == "principle"


def test_hd_me_het_han_vao_thieu_hut_rieng() -> None:
    m = _master(UNIT_A, "HDDH-HH", [_ln(800)], expiry="2026-05-31")          # hết hạn trong năm
    _contract(UNIT_A, "PL-HH1", [_ln(500)], ctype="long_term", master_id=m, delivered_at="2026-02-01")
    _contract(UNIT_A, "PL-HH2", [_ln(100)], ctype="long_term", master_id=m)  # nuốt vào HĐ mẹ
    # Hết hạn từ năm trước → KHÔNG được tính; phụ lục của nó quay về khối 3 như hợp đồng thường.
    old = _master(UNIT_A, "HDDH-CU", [_ln(900)], sign="2025-01-01", expiry="2025-12-31")
    _contract(UNIT_A, "PL-CU", [_ln(70)], ctype="long_term", master_id=old, sign="2025-06-01")

    b = _bl()
    assert b["master_committed"] == pytest.approx(800) and b["master_delivered"] == pytest.approx(500)
    # Hết hạn: phần chưa ký phụ lục (800 − 500 − 100 = 200) thôi phải giao, nhưng PL-HH2 đã ký 100 t
    # vẫn là nợ giao → tổng phải giao KHÔNG được nhỏ hơn "đã ký HĐ chưa giao" (100 + 70).
    assert b["master_remaining"] == pytest.approx(100)
    assert b["master_expired_short"] == pytest.approx(200)
    assert b["lt_unlinked_undelivered"] == pytest.approx(70)
    assert b["to_deliver"] == pytest.approx(170)
    assert [(i["code"], i["expired"], i["remaining"]) for i in b["items"]] == [("HDDH-HH", True, 100)]


def test_cam_ket_latex_khong_quy_kho_thi_tru_tren_mu_nuoc() -> None:
    """Hồ sơ cũ (trước luật bắt buộc quy khô) cam kết 1.000 t nước, không khô; giao 300 nước / 180
    khô → trừ trên NƯỚC: còn 700, không phải 820 (1.000 nước − 180 khô)."""
    m = _master(UNIT_A, "HDDH-LX-CU", [_ln(1000, LATEX)])
    _contract(UNIT_A, "PL-LX-CU", [_ln(300, LATEX, 180)], ctype="long_term", master_id=m,
              delivered_at="2026-05-05")

    b = _bl()
    assert (b["master_committed"], b["master_delivered"]) == (pytest.approx(1000), pytest.approx(300))
    assert b["master_remaining"] == pytest.approx(700)


def test_hd_me_vat_sang_nam_sau_bao_rieng() -> None:
    _master(UNIT_A, "HDDH-2027", [_ln(400)], expiry="2027-06-30")
    _master(UNIT_A, "HDNT-VOHAN", [_ln(100)], expiry=None, mtype="principle")
    _master(UNIT_A, "HDDH-2026", [_ln(50)])

    b = _bl()
    assert b["master_remaining"] == pytest.approx(550)
    assert b["master_remaining_after_year"] == pytest.approx(500)


def test_hd_me_khong_cam_ket_thi_phu_luc_tinh_nhu_hop_dong_thuong() -> None:
    m = _master(UNIT_A, "HDNT-TRONG", [_ln(None)], mtype="principle")        # chỉ chốt chủng loại
    _contract(UNIT_A, "PL-T", [_ln(120)], ctype="long_term", master_id=m)

    b = _bl()
    assert b["lt_unlinked_undelivered"] == pytest.approx(120)
    assert (b["masters"], b["items"], b["master_pct"]) == (0, [], None)
    assert b["lt_remaining"] == pytest.approx(120)


def test_quy_kho_latex_cung_goc_so_voi_tieu_thu() -> None:
    """Latex: cam kết 1.000 t nước / 600 t khô, giao 300 nước / 180 khô → còn 420 KHÔ, không 700."""
    m = _master(UNIT_A, "HDDH-LX", [_ln(1000, LATEX, 600)])
    _contract(UNIT_A, "PL-LX", [_ln(300, LATEX, 180)], ctype="long_term", master_id=m,
              delivered_at="2026-05-05")

    b = _bl()
    assert b["master_committed"] == pytest.approx(600)
    assert b["master_delivered"] == pytest.approx(180)
    assert b["master_remaining"] == pytest.approx(420)


def test_loc_chung_loai_ap_ca_cam_ket_lan_da_giao() -> None:
    m = _master(UNIT_A, "HDDH-2L", [_ln(600), _ln(1000, LATEX, 600)])
    _contract(UNIT_A, "PL-2L", [_ln(100), _ln(200, LATEX, 120)], ctype="long_term", master_id=m,
              delivered_at="2026-03-03")
    _contract(UNIT_A, "CH-SVR", [_ln(40)], ctype="spot")

    full = _bl()
    assert (full["master_committed"], full["master_delivered"]) == (
        pytest.approx(1200), pytest.approx(220))
    only = _bl(grades=[LATEX])
    assert (only["master_committed"], only["master_delivered"]) == (
        pytest.approx(600), pytest.approx(120))
    assert only["spot_undelivered"] == 0                     # HĐ chuyến chỉ bán SVR → bị lọc
    assert only["to_deliver"] == pytest.approx(480)


def test_ngay_tinh_moc_ky_giao_va_luy_ke_tu_ngay_ky() -> None:
    """Giao SAU ngày tính chưa phải đã giao; giao năm trước vẫn tính (lũy kế từ ngày ký);
    HĐ mẹ ký sau ngày tính chưa có hiệu lực."""
    m = _master(UNIT_A, "HDDH-LK", [_ln(1000)], sign="2025-10-01")
    _contract(UNIT_A, "PL-NAM-TRUOC", [_ln(200)], ctype="long_term", master_id=m, sign="2025-10-05",
              delivered_at="2025-11-15")
    _contract(UNIT_A, "PL-SAU", [_ln(300)], ctype="long_term", master_id=m, delivered_at="2026-07-15")
    _master(UNIT_A, "HDDH-TUONG-LAI", [_ln(500)], sign="2026-08-01")

    b = _bl()
    assert b["master_committed"] == pytest.approx(1000) and b["masters"] == 1
    assert b["master_delivered"] == pytest.approx(200)
    assert b["master_remaining"] == pytest.approx(800)       # max(800, PL-SAU còn 300 tại ngày tính)


def test_pham_vi_don_vi() -> None:
    _master(UNIT_B, "HDDH-B", [_ln(400)])
    _contract(UNIT_B, "CH-B", [_ln(60)], ctype="spot")
    _contract(UNIT_A, "CH-A", [_ln(10)], ctype="spot")

    only_a = contract_backlog.backlog_on(AS_OF, [UNIT_A])
    assert set(only_a) == {UNIT_A} and only_a[UNIT_A]["to_deliver"] == pytest.approx(10)
    assert contract_backlog.backlog_on(AS_OF, []) == {}
    assert _bl(UNIT_B)["to_deliver"] == pytest.approx(460)


def test_gop_don_vi_sap_nhap_tinh_lai_ty_le(monkeypatch) -> None:
    """Gộp B (đã sáp nhập) vào A: số cộng lại, % chia lại trên tổng — cộng hai tỷ lệ là sai."""
    _contract(UNIT_A, "PL-A", [_ln(50)], ctype="long_term", delivered_at="2026-03-01",
              master_id=_master(UNIT_A, "HDDH-A", [_ln(100)]))
    _master(UNIT_B, "HDDH-B", [_ln(300)])
    monkeypatch.setattr(member_unit_merge, "rollup_map", lambda: {UNIT_B: UNIT_A})

    rolled = contract_backlog.roll(contract_backlog.backlog_on(AS_OF, UNITS))
    assert set(rolled) == {UNIT_A}
    b = rolled[UNIT_A]
    assert (b["master_committed"], b["master_delivered"]) == (pytest.approx(400), pytest.approx(50))
    assert b["master_pct"] == pytest.approx(12.5) and b["masters"] == 2
    assert [i["code"] for i in b["items"]] == ["HDDH-B", "HDDH-A"]           # còn phải giao giảm dần
    assert b["to_deliver"] == pytest.approx(350)


def _admin() -> dict[str, str]:
    user_repo.seed_admin()
    tok = client.post("/api/auth/login", json={"username": "admin", "password": "admin"}).json()
    return {"Authorization": f"Bearer {tok['access_token']}"}


def test_api_tieu_thu_tra_backlog_va_file_excel_co_cot_phai_giao() -> None:
    from openpyxl import load_workbook

    h = _admin()
    client.post("/api/member-units", json={"name": UNIT_A}, headers=h)
    try:
        m = _master(UNIT_A, "HDDH-API", [_ln(1000)])
        _contract(UNIT_A, "PL-API", [_ln(400)], ctype="long_term", master_id=m,
                  delivered_at="2026-06-01")
        _contract(UNIT_A, "CH-API", [_ln(25)], ctype="spot")
        q = f"date_from=2026-06-01&date_to={AS_OF}&company={UNIT_A}"
        rep = client.get(f"/api/sales-contracts/consumption?{q}", headers=h).json()
        assert rep["backlog_as_of"] == AS_OF
        b = rep["backlog"][UNIT_A]
        assert b["to_deliver"] == pytest.approx(625) and b["master_pct"] == pytest.approx(40)
        # Khối 3 GIỮ NGUYÊN: chỉ hợp đồng đã ký (25 t chuyến), không gồm cam kết chưa ký phụ lục.
        assert rep["undelivered"][UNIT_A]["qty"] == pytest.approx(25)

        xls = client.get(f"/api/sales-contracts/consumption.xlsx?{q}", headers=h)
        ws = load_workbook(io.BytesIO(xls.content))["Đơn vị"]
        head = [c.value for c in ws[5]]
        row = dict(zip(head, [c.value for c in ws[7]], strict=True))
        assert row["Tổng phải giao"] == pytest.approx(625)
        assert row["HĐ dài hạn còn phải giao"] == pytest.approx(600)
        assert row["% thực hiện HĐ mẹ"] == pytest.approx(40)
        assert row["Đã ký HĐ chưa giao (khối 3)"] == pytest.approx(25)
    finally:
        _wipe()
        client.delete(f"/api/member-units/{UNIT_A}", headers=h)
