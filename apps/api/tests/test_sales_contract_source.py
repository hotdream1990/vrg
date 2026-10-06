"""NGUỒN TIÊU THỤ của lần giao — khai thác / thu mua / hàng hóa (chốt 03/10/2026; theo từng dòng
chủng loại từ 05/10/2026).

Luật: hợp đồng chưa giao chưa bắt buộc khai; có NGÀY GIAO (chốt thành tiêu thụ) là phải chọn
nguồn — ở chính hợp đồng giao 1 lần, hoặc ở từng đợt giao của hợp đồng giao nhiều lần. Báo cáo
tiêu thụ tách được sản lượng theo nguồn.
"""

from __future__ import annotations

import json
from datetime import date

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.core.db import db_healthy, session_scope
from app.main import app
from app.services import user_repo

pytestmark = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")

client = TestClient(app)

UNIT = "_zz_src_unit"
TODAY = date.today().isoformat()


def _admin() -> dict[str, str]:
    user_repo.seed_admin()
    token = client.post("/api/auth/login",
                        json={"username": "admin", "password": "admin"}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _cleanup(h: dict[str, str]) -> None:
    with session_scope() as db:
        db.execute(text("DELETE FROM sales_contract WHERE company = :u"), {"u": UNIT})
        db.execute(text("DELETE FROM unit_customer WHERE company = :u"), {"u": UNIT})
    client.delete(f"/api/member-units/{UNIT}", headers=h)


@pytest.fixture()
def env():
    h = _admin()
    _cleanup(h)
    client.post("/api/member-units", json={"name": UNIT}, headers=h)
    cus = client.put("/api/customers", json={"company": UNIT, "name": "KH nguồn"},
                     headers=h).json()["id"]
    yield h, cus
    _cleanup(h)


def _line(qty: float) -> dict:
    return {"grade": "SVR 10 / CSR 10", "qty": qty, "price": 40.0, "ccy": "VND"}


def _put(h: dict[str, str], body: dict):
    return client.put("/api/sales-contracts", json={"company": UNIT, **body}, headers=h)


def _single(cus: int, code: str, qty: float, **kw) -> dict:
    return {"code": code, "customer_id": cus, "delivery_type": "single", "contract_type": "spot",
            "sign_date": TODAY, "lines": [_line(qty)], **kw}


def test_delivery_requires_a_valid_source(env) -> None:
    h, cus = env
    # Chưa giao: chưa bắt buộc — hợp đồng khai lúc ký không cần biết hàng lấy từ đâu.
    pending = _put(h, _single(cus, "HD-SRC-0", 10.0))
    assert pending.status_code == 200, pending.text
    assert pending.json()["contract"]["source"] is None

    # Có ngày giao mà thiếu nguồn → chặn, KHÔNG tự gán mặc định.
    miss = _put(h, _single(cus, "HD-SRC-1", 10.0, delivered_at=TODAY, channel="domestic"))
    assert miss.status_code == 400 and "nguồn tiêu thụ" in miss.json()["detail"]

    bad = _put(h, _single(cus, "HD-SRC-2", 10.0, delivered_at=TODAY, channel="domestic",
                          source="khai_thac"))
    assert bad.status_code == 400 and "không hợp lệ" in bad.json()["detail"]

    ok = _put(h, _single(cus, "HD-SRC-3", 10.0, delivered_at=TODAY, channel="domestic",
                         source="purchase"))
    assert ok.status_code == 200, ok.text
    assert ok.json()["contract"]["source"] == "purchase"


def test_batches_carry_the_source_and_consumption_splits_by_it(env) -> None:
    h, cus = env
    parent = _put(h, {"code": "HD-SRC-M", "customer_id": cus, "delivery_type": "multi",
                      "contract_type": "spot", "sign_date": TODAY, "lines": [_line(100.0)],
                      # Hợp đồng giao nhiều lần không mang nguồn — khai ở từng đợt giao.
                      "source": "exploit"})
    assert parent.status_code == 200, parent.text
    pid = parent.json()["contract"]["id"]
    assert parent.json()["contract"]["source"] is None

    no_src = _put(h, {"parent_id": pid, "code": "1", "delivered_at": TODAY, "channel": "export",
                      "lines": [_line(30.0)]})
    assert no_src.status_code == 400 and "nguồn tiêu thụ" in no_src.json()["detail"]
    for code, qty, src in (("1", 30.0, "exploit"), ("2", 20.0, "purchase"), ("3", 10.0, "goods")):
        r = _put(h, {"parent_id": pid, "code": code, "delivered_at": TODAY, "channel": "export",
                     "source": src, "lines": [_line(qty)]})
        assert r.status_code == 200, r.text

    params = {"date_from": TODAY, "date_to": TODAY, "company": UNIT}
    rep = client.get("/api/sales-contracts/consumption", headers=h, params=params).json()
    by_source = rep["by_company"][UNIT]["by_source"]
    assert by_source == {"exploit": pytest.approx(30.0), "purchase": pytest.approx(20.0),
                         "goods": pytest.approx(10.0)}

    hist = client.get("/api/sales-contracts/consumption/deliveries", headers=h,
                      params=params).json()
    assert sorted(r["source"] for r in hist["rows"]) == ["Hàng hóa cao su", "Khai thác", "Thu mua"]

    # Màn Thống kê tiêu thụ đọc cùng nguồn: lọc theo nguồn chỉ còn phần thu mua.
    stats = client.get("/api/unit-daily/analytics/consumption", headers=h, params={
        "date_from": TODAY, "date_to": TODAY, "companies": UNIT, "source": "purchase",
        "group_by": "company", "split_merged": "true"}).json()
    assert stats["totals"]["qty"] == pytest.approx(20.0)
    assert stats["totals"]["qty_purchase"] == pytest.approx(20.0)
    assert not stats["totals"].get("qty_goods")

    goods = client.get("/api/unit-daily/analytics/consumption", headers=h, params={
        "date_from": TODAY, "date_to": TODAY, "companies": UNIT, "source": "goods",
        "group_by": "company", "split_merged": "true"}).json()
    assert goods["totals"]["qty"] == pytest.approx(10.0)
    assert goods["totals"]["qty_goods"] == pytest.approx(10.0)


def test_completing_a_single_contract_asks_for_the_source(env) -> None:
    h, cus = env
    cid = _put(h, _single(cus, "HD-SRC-C", 15.0)).json()["contract"]["id"]
    url = f"/api/sales-contracts/{cid}/completion"
    miss = client.put(url, headers=h, json={"completed_at": TODAY, "channel": "domestic"})
    assert miss.status_code == 400 and "Nguồn tiêu thụ" in miss.json()["detail"]

    ok = client.put(url, headers=h, json={"completed_at": TODAY, "channel": "domestic",
                                          "source": "purchase"})
    assert ok.status_code == 200, ok.text
    c = ok.json()["contract"]
    assert (c["delivered_at"], c["source"]) == (TODAY, "purchase")


def test_switching_to_multi_moves_the_source_to_the_first_batch(env) -> None:
    h, cus = env
    cid = _put(h, _single(cus, "HD-SRC-S", 12.0, delivered_at=TODAY, channel="domestic",
                          source="purchase")).json()["contract"]["id"]
    r = client.put(f"/api/sales-contracts/{cid}/delivery-type", headers=h,
                   json={"delivery_type": "multi"})
    assert r.status_code == 200, r.text
    assert r.json()["contract"]["source"] is None
    kids = client.get(f"/api/sales-contracts/{cid}", headers=h).json()["children"]
    assert [k["source"] for k in kids] == ["purchase"]


def test_legacy_edit_request_payload_keeps_the_stored_source(env) -> None:
    """Đề nghị sửa gửi TRƯỚC khi có ô nguồn không mang khoá `source` — Ban duyệt vẫn qua được."""
    from app.services.edit_request_ops_contract import _legacy_source

    h, cus = env
    c = _put(h, _single(cus, "HD-SRC-L", 8.0, delivered_at=TODAY, channel="domestic",
                        source="purchase")).json()["contract"]
    legacy = {k: v for k, v in c.items() if k != "source"}
    assert [ln["source"] for ln in _legacy_source(legacy)["lines"]] == ["purchase"]
    # Bản ghi mới (chưa có id) có ngày giao → tính khai thác như mọi lần giao cũ.
    fresh = _legacy_source({"delivered_at": TODAY, "lines": [_line(1.0)]})
    assert [ln["source"] for ln in fresh["lines"]] == ["exploit"]
    # Lần giao đang LẪN nguồn: điền theo dòng cùng vị trí, không gom hết về khai thác.
    mixed = _put(h, _single(cus, "HD-SRC-LM", 1.0, delivered_at=TODAY, channel="domestic",
                            lines=[{**_line(2.0), "source": "purchase"},
                                   {**_line(3.0), "grade": "SVR 3L", "source": "goods"}])).json()
    old = {k: v for k, v in mixed["contract"].items() if k != "source"}
    old["lines"] = [{k: v for k, v in ln.items() if k != "source"} for ln in old["lines"]]
    assert [ln["source"] for ln in _legacy_source(old)["lines"]] == ["purchase", "goods"]
    # Payload mới luôn mang khoá — giữ nguyên, kể cả khi để trống (luật bắt buộc vẫn áp).
    assert _legacy_source({**c, "source": None})["source"] is None
    # Chưa giao thì không đoán gì.
    assert "source" not in _legacy_source({"code": "x"})


# ── Nguồn theo TỪNG DÒNG chủng loại (chốt 05/10/2026) ────────────────────────────────────────
def _svr3l(qty: float, **kw) -> dict:
    return {**_line(qty), "grade": "SVR 3L", **kw}


def test_each_line_carries_its_own_source_and_reports_split_by_line(env) -> None:
    """Một lần giao gồm SVR 10 khai thác + SVR 3L hàng hóa: mọi báo cáo tách đúng theo dòng."""
    h, cus = env
    lines = [{**_line(30.0), "source": "exploit"}, _svr3l(10.0, source="goods")]
    base = {"delivered_at": TODAY, "channel": "domestic"}

    # Có ngày giao mà một dòng chưa có nguồn (và không có nguồn chung) → chặn, chỉ rõ dòng.
    miss = _put(h, _single(cus, "HD-LN-0", 1.0, **base, lines=[lines[0], _svr3l(10.0)]))
    assert miss.status_code == 400 and "Dòng 2 (SVR 3L)" in miss.json()["detail"]
    bad = _put(h, _single(cus, "HD-LN-X", 1.0, **base, lines=[lines[0], _svr3l(10.0, source="x")]))
    assert bad.status_code == 400 and "không hợp lệ" in bad.json()["detail"]

    ok = _put(h, _single(cus, "HD-LN-1", 1.0, **base, lines=lines))
    assert ok.status_code == 200, ok.text
    c = ok.json()["contract"]
    assert [ln["source"] for ln in c["lines"]] == ["exploit", "goods"]
    assert c["source"] is None                       # lẫn nguồn → không có nguồn chung

    params = {"date_from": TODAY, "date_to": TODAY, "company": UNIT}
    rep = client.get("/api/sales-contracts/consumption", headers=h,
                     params=params).json()["by_company"][UNIT]
    assert rep["by_source"] == {"exploit": pytest.approx(30.0), "goods": pytest.approx(10.0)}
    assert rep["by_source_grade"] == {"exploit": {"SVR 10 / CSR 10": pytest.approx(30.0)},
                                      "goods": {"SVR 3L": pytest.approx(10.0)}}

    hist = client.get("/api/sales-contracts/consumption/deliveries", headers=h,
                      params=params).json()
    assert [r["source"] for r in hist["rows"]] == ["Khai thác · Hàng hóa cao su"]

    # Excel Báo cáo tiêu thụ có sheet "Nguồn × chủng loại" khớp số với API.
    from io import BytesIO

    from openpyxl import load_workbook

    xlsx = client.get("/api/sales-contracts/consumption.xlsx", headers=h, params=params)
    assert xlsx.status_code == 200
    ws = load_workbook(BytesIO(xlsx.content))["Nguồn × chủng loại"]
    cells = {ws.cell(r, 1).value: [ws.cell(r, c).value for c in range(2, 9)]
             for r in range(1, ws.max_row + 1)}
    # Cột: KT (tấn, %) · TM (tấn, %) · HH (tấn, %) · Tổng
    assert cells["SVR 10 / CSR 10"] == [30.0, 100.0, None, None, None, None, 30.0]
    assert cells["SVR 3L"] == [None, None, None, None, 10.0, 100.0, 10.0]
    assert cells["Tổng cộng"] == [30.0, 75.0, None, None, 10.0, 25.0, 40.0]

    # Thống kê tiêu thụ: lọc theo nguồn chỉ còn đúng dòng đó; nhóm theo chủng loại = ma trận.
    stats = {"date_from": TODAY, "date_to": TODAY, "companies": UNIT, "split_merged": "true"}
    goods = client.get("/api/unit-daily/analytics/consumption", headers=h,
                       params={**stats, "source": "goods", "group_by": "company"}).json()
    assert goods["totals"]["qty"] == pytest.approx(10.0)
    by_grade = client.get("/api/unit-daily/analytics/consumption", headers=h,
                          params={**stats, "group_by": "grade"}).json()
    rows = {r["key"]: r for r in by_grade["rows"]}
    assert rows["SVR 3L"]["qty_goods"] == pytest.approx(10.0) and not rows["SVR 3L"]["qty_exploit"]
    assert rows["SVR 10 / CSR 10"]["qty_exploit"] == pytest.approx(30.0)


def test_completion_takes_a_source_per_line(env) -> None:
    h, cus = env
    cid = _put(h, _single(cus, "HD-LN-C", 1.0,
                          lines=[_line(5.0), _svr3l(3.0)])).json()["contract"]["id"]
    url = f"/api/sales-contracts/{cid}/completion"
    part = client.put(url, headers=h, json={"completed_at": TODAY, "channel": "domestic",
                                            "line_sources": ["purchase", None]})
    assert part.status_code == 400 and "Nguồn tiêu thụ" in part.json()["detail"]

    ok = client.put(url, headers=h, json={"completed_at": TODAY, "channel": "domestic",
                                          "line_sources": ["purchase", "goods"]})
    assert ok.status_code == 200, ok.text
    c = ok.json()["contract"]
    assert c["delivered_at"] == TODAY
    assert [ln["source"] for ln in c["lines"]] == ["purchase", "goods"]


def test_migration_writes_the_delivery_source_down_to_each_line_once(env, monkeypatch) -> None:
    """Bước dữ liệu một lần (`core/db.py`, mốc `schema_once`): lần giao cũ chỉ có nguồn ở cấp lần
    giao → ghi xuống từng dòng; đã có mốc thì khởi động lại không chạy nữa."""
    from app.core import db as core_db

    h, cus = env
    cid = _put(h, _single(cus, "HD-MIG", 6.0, delivered_at=TODAY, channel="domestic",
                          source="purchase")).json()["contract"]["id"]

    def strip_line_source() -> None:
        with session_scope() as s:
            s.execute(text("UPDATE sales_contract SET lines = (SELECT jsonb_agg(e - 'source') "
                           "FROM jsonb_array_elements(lines) e) WHERE id = :i"), {"i": cid})

    def line_sources() -> list:
        with session_scope() as s:
            lines = s.execute(text("SELECT lines FROM sales_contract WHERE id = :i"),
                              {"i": cid}).scalar()
        return [ln.get("source") for ln in lines]

    strip_line_source()
    # Đề nghị sửa ĐANG CHỜ gửi trước bản này: ảnh chụp "lúc gửi" chưa có nguồn ở dòng.
    before = {"id": cid, "source": "purchase", "lines": [{"grade": "SVR 10 / CSR 10", "qty": 6.0}]}
    with session_scope() as s:
        s.execute(text("DELETE FROM schema_once WHERE name = 'sales_contract_line_source'"))
        req = s.execute(text(
            "INSERT INTO edit_request (company, op, target_key, payload, before, requested_by) "
            "VALUES (:c, 'contract_save', :k, '{}'::jsonb, CAST(:b AS jsonb), 'zz') RETURNING id"),
            {"c": UNIT, "k": f"zz-mig-{cid}", "b": json.dumps(before)}).scalar()
    monkeypatch.setattr(core_db, "_schema_ready", False)
    try:
        core_db.ensure_schema()
        assert line_sources() == ["purchase"]
        with session_scope() as s:
            snap = s.execute(text("SELECT before FROM edit_request WHERE id = :i"),
                             {"i": req}).scalar()
        assert [ln.get("source") for ln in snap["lines"]] == ["purchase"]
    finally:
        with session_scope() as s:
            s.execute(text("DELETE FROM edit_request WHERE id = :i"), {"i": req})

    strip_line_source()
    monkeypatch.setattr(core_db, "_schema_ready", False)
    core_db.ensure_schema()
    assert line_sources() == [None]                  # mốc đã có → không chạy lại


def test_old_client_single_source_applies_to_every_line(env) -> None:
    """Trình duyệt còn bản web CŨ gửi một nguồn cấp lần giao kèm dòng (đã có nguồn sau migration):
    nguồn đó áp cho MỌI dòng — chỉ điền dòng trống thì thay đổi bị nuốt mà vẫn báo lưu thành công."""
    h, cus = env
    c = _put(h, _single(cus, "HD-LN-OLD", 1.0, delivered_at=TODAY, channel="domestic",
                        lines=[{**_line(2.0), "source": "exploit"},
                               _svr3l(3.0, source="goods")])).json()["contract"]
    r = _put(h, {**{k: v for k, v in c.items() if k not in ("qty", "revenue")},
                 "source": "purchase"})
    assert r.status_code == 200, r.text
    assert [ln["source"] for ln in r.json()["contract"]["lines"]] == ["purchase", "purchase"]


def test_switching_an_undelivered_single_to_multi_makes_no_ghost_batch(env) -> None:
    """Hợp đồng giao 1 lần CHƯA giao, đã chọn nguồn/hình thức sẵn → chuyển sang giao nhiều lần không
    được sinh "đợt" chờ giao mang 100% sản lượng (lỗi F2); hợp đồng cha không còn nguồn ở dòng."""
    h, cus = env
    cid = _put(h, _single(cus, "HD-LN-F2", 1.0, channel="domestic",
                          lines=[{**_line(8.0), "source": "exploit"}])).json()["contract"]["id"]
    r = client.put(f"/api/sales-contracts/{cid}/delivery-type", headers=h,
                   json={"delivery_type": "multi"})
    assert r.status_code == 200, r.text
    detail = client.get(f"/api/sales-contracts/{cid}", headers=h).json()
    assert detail["children"] == []
    assert all("source" not in ln for ln in detail["contract"]["lines"])
    assert detail["contract"]["channel"] is None
