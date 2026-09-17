"""Test Cảnh báo bất thường cho LÃNH ĐẠO ĐƠN VỊ: chỉ thấy đơn vị mình, không thấy số gộp Tập đoàn."""

from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.db import db_healthy
from app.main import app
from app.services import anomaly_scope, user_repo

client = TestClient(app)

UNIT_A, UNIT_B = "_zz_an_don_vi_a", "_zz_an_don_vi_b"
ACCOUNTS = ("an_lead_a", "an_mem_a")


def _group(key: str, severity: str, units: list[str]) -> dict:
    rows = [{"don_vi": u, "ghi_chu": f"dòng của {u}"} for u in units]
    return {"key": key, "label": key, "desc": "", "severity": severity, "columns": [],
            "rows": rows, "count": len(rows), "units": len(set(units))}


def test_for_units_chi_giu_don_vi_duoc_gan_va_tinh_lai_tong_quan():
    """Không cần DB: kiểm thuần hàm thu hẹp."""
    result = {"date_from": "2026-01-01", "date_to": "2026-09-16", "groups": [
        _group("wrong_sale_price", "high", ["A", "B"]),
        _group("revenue_outlier", "high", ["A"]),          # số gộp Tập đoàn → phải bỏ hẳn
        _group("not_submitted", "medium", ["B"]),
        _group("plan_missing", "low", ["A", "A"]),
    ]}
    out = anomaly_scope.for_units(result, ["A"])
    keys = [g["key"] for g in out["groups"]]
    assert "revenue_outlier" not in keys
    assert all(r["don_vi"] == "A" for g in out["groups"] for r in g["rows"])
    # Nhóm có cảnh báo lên trước, nhóm rỗng (not_submitted chỉ có B) xuống cuối.
    assert keys == ["wrong_sale_price", "plan_missing", "not_submitted"]
    assert out["summary"] == {"total": 3, "high": 1, "medium": 0, "low": 2, "units": 1}
    by_key = {g["key"]: g for g in out["groups"]}
    assert (by_key["plan_missing"]["count"], by_key["plan_missing"]["units"]) == (2, 1)


def _bearer(username: str, password: str) -> dict[str, str]:
    token = client.post("/api/auth/login",
                        json={"username": username, "password": password}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def env():
    """2 đơn vị (chưa nộp gì → chắc chắn có cảnh báo) + lãnh đạo và nhập liệu của đơn vị A."""
    if not db_healthy():
        pytest.skip("DB không sẵn sàng")
    user_repo.seed_admin()
    h = _bearer("admin", "admin")
    for u in ACCOUNTS:
        client.delete(f"/api/users/{u}", headers=h)
    for unit in (UNIT_A, UNIT_B):
        client.post("/api/member-units", json={"name": unit}, headers=h)
    client.post("/api/users", json={"username": "an_lead_a", "password": "pass123",
                                    "role": "leader", "member_units": [UNIT_A]}, headers=h)
    client.post("/api/users", json={"username": "an_mem_a", "password": "pass123",
                                    "role": "member", "member_units": [UNIT_A]}, headers=h)
    yield {"admin": h, "lead": _bearer("an_lead_a", "pass123"),
           "mem": _bearer("an_mem_a", "pass123")}
    for u in ACCOUNTS:
        client.delete(f"/api/users/{u}", headers=h)
    for unit in (UNIT_A, UNIT_B):
        client.delete(f"/api/member-units/{unit}", headers=h)


def test_lanh_dao_chi_thay_don_vi_minh(env):
    r = client.get("/api/member/anomalies", headers=env["lead"])
    assert r.status_code == 200
    body = r.json()
    units = {row.get("don_vi") for g in body["groups"] for row in g["rows"]}
    assert units == {UNIT_A}                       # có cảnh báo của A, tuyệt đối không có B
    assert "revenue_outlier" not in {g["key"] for g in body["groups"]}
    assert "thresholds" not in body                # ngưỡng cấu hình là việc của quản trị
    assert body["summary"]["units"] == 1

    # Quản trị vẫn thấy đủ cả hai đơn vị ở màn của mình.
    admin = client.get("/api/anomalies", headers=env["admin"]).json()
    admin_units = {row.get("don_vi") for g in admin["groups"] for row in g["rows"]}
    assert {UNIT_A, UNIT_B} <= admin_units


def test_xuat_excel_chi_co_don_vi_minh(env):
    r = client.get("/api/member/anomalies/export.xlsx", headers=env["lead"])
    assert r.status_code == 200
    cells = {str(c.value) for ws in load_workbook(BytesIO(r.content)).worksheets
             for row in ws.iter_rows() for c in row if c.value is not None}
    assert UNIT_A in cells and UNIT_B not in cells


def test_chi_lanh_dao_don_vi_vao_duoc(env):
    # Tài khoản nhập liệu, quản trị: không phải lãnh đạo đơn vị → 403 ở đường của lãnh đạo.
    assert client.get("/api/member/anomalies", headers=env["mem"]).status_code == 403
    assert client.get("/api/member/anomalies", headers=env["admin"]).status_code == 403
    # Lãnh đạo đơn vị KHÔNG vào được màn toàn hệ thống, cũng không chạm được cấu hình ngưỡng.
    assert client.get("/api/anomalies", headers=env["lead"]).status_code == 403
    assert client.get("/api/anomalies/config", headers=env["lead"]).status_code == 403
    assert client.put("/api/anomalies/config", json={"values": {}},
                      headers=env["lead"]).status_code == 403
    assert client.get("/api/member/anomalies").status_code == 401
    assert client.get("/api/member/anomalies", headers=env["lead"],
                      params={"date_from": "01/01/2026"}).status_code == 422


def test_lanh_dao_don_vi_nhan_thay_ca_canh_bao_cu_cua_don_vi_da_sap_nhap(env):
    """Số liệu trước sáp nhập đứng tên đơn vị CŨ → cảnh báo cũ cũng vậy; đơn vị nhận phải thấy."""
    from datetime import date, timedelta

    from sqlalchemy import text as sql

    from app.core.db import session_scope
    from app.core.market_meta import PURCHASE_SOURCE_UNIT
    from app.services import member_unit_merge, price_repo

    old = "_zz_an_don_vi_cu"
    day = (date.today() - timedelta(days=5)).isoformat()
    client.post("/api/member-units", json={"name": old}, headers=env["admin"])
    try:
        # Giá mủ gõ nhầm đồng/kg (45.000) khi đơn vị cũ còn hoạt động, rồi sáp nhập vào đơn vị A.
        price_repo.upsert_record({"as_of": day, "source": PURCHASE_SOURCE_UNIT, "grade": old,
                                  "contract": "", "price_type": "purchase", "price": 45000,
                                  "currency": "VND", "unit": "đồng/độ TSC"})
        member_unit_merge.merge(old, UNIT_A, (date.today() - timedelta(days=2)).isoformat())

        body = client.get("/api/member/anomalies", headers=env["lead"]).json()
        wrong = next(g for g in body["groups"] if g["key"] == "wrong_raw_price")
        assert {r["don_vi"] for r in wrong["rows"]} == {old}
        units = {r.get("don_vi") for g in body["groups"] for r in g["rows"]}
        assert UNIT_B not in units                  # mở theo sáp nhập, không mở sang đơn vị khác
    finally:
        with session_scope() as db:
            db.execute(sql("DELETE FROM fact_price WHERE grade = :g"), {"g": old})
            db.execute(sql("UPDATE member_unit SET merged_into = NULL, merged_at = NULL, "
                           "is_active = true WHERE name = :g"), {"g": old})
        client.delete(f"/api/member-units/{old}", headers=env["admin"])
