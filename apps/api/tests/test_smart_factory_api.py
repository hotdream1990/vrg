"""Test API Nhà máy thông minh: cấu hình kết nối (admin) + chỉ số theo ngày (quyền `smart_factory`).

Không có SQL Server thật trong CI → lớp đọc SCADA (`scada_client.read_meters` / `probe`) được
monkeypatch trả mẫu tổng hợp. Dữ liệu test (nhà máy `_zz_sf…`, tài khoản `sf_…`) + cache đọc SCADA
dọn sạch trước/sau test. Chặn 429/502 theo vai trò/400 chưa khai tag: test_smart_factory_api_hardening.
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.db import db_healthy, session_scope
from app.core.permissions import effective_caps, has_cap
from app.main import app
from app.services import scada_client, scada_daily_meters as dm, scada_factory_repo, user_repo
from app.services import scada_historian_sql as hsql
from app.services import scada_read_guard as guard

client = TestClient(app)
PREFIX = "_zz_sf"
USERS = ("sf_cap", "sf_nocap")
SECRET = "S3cret-pw!"
_db = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")


def body(name: str = f"{PREFIX} Phú Riềng", **kw) -> dict:
    return {"name": name, "host": "10.0.0.5", "port": 1433, "username": "scada_ro",
            "password": SECRET, "database": "Runtime", "linked_server": "INSQL",
            "energy_tags": ["PM_EnergyReal0", "PM_EnergyReal1", "PM_EnergyReal2", "PM_EnergyReal3"],
            "water_tag": "Water_TotalVolume", "bales_tag": "Bales_Count", "enabled": True, **kw}


def _bearer(username: str, password: str) -> dict[str, str]:
    res = client.post("/api/auth/login", json={"username": username, "password": password})
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def _cleanup() -> None:
    guard.clear_cache()
    with session_scope() as db:
        db.execute(text("DELETE FROM scada_factory WHERE name LIKE :p"), {"p": f"{PREFIX}%"})
    for u in USERS:
        try:
            user_repo.delete_user(u)
        except ValueError:
            pass


def setup_users() -> dict[str, str]:
    """Dọn dữ liệu test cũ + tạo 2 tài khoản thử → header Bearer của admin (dùng chung file test khác)."""
    user_repo.seed_admin()
    _cleanup()
    h = _bearer("admin", "admin")
    client.post("/api/users", json={"username": "sf_cap", "password": "pass123", "role": "editor",
                                    "permissions": ["smart_factory"]}, headers=h)
    client.post("/api/users", json={"username": "sf_nocap", "password": "pass123", "role": "editor",
                                    "permissions": ["inventory"]}, headers=h)
    return h


@pytest.fixture()
def admin():
    yield setup_users()
    _cleanup()


# ── Thuần (không cần DB) ──

def test_cap_is_single_level_and_given_to_executive() -> None:
    assert has_cap(effective_caps("executive", None), "smart_factory")
    assert has_cap(effective_caps("admin", None), "smart_factory")
    assert not has_cap(effective_caps("viewer", ["smart_factory"]), "smart_factory")


class _FakeErr(Exception):
    pass


_REFUSED = (b"DB-Lib error message 20009, severity 9:\nUnable to connect: TDS server is unavailable "
            b"or does not exist (10.0.0.5)\nNet-Lib error during Connection refused (61)\n"
            b"DB-Lib error message 20009, severity 9:\nUnable to connect: TDS server is unavailable "
            b"or does not exist (10.0.0.5)\nNet-Lib error during Connection refused (61)\n")


@pytest.mark.parametrize("arg, connected, expect", [
    # Hình dạng THẬT của pymssql: một tuple (mã, bytes) bọc trong args (đo trên SQL Server 2022).
    ((18456, b"Login failed for user 'scada_ro'. S3cret-pw!DB-Lib error message 20018, severity 14:"
              b"\nGeneral SQL Server error: Check messages from the SQL Server\n"),
     False, "Sai tài khoản hoặc mật khẩu"),
    ((20009, _REFUSED), False, "Không kết nối được tới máy chủ 10.0.0.5:1433"),
    # pymssql giữ mã của lần lỗi TRƯỚC: bị từ chối kết nối nhưng mang mã 18456 → vẫn là "không tới được".
    ((18456, _REFUSED), False, "Không kết nối được tới máy chủ 10.0.0.5:1433"),
    ((7202, b"Could not find server 'INSQL' in sys.servers."), True, "Không tìm thấy linked server «INSQL»"),
    ((207, b"Invalid column name 'Bales_Count'."), True, "Tag không tồn tại trong Historian"),
    ((4060, b"Cannot open database \"Runtime\" requested by the login."), False, "Không mở được database"),
    ((18456, b"Cannot open database \"Runtime\" requested by the login. The login failed.\n"
              b"Login failed for user 'scada_ro'."), False, "Không mở được database"),
    # Lỗi lúc truy vấn mang mã cũ 18456 → không được gán nhầm thành sai mật khẩu.
    ((18456, b"Some provider error"), True, "Lỗi khi đọc SQL Server của SCADA"),
    ((20003, b"Adaptive Server connection timed out"), True, "Truy vấn SCADA quá thời gian chờ"),
])
def test_friendly_error_readable_and_hides_password(arg, connected, expect) -> None:
    f = {"host": "10.0.0.5", "port": 1433, "username": "scada_ro", "password": SECRET,
         "database_name": "Runtime", "linked_server": "INSQL"}
    msg = scada_client.friendly_error(_FakeErr(arg), f, connected=connected)
    assert msg.startswith(expect) and SECRET not in msg
    # Thông điệp gốc lặp lại mật khẩu → bỏ HẲN phần chi tiết (không che bằng ***).
    assert ("Chi tiết:" in msg) is (SECRET.encode() not in arg[1]) and "***" not in msg
    assert "b'" not in msg and 'b"' not in msg  # đã decode, không lộ repr bytes
    assert msg.count("Unable to connect") <= 1  # dòng trùng đã gộp


@pytest.mark.parametrize("password", ["able", "ABLE"])
def test_friendly_error_weak_password_drops_detail_instead_of_masking(password) -> None:
    """*** làm lộ mật khẩu yếu ("able" → "Un***") → thấy mật khẩu trong thông điệp là bỏ chi tiết."""
    f = {"host": "10.0.0.5", "port": 1433, "username": "scada_ro", "password": password,
         "database_name": "Runtime", "linked_server": "INSQL"}
    msg = scada_client.friendly_error(_FakeErr((20009, _REFUSED)), f, connected=False)
    assert msg.startswith("Không kết nối được tới máy chủ 10.0.0.5:1433")
    assert "Chi tiết" not in msg and "***" not in msg and "able" not in msg.lower()


def test_friendly_error_uses_server_messages_when_exception_is_generic() -> None:
    # Đo thật trên SQL Server 2022 (30/09/2026): tag không tồn tại → exception chỉ giữ câu CUỐI,
    # câu có ích nằm ở thông điệp server hứng qua msghandler.
    f = {"host": "10.0.0.5", "port": 1433, "username": "scada_ro", "password": SECRET,
         "database_name": "Runtime", "linked_server": "INSQL"}
    exc = _FakeErr((8180, b"Statement(s) could not be prepared.DB-Lib error message 20018, "
                          b"severity 16:\nGeneral SQL Server error: Check messages from the SQL Server\n"))
    msgs = ['OLE DB provider "INSQL" for linked server "INSQL" returned message '
            '"Deferred prepare could not be completed.".',
            "Statement(s) could not be prepared.", "Invalid column name 'Khong_Co_Tag'."]
    msg = scada_client.friendly_error(exc, f, connected=True, server_msgs=msgs)
    assert msg.startswith("Tag không tồn tại trong Historian")
    assert "Khong_Co_Tag" in msg


def test_map_rows_case_insensitive_and_truncates_millis() -> None:
    rows = [{"DateTime": datetime(2026, 9, 1, 1, 0, 0, 3000), "WATER_TOTALVOLUME": 5.5},
            {"DateTime": "2026-09-01 00:00:00.0000000", "WATER_TOTALVOLUME": None}]
    out = hsql.map_rows(rows, ["Water_TotalVolume"])
    assert out == [(datetime(2026, 9, 1), {"Water_TotalVolume": None}),
                   (datetime(2026, 9, 1, 1), {"Water_TotalVolume": 5.5})]


# ── Cấu hình kết nối (admin) ──

@_db
def test_admin_crud_masks_password(admin) -> None:
    res = client.post("/api/smart-factory/admin/factories", json=body(), headers=admin)
    assert res.status_code == 200, res.text
    f = res.json()["factory"]
    assert "password" not in f and f["password_set"] is True and f["database"] == "Runtime"
    assert f["updated_by"] == "admin" and len(f["updated_at"]) == 19
    listed = client.get("/api/smart-factory/admin/factories", headers=admin).json()["factories"]
    assert all("password" not in x for x in listed) and SECRET not in str(listed)

    # PUT bỏ trống mật khẩu → giữ mật khẩu cũ; đổi các trường khác (không phải host/cổng/tài khoản)
    # vẫn lưu. Host/tài khoản chỉ khác hoa thường/khoảng trắng = không đổi.
    res = client.put(f"/api/smart-factory/admin/factories/{f['id']}",
                     json=body(password="", host=" 10.0.0.5 ", username="SCADA_RO", bales_tag=""),
                     headers=admin)
    assert res.status_code == 200, res.text
    assert res.json()["factory"]["bales_tag"] is None
    assert scada_factory_repo.get_factory(f["id"])["password"] == SECRET
    # Đổi host/cổng/tài khoản mà không nhập lại mật khẩu → 400 (chặn hứng mật khẩu ở máy lạ).
    for patch in ({"host": "10.0.0.66"}, {"port": 1434}, {"username": "sa"}):
        for pw in ("", None):
            res = client.put(f"/api/smart-factory/admin/factories/{f['id']}",
                             json=body(password=pw, **patch), headers=admin)
            assert res.status_code == 400, (patch, res.text)
            assert res.json()["detail"] == ("Đổi máy chủ, cổng hoặc tài khoản thì phải nhập lại "
                                            "mật khẩu.")
    assert scada_factory_repo.get_factory(f["id"])["host"] == "10.0.0.5"
    res = client.put(f"/api/smart-factory/admin/factories/{f['id']}",
                     json=body(host="10.0.0.6"), headers=admin)
    assert res.status_code == 200 and res.json()["factory"]["host"] == "10.0.0.6"
    client.put(f"/api/smart-factory/admin/factories/{f['id']}", json=body(password=None,
               host="10.0.0.6"), headers=admin)
    assert scada_factory_repo.get_factory(f["id"])["password"] == SECRET
    client.put(f"/api/smart-factory/admin/factories/{f['id']}", json=body(password="new-pw"), headers=admin)
    assert scada_factory_repo.get_factory(f["id"])["password"] == "new-pw"

    with session_scope() as db:
        logs = db.execute(text("SELECT action, data_before::text AS b, data_after::text AS a "
                               "FROM audit_log WHERE entity = 'scada_factory' AND entity_key = :k "
                               "ORDER BY id"), {"k": f["name"]}).mappings().all()
    assert [r["action"] for r in logs][:1] == ["create"]
    assert not any(SECRET in (r["a"] or "") + (r["b"] or "") or "new-pw" in (r["a"] or "") for r in logs)
    assert '"password_changed": true' in logs[-1]["a"]

    assert client.delete(f"/api/smart-factory/admin/factories/{f['id']}", headers=admin).json() == {"ok": True}
    assert client.delete(f"/api/smart-factory/admin/factories/{f['id']}", headers=admin).status_code == 404
    assert client.put(f"/api/smart-factory/admin/factories/{f['id']}", json=body(), headers=admin).status_code == 404


@_db
@pytest.mark.parametrize("patch", [
    {"energy_tags": ["A", "B", "C"]}, {"energy_tags": ["A", "B"]}, {"water_tag": "Water]; DROP"},
    {"bales_tag": "Bales'"}, {"linked_server": "IN-SQL"}, {"database": "Run time"},
    {"energy_tags": ["R0", "R1", "R2", "R 3"]}, {"password": ""}, {"name": "  "}, {"host": ""},
    {"port": 70000}, {"host": "10.0.0.5 extra"},
    # Tag trùng nhau (không phân biệt hoa thường), giữa mọi tag điện/nước/bành.
    {"water_tag": "pm_energyreal0"}, {"bales_tag": "WATER_TOTALVOLUME"},
    {"energy_tags": ["R0", "R1", "r0", "R3"]},
])
def test_admin_validation_400(admin, patch) -> None:
    res = client.post("/api/smart-factory/admin/factories", json=body(**patch), headers=admin)
    assert res.status_code == 400, res.text


@_db
def test_admin_duplicate_name_409(admin) -> None:
    a = client.post("/api/smart-factory/admin/factories", json=body(f"{PREFIX} A"), headers=admin).json()
    client.post("/api/smart-factory/admin/factories", json=body(f"{PREFIX} B"), headers=admin)
    assert client.post("/api/smart-factory/admin/factories", json=body(f"{PREFIX} A"),
                       headers=admin).status_code == 409
    assert client.put(f"/api/smart-factory/admin/factories/{a['factory']['id']}",
                      json=body(f"{PREFIX} B"), headers=admin).status_code == 409


@_db
def test_admin_endpoints_blocked_for_non_admin(admin) -> None:
    h = _bearer("sf_cap", "pass123")
    fid = client.post("/api/smart-factory/admin/factories", json=body(), headers=admin).json()["factory"]["id"]
    assert client.get("/api/smart-factory/admin/factories", headers=h).status_code == 403
    assert client.post("/api/smart-factory/admin/factories", json=body(f"{PREFIX} X"), headers=h).status_code == 403
    assert client.put(f"/api/smart-factory/admin/factories/{fid}", json=body(), headers=h).status_code == 403
    assert client.delete(f"/api/smart-factory/admin/factories/{fid}", headers=h).status_code == 403
    assert client.post(f"/api/smart-factory/admin/factories/{fid}/test", headers=h).status_code == 403
    assert client.get("/api/smart-factory/admin/factories").status_code == 401


@_db
def test_admin_test_connection_always_200(admin, monkeypatch) -> None:
    fid = client.post("/api/smart-factory/admin/factories", json=body(), headers=admin).json()["factory"]["id"]
    raw = {"PM_EnergyReal0": 0.0, "PM_EnergyReal1": 0.0, "PM_EnergyReal2": 3869.0,
           "PM_EnergyReal3": -9724.0, "Water_TotalVolume": 12473.64, "Bales_Count": 21397.0}
    server_time = dm.vn_now().isoformat(timespec="seconds") + ".1234567+07:00"
    monkeypatch.setattr(scada_client, "probe", lambda f: (
        "Microsoft SQL Server 2016", server_time, [(datetime(2026, 9, 30, 15, 39), raw)]))
    res = client.post(f"/api/smart-factory/admin/factories/{fid}/test", headers=admin).json()
    assert res["ok"] is True and res["server_version"] == "Microsoft SQL Server 2016"
    assert res["latest_at"] == "2026-09-30T15:39:00" and res["raw"] == raw
    assert res["values"] == {"energy_kwh": 253614.596, "water_m3": 12473.64, "bales": 21397}
    assert res["server_time"] == server_time[:19] + "+07:00" and res["warnings"] == []

    def boom(f):
        raise scada_client.ScadaError("Không kết nối được tới máy chủ 10.0.0.5:1433")
    monkeypatch.setattr(scada_client, "probe", boom)
    res = client.post(f"/api/smart-factory/admin/factories/{fid}/test", headers=admin)
    assert res.status_code == 200 and res.json() == {
        "ok": False, "detail": "Không kết nối được tới máy chủ 10.0.0.5:1433"}
    assert client.post("/api/smart-factory/admin/factories/999999999/test", headers=admin).status_code == 404


# ── Chỉ số theo ngày (quyền smart_factory) ──

def _fake_reader(calls: list):
    def read(factory, start, end):
        calls.append((start, end))
        hours = int(((end or start + timedelta(days=2)) - start).total_seconds() // 3600)
        rows = [(start + timedelta(hours=h),
                 {"PM_EnergyReal0": 0, "PM_EnergyReal1": 0, "PM_EnergyReal2": 1,
                  "PM_EnergyReal3": 10000 + 10 * h, "Water_TotalVolume": 50.0 + h,
                  "Bales_Count": 100 + 2 * h}) for h in range(hours + 1)]
        return rows, []
    return read


@_db
def test_meters_daily_with_fake_scada(admin, monkeypatch) -> None:
    h = _bearer("sf_cap", "pass123")
    f = client.post("/api/smart-factory/admin/factories", json=body(), headers=admin).json()["factory"]
    off = client.post("/api/smart-factory/admin/factories",
                      json=body(f"{PREFIX} tắt", enabled=False), headers=admin).json()["factory"]
    listed = client.get("/api/smart-factory/factories", headers=h).json()["factories"]
    mine = [x for x in listed if x["name"].startswith(PREFIX)]
    assert mine == [{"id": f["id"], "name": f["name"], "metrics": ["energy", "water", "bales"]}]

    calls: list = []
    monkeypatch.setattr(scada_client, "read_meters", _fake_reader(calls))
    q = f"factory_id={f['id']}&date_from=2026-09-01&date_to=2026-09-02"
    res = client.get(f"/api/smart-factory/meters/daily?{q}", headers=h)
    assert res.status_code == 200, res.text
    rep = res.json()
    # Kỳ đã qua → vế trên = D+1 00:00 + 1 giờ; mẫu sau 00:00 (reader giả trả tới 01:00) bị bỏ.
    assert calls[-1] == (datetime(2026, 9, 1), datetime(2026, 9, 3, 1))
    assert [r["date"] for r in rep["rows"]] == ["2026-09-01", "2026-09-02"]
    e = rep["rows"][0]["energy"]
    assert e["flag"] is None and e["used"] == 0.24 and e["close_at"] == "2026-09-02T00:00:00"
    assert rep["rows"][0]["bales"]["used"] == 48 and rep["summary"]["water"]["total"] == 48.0
    assert rep["rows"][1]["energy"]["close_at"] == "2026-09-03T00:00:00"
    assert rep["summary"]["water"]["days_complete"] == 2
    assert set(rep) == {"factory", "date_from", "date_to", "fetched_at", "metrics", "rows",
                        "summary", "intensity"}

    n_calls = len(calls)
    xlsx = client.get(f"/api/smart-factory/meters/daily.xlsx?{q}", headers=h)
    assert xlsx.status_code == 200 and len(calls) == n_calls  # dùng lại cache 60s, không đọc lại SCADA
    assert "chi-so-dien-nuoc-banh_zz-sf-phu-rieng_2026-09-01_2026-09-02.xlsx" in \
        xlsx.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    assert "Phú Riềng" in ws["A1"].value

    # Không truyền kỳ → 30 ngày gần nhất tính cả hôm nay, vế cuối = GetDate() (end=None).
    client.get(f"/api/smart-factory/meters/daily?factory_id={f['id']}", headers=h)
    today = dm.vn_now().date()
    assert calls[-1] == (datetime.combine(today - timedelta(days=29), datetime.min.time()), None)

    base = "/api/smart-factory/meters/daily?factory_id="
    assert client.get(f"{base}{off['id']}", headers=h).status_code == 404
    assert client.get(f"{base}999999999", headers=h).status_code == 404
    assert client.get(f"{base}{f['id']}&date_from=2026-09-10&date_to=2026-09-01",
                      headers=h).status_code == 400
    assert client.get(f"{base}{f['id']}&date_from=2026-01-01&date_to=2026-06-30",
                      headers=h).status_code == 400
    assert client.get(f"{base}{f['id']}&date_from=2026-13-01", headers=h).status_code == 400
    # date_to tương lai → kẹp về hôm nay (không 400).
    future = (today + timedelta(days=5)).isoformat()
    res = client.get(f"{base}{f['id']}&date_from={today.isoformat()}&date_to={future}", headers=h)
    assert res.status_code == 200 and res.json()["date_to"] == today.isoformat()
    # date_from tương lai → 400 đúng câu (không phải "từ ngày sau đến ngày").
    res = client.get(f"{base}{f['id']}&date_from={future}", headers=h)
    assert res.status_code == 400 and res.json()["detail"] == "Từ ngày không được sau hôm nay."


@_db
def test_meters_blocked_without_cap(admin) -> None:
    h = _bearer("sf_nocap", "pass123")
    assert client.get("/api/smart-factory/factories", headers=h).status_code == 403
    assert client.get("/api/smart-factory/meters/daily?factory_id=1", headers=h).status_code == 403
    assert client.get("/api/smart-factory/meters/daily.xlsx?factory_id=1", headers=h).status_code == 403
    assert client.get("/api/smart-factory/meters/daily?factory_id=1").status_code == 401
