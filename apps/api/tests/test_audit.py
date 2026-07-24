"""Test Nhật ký hoạt động: ghi vết khi sửa/xoá số liệu + tra cứu + không lộ bí mật."""

from __future__ import annotations

from datetime import date, timedelta

import pytest
from sqlalchemy import text

from app.core import request_ctx
from app.core.db import db_healthy, session_scope
from app.services import audit_repo, inventory_repo, price_repo

pytestmark = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")

_DAY = (date.today() - timedelta(days=2)).isoformat()
_GRADE = "__audit_test__"


@pytest.fixture(autouse=True)
def _cleanup():
    """Dọn dữ liệu + nhật ký của riêng test này (không đụng dữ liệu thật)."""
    yield
    with session_scope() as db:
        db.execute(text("DELETE FROM fact_price WHERE grade = :g"), {"g": _GRADE})
        db.execute(text("DELETE FROM fact_inventory WHERE as_of = CAST(:d AS date) "
                        "AND note = 'audit-test'"), {"d": _DAY})
        db.execute(text("DELETE FROM audit_log WHERE entity_key LIKE :k OR note = 'audit-test'"),
                   {"k": f"%{_GRADE}%"})


def _rec(price: float) -> dict:
    return {"as_of": _DAY, "source": "vrg", "grade": _GRADE, "contract": "",
            "price_type": "purchase", "price": price, "currency": "VND", "unit": "đồng/độ TSC"}


def _rows(**filters) -> list[dict]:
    return audit_repo.search(q=_GRADE, limit=50, **filters)["items"]


def test_create_update_delete_leaves_full_trail() -> None:
    """Thêm → sửa → xoá đều để lại vết, có giá trị trước/sau đúng."""
    with request_ctx.use_actor("tester"):
        price_repo.upsert_record(_rec(385))
        price_repo.upsert_record(_rec(405))
        price_repo.delete_record(_DAY, "vrg", _GRADE, "", "purchase")

    rows = _rows()
    assert [r["action"] for r in rows] == ["delete", "update", "create"]  # mới nhất trước
    assert all(r["actor"] == "tester" for r in rows)

    created, updated, deleted = rows[2], rows[1], rows[0]
    assert created["before"] is None and created["after"]["price"] == 385
    assert updated["before"]["price"] == 385 and updated["after"]["price"] == 405
    assert deleted["before"]["price"] == 405 and deleted["after"] is None  # xoá vẫn giữ số cũ
    assert deleted["entity"] == "raw_material" and deleted["company"] == _GRADE
    assert deleted["as_of"] == _DAY


def test_save_without_change_is_not_logged() -> None:
    """Bấm lưu nhưng không đổi gì → không sinh dòng nhật ký (tránh nhiễu)."""
    with request_ctx.use_actor("tester"):
        price_repo.upsert_record(_rec(400))
        price_repo.upsert_record(_rec(400))
    assert len(_rows()) == 1


def test_search_filters_by_actor_entity_and_action() -> None:
    """Lọc theo người · nhóm số liệu · loại thao tác đều thu hẹp đúng kết quả."""
    with request_ctx.use_actor("nguoi_a"):
        price_repo.upsert_record(_rec(300))
    with request_ctx.use_actor("nguoi_b"):
        price_repo.upsert_record(_rec(310))

    assert len(_rows()) == 2
    assert [r["actor"] for r in _rows(actor="nguoi_b")] == ["nguoi_b"]
    assert len(_rows(entity="raw_material")) == 2
    assert _rows(entity="inventory") == []
    assert [r["action"] for r in _rows(action="create")] == ["create"]


def test_note_context_marks_source_of_data() -> None:
    """`use_note` gắn nguồn thao tác (vd nhập từ file) cho mọi dòng bên trong."""
    with request_ctx.use_actor("tester"), request_ctx.use_note("Nhập từ file Excel"):
        price_repo.upsert_record(_rec(320))
    assert _rows()[0]["note"] == "Nhập từ file Excel"


def test_paused_context_skips_logging() -> None:
    """Dữ liệu phái sinh (mirror) không sinh nhật ký."""
    with request_ctx.use_actor("tester"), request_ctx.paused():
        price_repo.upsert_record(_rec(330))
    assert _rows() == []


def test_other_entities_are_logged() -> None:
    """Tồn kho (bảng khác) cũng được ghi vết với đúng nhóm số liệu."""
    with request_ctx.use_actor("tester"):
        inventory_repo.upsert(_DAY, 1000.0, 200.0, note="audit-test")
    rows = audit_repo.search(entity="inventory", date_from=_DAY, limit=10)["items"]
    assert any(r["as_of"] == _DAY and r["after"]["ton_kho"] == 1000.0 for r in rows)


def test_export_xlsx_has_readable_change_column() -> None:
    """Xuất Excel: đúng số dòng + cột 'Nội dung thay đổi' đọc được (kèm biên bản đối chiếu)."""
    import io

    from openpyxl import load_workbook

    from app.services import audit_export

    rows = [
        {"at": "2026-07-23T14:02:11+00:00", "actor": "nguyenvana", "actor_role": "editor",
         "on_behalf": "", "entity_label": "Giá mủ nguyên liệu", "action": "update",
         "action_label": "Sửa", "entity_key": "2026-07-23|vrg|X|purchase", "as_of": "2026-07-23",
         "company": "X", "before": {"price": 385}, "after": {"price": 405}, "ip": "10.0.0.1",
         "note": ""},
        {"at": "2026-07-23T15:00:00+00:00", "actor": "admin", "actor_role": "admin",
         "on_behalf": "", "entity_label": "Tồn kho", "action": "delete", "action_label": "Xoá",
         "entity_key": "2026-07-23", "as_of": "2026-07-23", "company": "",
         "before": {"ton_kho": 1000}, "after": None, "ip": "", "note": "xoá tuần"},
    ]
    ws = load_workbook(io.BytesIO(audit_export.build_xlsx(rows))).active
    assert ws.max_row == 3  # 1 dòng tiêu đề + 2 dòng dữ liệu
    assert ws.cell(row=1, column=1).value == "Thời điểm"
    assert ws.cell(row=2, column=10).value == "price: 385 → 405"
    assert ws.cell(row=3, column=10).value == "ton_kho: 1000"  # xoá → chỉ liệt kê số đã mất
