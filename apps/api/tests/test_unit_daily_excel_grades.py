"""Chủng loại mủ nguyên liệu trên ĐƯỜNG EXCEL của biểu Thu mua (xuất · nhập).

Khoá lại 3 điều quan trọng nhất:
1. Xuất số liệu ra file rồi nhập lại KHÔNG đổi số (kể cả khi đã tách chủng loại).
2. File theo mẫu CŨ (không có cột chủng loại) vẫn nhập được y như trước — 67 đơn vị đang dùng
   mẫu cũ, gãy chỗ này là gãy toàn bộ đường nhập liệu.
3. Dòng trùng chủng loại được GỘP, dòng trống bị bỏ, chủng loại lạ bị BÁO LỖI (không im lặng bỏ).

Các test không cần DB dùng monkeypatch cho 2 hàm tra DB của bộ đọc; chỉ test end-to-end cuối cùng
(ghi thật qua `commit_rows`) mới cần DB.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.db import db_healthy, session_scope
from app.services import (member_unit_repo, unit_daily_excel_io, unit_daily_excel_purchase,
                          unit_daily_fields, unit_daily_repo)

UNIT = "_zzX_excel_grades"
AS_OF = "2026-09-18"
AS_OF_OLD = "2026-09-17"
GRADE_TITLE = "Chủng loại mủ nguyên liệu"


@pytest.fixture()
def offline(monkeypatch: pytest.MonkeyPatch):
    """Bộ đọc chỉ chạm DB ở 2 chỗ (danh mục đơn vị + đánh dấu tạo/ghi đè) → cắt ra để test thuần."""
    monkeypatch.setattr(member_unit_repo, "active_names", lambda: [UNIT])
    monkeypatch.setattr(unit_daily_repo, "has_entry", lambda *a, **k: False)


def _xlsx(rows: list[dict]) -> bytes:
    """Các dòng → file Excel ĐÚNG mẫu (tài khoản 1 đơn vị nên mẫu không có cột 'Đơn vị')."""
    return unit_daily_excel_io.build_template("purchase", allowed_units=[UNIT], rows=rows)


def _parse(data: bytes) -> dict:
    return unit_daily_excel_io.parse_upload("purchase", data, allowed_units=[UNIT])


def _drop_grade_column(data: bytes) -> bytes:
    """Dựng lại file theo MẪU CŨ: xoá hẳn cột chủng loại mủ nguyên liệu."""
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Thu mua"]
    col = next(i for i in range(1, ws.max_column + 1)
               if str(ws.cell(row=5, column=i).value or "").strip() == GRADE_TITLE)
    ws.delete_cols(col)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _saved(rows: list[dict], current: dict | None = None) -> dict:
    """Dòng đã đọc → payload đúng như lúc lưu (upsert luôn đi qua `clean_fields`)."""
    fields, _ = unit_daily_excel_purchase.fields_from_rows(rows, current)
    return unit_daily_fields.clean_fields("purchase", fields)


def test_xuat_roi_nhap_lai_khong_doi_so(offline) -> None:
    """Xuất số liệu ĐÃ TÁCH CHỦNG LOẠI ra file mẫu → nhập lại → payload y hệt."""
    stored = unit_daily_fields.clean_fields("purchase", {
        "latex_grades": [{"grade": "SVR 3L", "qty": 10.0}, {"grade": "SVR 10 Mix", "qty": 5.5}],
        "cup_grades": [{"grade": "SVR 10 Mix", "qty": 3.25}],
        "lace": 2.5,                                    # mủ dây chưa tách chủng loại
        "finished": [{"grade": "SVR CV 50", "qty": 30.0, "price": 41.2, "ccy": "VND"}],
    })
    rows = unit_daily_excel_purchase.export_rows(
        stored, {"price_latex": 540, "price_cup": 510, "price_lace": 260})
    # Mỗi chủng loại 1 dòng (mủ nước 2 + mủ chén dùng chung dòng "SVR 10 Mix") + 1 dòng tổng mủ dây.
    assert [r.get("material_grade") for r in rows] == ["SVR 3L", "SVR 10 Mix", None]
    assert rows[1]["coagulum"] == 3.25    # mủ chén nằm đúng cột của nó
    assert rows[2] == {"lace": 2.5}

    got = _parse(_xlsx([dict(r, as_of=AS_OF) for r in rows]))
    assert got["summary"] == {"total": 3, "ok": 3, "error": 0}, got["rows"]
    assert _saved(got["rows"]) == stored
    # Đơn giá là số của CẢ NGÀY, đi ra ngoài payload (kho giá mủ nguyên liệu) → đọc lại phải còn.
    assert got["rows"][0]["price_latex"] == 540 and got["rows"][0]["price_lace"] == 260


def test_ngay_chua_tach_chung_loai_xuat_dung_mot_dong_tong(offline) -> None:
    """Số liệu cũ chỉ có ô tổng → xuất ra ĐÚNG 1 dòng tổng, không bịa chủng loại nào."""
    rows = unit_daily_excel_purchase.export_rows({"latex_wet": 120.0, "coagulum": 50.0})
    assert rows == [{"latex_wet": 120.0, "coagulum": 50.0}]
    got = _parse(_xlsx([dict(rows[0], as_of=AS_OF)]))
    assert got["summary"]["error"] == 0
    assert _saved(got["rows"]) == {"latex_wet": 120.0, "coagulum": 50.0}


def test_file_mau_cu_khong_co_cot_chung_loai_van_nhap_duoc(offline) -> None:
    """Mẫu CŨ (chưa có cột chủng loại) → mọi dòng là dòng tổng, ô đã nhập trên web được giữ."""
    data = _drop_grade_column(_xlsx([{"as_of": AS_OF, "latex_wet": 120.0, "coagulum": 50.0,
                                      "lace": 15.0, "price_latex": 540}]))
    got = _parse(data)
    assert got["summary"] == {"total": 1, "ok": 1, "error": 0}, got["rows"]
    assert got["rows"][0]["material_grade"] is None     # thiếu cột ⇒ dòng tổng, không phải lỗi
    # `fx_purchase` chỉ có trên web (file không có cột này) → MERGE phải giữ nguyên.
    saved = _saved(got["rows"], {"fx_purchase": 1.25, "latex_grades": [{"grade": "SVR 3L",
                                                                       "qty": 999.0}]})
    assert saved["latex_wet"] == 120.0 and saved["coagulum"] == 50.0 and saved["lace"] == 15.0
    assert saved["fx_purchase"] == 1.25
    # Khai lại bằng ô tổng ⇒ bảng chủng loại cũ bị bỏ, không để tổng lệch tổng các dòng.
    assert "latex_grades" not in saved


def test_dong_trung_chung_loai_duoc_gop_va_dong_trong_bi_bo(offline) -> None:
    """2 dòng cùng chủng loại → cộng lại thành 1; dòng chỉ có chủng loại mà không có số → bỏ."""
    rows = [{"as_of": AS_OF, "material_grade": "SVR 3L", "latex_wet": 10.0},
            {"as_of": AS_OF, "material_grade": "SVR 3L", "latex_wet": 2.5},
            {"as_of": AS_OF, "material_grade": "SVR 10 Mix"},          # không có số → bỏ
            {"as_of": AS_OF, "material_grade": "SVR 3L", "coagulum": 4.0}]
    got = _parse(_xlsx(rows))
    assert got["summary"]["error"] == 0, got["rows"]
    saved = _saved(got["rows"])
    assert saved["latex_grades"] == [{"grade": "SVR 3L", "qty": 12.5}]
    assert saved["latex_wet"] == 12.5                  # ô tổng là số SUY RA = tổng các dòng
    assert saved["cup_grades"] == [{"grade": "SVR 3L", "qty": 4.0}]
    assert saved["coagulum"] == 4.0


def test_chung_loai_la_bi_bao_loi_khong_im_lang_bo(offline) -> None:
    """Gõ sai tên chủng loại → báo lỗi ĐÚNG dòng đó, không lặng lẽ bỏ số liệu."""
    got = _parse(_xlsx([{"as_of": AS_OF, "material_grade": "SVR 999", "latex_wet": 10.0}]))
    assert got["summary"] == {"total": 1, "ok": 0, "error": 1}
    assert any("SVR 999" in e for e in got["rows"][0]["_errors"]), got["rows"][0]["_errors"]


def test_vua_co_dong_chung_loai_vua_co_o_tong_thi_canh_bao(offline) -> None:
    """Khai cả 2 kiểu cho cùng loại mủ → lấy tổng các dòng chủng loại và NÓI RÕ cho người nhập."""
    rows = [{"material_grade": "SVR 3L", "latex_wet": 10.0}, {"latex_wet": 99.0}]
    fields, warnings = unit_daily_excel_purchase.fields_from_rows(rows, None)
    assert unit_daily_fields.clean_fields("purchase", fields)["latex_wet"] == 10.0
    assert warnings and "mủ nước" in warnings[0]


@pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")
def test_ghi_that_qua_commit_rows() -> None:
    """End-to-end: đọc file có dòng chủng loại → GHI THẬT → đọc lại đúng bảng + ô tổng suy ra."""
    member_unit_repo.add_unit(UNIT)
    try:
        data = _xlsx([{"as_of": AS_OF, "material_grade": "SVR 3L", "latex_wet": 10.0,
                       "price_latex": 540},
                      {"as_of": AS_OF, "material_grade": "SVR 10 Mix", "latex_wet": 5.5,
                       "coagulum": 3.0}])
        got = _parse(data)
        assert got["summary"]["error"] == 0, got["rows"]
        res = unit_daily_excel_io.commit_rows("purchase", got["rows"], "test",
                                              allowed_units=[UNIT])
        assert res["saved"] == 1 and res["warnings"] == []
        fields = unit_daily_repo.entries_on("purchase", AS_OF)[UNIT]["fields"]
        assert fields["latex_grades"] == [{"grade": "SVR 3L", "qty": 10.0},
                                          {"grade": "SVR 10 Mix", "qty": 5.5}]
        assert fields["latex_wet"] == 15.5 and fields["coagulum"] == 3.0
        assert fields["cup_grades"] == [{"grade": "SVR 10 Mix", "qty": 3.0}]

        # Cùng bộ mã đó phải ghi được file theo MẪU CŨ (không cột chủng loại): ô tổng + đơn giá.
        old = _parse(_drop_grade_column(_xlsx([{"as_of": AS_OF_OLD, "latex_wet": 120.0,
                                                "lace": 15.0, "price_latex": 540,
                                                "price_lace": 260}])))
        assert old["summary"]["error"] == 0, old["rows"]
        assert unit_daily_excel_io.commit_rows("purchase", old["rows"], "test",
                                               allowed_units=[UNIT])["saved"] == 1
        prev = unit_daily_repo.entries_on("purchase", AS_OF_OLD)[UNIT]["fields"]
        assert prev["latex_wet"] == 120.0 and prev["lace"] == 15.0
        assert "latex_grades" not in prev                 # không bịa chủng loại cho dòng tổng
        with session_scope() as db:
            px = db.execute(text(
                "SELECT price_type, price FROM fact_price WHERE grade = :u "
                "AND as_of = CAST(:d AS date) ORDER BY price_type"),
                {"u": UNIT, "d": AS_OF_OLD}).all()
        assert [(t, float(p)) for t, p in px] == [("purchase", 540.0), ("purchase_lace", 260.0)]
    finally:
        with session_scope() as db:
            db.execute(text("DELETE FROM unit_daily_report WHERE company = :u"), {"u": UNIT})
            db.execute(text("DELETE FROM fact_price WHERE grade = :u"), {"u": UNIT})
        member_unit_repo.delete_unit(UNIT)
