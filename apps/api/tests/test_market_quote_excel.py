"""Test xuất Excel Phiếu báo giá mủ: nội dung theo mẫu Ban TTKD + endpoint chỉ xuất phiếu ĐÃ LƯU."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from app.core.db import db_healthy, session_scope
from app.core.security import create_access_token
from app.main import app
from app.services import market_quote_excel as xl
from app.services import market_quote_excel_extra as extra
from app.services import user_repo

client = TestClient(app)
_DAY = "2099-02-03"
_db = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")


def _quote() -> dict:
    return {
        "as_of": _DAY, "fx": {"mua_tm": 26050, "mua_ck": 26080, "ban": 26370},
        "domestic_private": {"prices": {"SVR 3L": 68500000, "LATEX": 44300000},
                             "packaging": {"SVR 3L": "Hàng rời"}, "shipping": {"SVR 3L": "Tại Kho"},
                             "status": {"SVR 3L": "Sôi động"}, "note": "ghi chú 1"},
        "domestic_export": {"prices": {"LATEX": 1740.5}},
        "export_vrg": {"prices": {"LATEX": 1720}},
        "domestic_vrg": {"prices": {"SVR 3L": 68400000}, "note": '=HYPERLINK("http://x","click")'},
        "customer_proposal": {"items": [
            {"grade": "SVR 3L", "customer": "KH A", "qty": 50, "price_vnd": 68200000, "price_usd": 2700.5,
             "source": "Công ty X"},
            {"grade": "SVR 3L", "customer": "KH B", "qty": 150, "price_vnd": 68800000},
            {"grade": "", "customer": "", "qty": None, "price_vnd": None}],  # dòng trống → không xuất
            "note": "ghi chú KH"},
        "private_prices": {"Long Hòa": {"price": 585, "price_max": None}, "Mai Thảo": {"price": 586, "price_max": 589}},
        "private_processing_cost": None, "footer": "chân phiếu",
    }


def _sheet(q: dict, units: list[str] | None = None):
    return load_workbook(io.BytesIO(xl.build_xlsx(q, units or ["Long Hòa", "Mai Thảo"]))).active


def _row_of(ws, text_: str, col: str = "A") -> int:
    return next(c.row for c in ws[col] if isinstance(c.value, str) and c.value.startswith(text_))


def test_header_sections_and_units_follow_template() -> None:
    ws = _sheet(_quote())
    assert ws["A1"].value == "PHIẾU BÁO GIÁ MỦ THỊ TRƯỜNG" and ws["A2"].value == "Phiếu 03/02/2099"
    assert (ws["B5"].value, ws["B6"].value, ws["B7"].value) == (26050, 26080, 26370)
    # Mục 1: SVR 3L đúng hàng, đủ cột bao bì · vận chuyển · tình trạng · đơn giá
    r = _row_of(ws, "SVR 3L")
    assert [ws.cell(r, c).value for c in range(1, 6)] == ["SVR 3L", "Hàng rời", "Tại Kho", "Sôi động", 68500000]
    assert "USD/tấn FOB" in ws[f"A{_row_of(ws, 'MỤC 2.')}"].value
    # Mục 6: giá thành = giá × 1,08 × 100.000 + chi phí gia công mặc định 2.000.000
    r = _row_of(ws, "Long Hòa")
    assert (ws[f"B{r}"].value, ws[f"D{r}"].value, ws[f"E{r}"].value) == (585, 65_180_000, None)
    r = _row_of(ws, "Mai Thảo")
    assert (ws[f"D{r}"].value, ws[f"E{r}"].value) == (65_288_000, 65_612_000)


def test_proposal_rows_amount_and_comparison_use_weighted_average() -> None:
    ws = _sheet(_quote())
    head = _row_of(ws, "MỤC 5.")
    rows = [[ws.cell(r, c).value for c in range(1, 8)] for r in (head + 2, head + 3)]
    assert rows[0] == ["SVR 3L", "KH A", 50, 68200000, 2700.5, "Công ty X", 3_410_000_000]
    assert rows[1][:4] == ["SVR 3L", "KH B", 150, 68800000] and rows[1][6] == 10_320_000_000
    assert ws.cell(head + 4, 1).value == "Ghi chú:"            # dòng đề xuất trống đã bị bỏ
    # Bảng đối chiếu: đề xuất SVR 3L = (50×68,2tr + 150×68,8tr) / 200 = 68,65tr; chênh với giá tư nhân 68,5tr
    r = max(c.row for c in ws["A"] if c.value == "SVR 3L")       # lần xuất hiện cuối = bảng đối chiếu
    assert ws[f"J{r}"].value == 68_650_000 and ws[f"K{r}"].value == 150_000
    assert ws[f"D{r}"].value == -100_000 and ws[f"E{r}"].value == pytest.approx(-100_000 / 68_500_000)
    lr = max(c.row for c in ws["A"] if c.value == "LATEX")
    assert ws[f"H{lr}"].value == pytest.approx(-20.5)           # xuất khẩu: VRG − tư nhân


def test_numbers_are_values_not_formulas_and_text_is_neutralised() -> None:
    """Ô tự tính là GIÁ TRỊ (xem trước/pandas không trống); chữ người dùng gõ không thành công thức."""
    ws = _sheet(_quote())
    for row in ws.iter_rows():
        for c in row:
            assert not (isinstance(c.value, str) and c.value.startswith("=")), c.coordinate
    assert any(c.value == '\'=HYPERLINK("http://x","click")' for row in ws.iter_rows() for c in row)


def test_empty_quote_still_builds_all_sections() -> None:
    ws = _sheet({"as_of": _DAY}, units=[])
    assert [_row_of(ws, f"MỤC {i}.") for i in range(1, 7)] == sorted(_row_of(ws, f"MỤC {i}.") for i in range(1, 7))
    assert any(c.value == "Chưa có đề xuất nào trong phiếu." for row in ws.iter_rows() for c in row)
    assert any(c.value == "Chưa có đơn vị tư nhân nào báo giá." for row in ws.iter_rows() for c in row)


def test_proposal_price_by_grade_rules() -> None:
    items = [{"grade": "A", "qty": 100, "price_vnd": 10}, {"grade": "A", "qty": 300, "price_vnd": 20},
             {"grade": "A", "qty": None, "price_vnd": 999},   # không số lượng → bỏ qua khi có dòng có số lượng
             {"grade": "B", "qty": None, "price_vnd": 10}, {"grade": "B", "qty": None, "price_vnd": 20},
             {"grade": "C", "qty": 5, "price_vnd": None}, {"grade": "", "qty": 1, "price_vnd": 5}]
    assert extra.proposal_price_by_grade(items) == {"A": 18, "B": 15}   # A: 17,5 làm tròn nửa LÊN


# ── Qua HTTP (cần DB) ──

@pytest.fixture
def _admin():
    if db_healthy():
        user_repo.seed_admin()
    yield {"Authorization": f"Bearer {create_access_token('admin')}"}
    if db_healthy():
        with session_scope() as db:
            db.execute(text("DELETE FROM market_quote WHERE as_of = CAST(:d AS date)"), {"d": _DAY})
            db.execute(text("DELETE FROM fact_price WHERE as_of = CAST(:d AS date)"), {"d": _DAY})


@_db
def test_endpoint_exports_only_the_saved_quote(_admin) -> None:
    url = f"/api/market-quote/{_DAY}/export.xlsx"
    res = client.get(url, headers=_admin)
    assert res.status_code == 404 and "lưu phiếu trước" in res.json()["detail"]   # chưa lưu → không xuất
    assert client.get("/api/market-quote/khong-phai-ngay/export.xlsx", headers=_admin).status_code == 400

    assert client.put("/api/market-quote", json=_quote(), headers=_admin).status_code == 200
    res = client.get(url, headers=_admin)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert f"bao-gia-mu-thi-truong-{_DAY}.xlsx" in res.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws["A2"].value == "Phiếu 03/02/2099" and ws["B5"].value == 26050
    assert client.get(url).status_code in (401, 403)                              # phải đăng nhập
