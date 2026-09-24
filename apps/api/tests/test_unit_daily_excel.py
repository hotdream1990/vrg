"""Test nhập liệu bằng Excel cho báo cáo đơn vị — tải mẫu · xem trước · ghi.

Hai điểm quan trọng nhất được khoá lại ở đây:
1. Tài khoản chỉ quản 1 đơn vị → mẫu KHÔNG có cột 'Đơn vị', khi nhập server tự gán đúng đơn vị đó.
2. Không thể lách quyền: client sửa payload (xoá cờ lỗi) để ghi cho đơn vị khác vẫn bị server chặn.
"""

from __future__ import annotations

import io
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from sqlalchemy import text

from app.core.db import db_healthy, session_scope
from app.core.feature_flags import EXCEL_IMPORT_ENABLED
from app.main import app
from app.services import user_repo

pytestmark = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")

# Nhập Excel đang TẠM TẮT (app/core/feature_flags.py). Giữ nguyên các test luồng nhập —
# chúng tự chạy lại khi bật cờ; khi tắt thì chỉ còn test khẳng định endpoint đã đóng.
_excel_on = pytest.mark.skipif(not EXCEL_IMPORT_ENABLED,
                               reason="Nhập liệu bằng Excel đang tạm tắt")

client = TestClient(app)
UNIT = "_zz_xl_unit"
OTHER = "_zz_xl_other"
MEMBER = "xl_mem"


@pytest.fixture(autouse=True)
def _seed():
    user_repo.seed_admin()


def _admin() -> dict[str, str]:
    tok = client.post("/api/auth/login", json={"username": "admin", "password": "admin"}).json()
    return {"Authorization": f"Bearer {tok['access_token']}"}


# Tiêu đề cột đúng như file mẫu (bộ đọc dò cột THEO TIÊU ĐỀ nên bắt buộc phải có dòng này).
_PURCHASE_HEAD = ["Đơn vị", "Ngày", "SL thu mua mủ nước", "SL thu mua mủ chén",
                  "Đơn giá mủ nước", "Đơn giá mủ chén",
                  "Chủng loại thành phẩm", "SL thu mua thành phẩm", "Đơn giá thành phẩm",
                  "Đơn giá thành phẩm bằng"]


def _rows_to_xlsx(sheet: str, head: list[str], rows: list[tuple], header_row: int = 5) -> bytes:
    """Dựng file .xlsx giống mẫu: tiêu đề ở `header_row`, dữ liệu từ header_row+2."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for c, t in enumerate(head, start=1):
        ws.cell(row=header_row, column=c, value=t)
    for i, r in enumerate(rows, start=header_row + 2):
        for c, v in enumerate(r, start=1):
            ws.cell(row=i, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@_excel_on
def test_excel_import_single_unit_and_permission() -> None:
    h = _admin()
    today = date.today().isoformat()
    client.delete(f"/api/users/{MEMBER}", headers=h)
    for u in (UNIT, OTHER):
        client.post("/api/member-units", json={"name": u}, headers=h)
    client.post("/api/users", json={"username": MEMBER, "password": "pass123", "role": "member",
                                    "member_units": [UNIT]}, headers=h)
    mh = {"Authorization": "Bearer " + client.post(
        "/api/auth/login", json={"username": MEMBER, "password": "pass123"}
    ).json()["access_token"]}

    # 1) Mẫu của tài khoản 1 đơn vị: KHÔNG còn cột 'Đơn vị', ghi rõ áp dụng cho đơn vị nào.
    tpl = client.get("/api/member/import/template?kind=purchase", headers=mh)
    assert tpl.status_code == 200
    ws = load_workbook(io.BytesIO(tpl.content))["Thu mua"]
    titles = [ws.cell(row=5, column=c).value for c in range(1, 10)]
    assert "Đơn vị" not in titles and "Ngày" in titles
    assert UNIT in str(ws.cell(row=2, column=1).value)

    # 2) Nhập file không có cột 'Đơn vị' → server tự gán đơn vị của tài khoản.
    data = _rows_to_xlsx("Thu mua", _PURCHASE_HEAD[1:],   # mẫu 1 đơn vị: bỏ cột "Đơn vị"
                          [(today, 12.5, 4.5, 500, 450, "SVR CV 50", 9, 40.5, "VND")])
    prev = client.post("/api/member/import/preview?kind=purchase", headers=mh,
                       files={"file": ("f.xlsx", data)}).json()
    assert prev["summary"] == {"total": 1, "ok": 1, "error": 0}
    assert prev["rows"][0]["company"] == UNIT
    assert client.post("/api/member/import/commit", headers=mh,
                       json={"kind": "purchase", "rows": prev["rows"]}).json()["saved"] == 1

    # 3) Không lách được quyền: dùng mẫu có cột 'Đơn vị' điền đơn vị KHÁC.
    admin_tpl = client.get("/api/unit-daily/import/template?kind=purchase", headers=h)
    ws2 = load_workbook(io.BytesIO(admin_tpl.content))["Thu mua"]
    assert ws2.cell(row=5, column=1).value == "Đơn vị"      # mẫu chuyên viên vẫn có cột này
    bad = _rows_to_xlsx("Thu mua", _PURCHASE_HEAD,
                         [(OTHER, today, 99, 99, 900, 900, None, "SVR 3L", 99, 9.9, "VND")])
    prev2 = client.post("/api/member/import/preview?kind=purchase", headers=mh,
                        files={"file": ("f.xlsx", bad)}).json()
    assert prev2["rows"][0]["_errors"], "phải báo lỗi đơn vị ngoài quyền"

    # Kẻ xấu xoá cờ lỗi ở client rồi ép ghi → server vẫn phải chặn.
    hacked = [dict(r, _errors=[]) for r in prev2["rows"]]
    res = client.post("/api/member/import/commit", headers=mh,
                      json={"kind": "purchase", "rows": hacked}).json()
    assert res["saved"] == 0 and res["skipped"] == 1

    # Dọn SẠCH dữ liệu test (kể cả bản ghi đã ghi vào bảng số liệu) — không để rác lại DB.
    client.delete(f"/api/users/{MEMBER}", headers=h)
    for u in (UNIT, OTHER):
        client.delete(f"/api/member-units/{u}", headers=h)
    with session_scope() as db:
        db.execute(text("DELETE FROM unit_daily_report WHERE company IN (:a, :b)"),
                   {"a": UNIT, "b": OTHER})
        db.execute(text("DELETE FROM fact_price WHERE source = 'vrg' AND grade IN (:a, :b)"),
                   {"a": UNIT, "b": OTHER})


@_excel_on
def test_template_download_then_import_roundtrip() -> None:
    """Tải mẫu → điền vào ĐÚNG file mẫu → xem trước → ghi → đọc lại: số phải khớp.

    Khoá lại 2 lỗi từng gặp: (1) thêm cột vào mẫu nhưng bộ ghi bỏ qua (đơn giá mủ chén từng bị
    gán SAU lệnh upsert nên không lưu); (2) mẫu và bộ đọc lệch cột.
    """
    from datetime import timedelta
    h = _admin()
    unit = "_zz_xl_rt"
    client.post("/api/member-units", json={"name": unit}, headers=h)
    d = date.today() - timedelta(days=2)
    dmy, iso = d.strftime("%d/%m/%Y"), d.isoformat()
    later = (d + timedelta(days=20)).strftime("%d/%m/%Y")   # lịch giao — phải SAU ngày bắt đầu

    cases = {
        # Thu mua thành phẩm theo CHỦNG LOẠI: 2 dòng cho cùng (đơn vị, ngày) — số mủ nguyên liệu
        # chỉ điền ở dòng đầu, dòng sau để trống (đúng cách hướng dẫn trong file mẫu).
        # Không còn cột "Đơn giá mủ chén tính theo": mủ chén LUÔN theo độ DRC (chốt 17/08/2026).
        # Thứ tự cột: đơn vị · ngày · SL mủ nước · SL mủ chén · SL mủ dây
        #             · đơn giá nước · đơn giá chén · đơn giá dây · chủng loại TP · SL TP · giá TP · tiền TP
        "purchase": [(unit, dmy, 120.0, 50.0, 15.0, 540, 510, 260,
                      "SVR CV 50", 30, 41.2, "VND"),
                     (unit, dmy, None, None, None, None, None, None,
                      "SVR 3L", 20, 1800, "USD")],
        # Cột "Nguồn mủ" tách mủ thu mua / mủ khai thác thành 2 bảng lưu riêng.
        # Cột "Số HĐ/PL" = số hợp đồng / phụ lục của dòng bán (và của dòng tồn kho đã ký HĐ).
        "sales": [(unit, dmy, "HĐ-01/2026", "Mủ thu mua", "Dài hạn", "XK / UTXK", "SVR CV 50",
                   25, 1800, "USD", dmy, dmy),
                  (unit, dmy, "PL-02/2026", "Mủ khai thác", "Chuyến", "Nội tiêu", "SVR 3L",
                   10, 1700, "USD", dmy, dmy)],
        # Nhóm "Đã ký HĐ": cột Ngày = ngày BẮT ĐẦU tồn kho, thêm cột Lịch giao + Ngày giao.
        "stock": [(unit, dmy, "Chế biến chưa nhập kho", "SVR CV 50", None, 33,
                   None, None, None, None, None),
                  (unit, dmy, "Đã nhập kho", "SVR 3L", None, 66, None, None, None, None, None),
                  (unit, dmy, "Đã ký HĐ", "RSS 3", "HĐ-03/2026", 12, 1750, "USD",
                   later, None, 21)],
    }
    for kind, rows in cases.items():
        tpl = client.get(f"/api/unit-daily/import/template?kind={kind}", headers=h)
        assert tpl.status_code == 200
        wb = load_workbook(io.BytesIO(tpl.content))
        ws = wb.active
        for i, r in enumerate(rows, start=7):
            for j, v in enumerate(r, start=1):
                if v is not None:
                    ws.cell(row=i, column=j, value=v)
        buf = io.BytesIO()
        wb.save(buf)
        prev = client.post(f"/api/unit-daily/import/preview?kind={kind}", headers=h,
                           files={"file": ("f.xlsx", buf.getvalue())}).json()
        assert prev["summary"]["error"] == 0, (kind, prev["rows"])
        assert client.post("/api/unit-daily/import/commit", headers=h,
                           json={"kind": kind, "rows": prev["rows"]}).json()["saved"] >= 1

    pur = client.get(f"/api/unit-daily/day?kind=purchase&as_of={iso}",
                     headers=h).json()["entries"][unit]["fields"]
    assert pur["latex_wet"] == 120.0 and pur["coagulum"] == 50.0
    assert pur["lace"] == 15.0
    with session_scope() as db:
        lace_px = db.execute(text(
            "SELECT price FROM fact_price WHERE grade = :g AND price_type = 'purchase_lace' "
            "AND as_of = CAST(:d AS date)"), {"g": unit, "d": iso}).scalar()
    assert lace_px is not None and float(lace_px) == 260
    # Nhiều dòng thành phẩm của cùng 1 ngày phải gom vào MẢNG, không đè lẫn nhau.
    assert [(r["grade"], r["qty"], r["ccy"]) for r in pur["finished"]] == [
        ("SVR CV 50", 30, "VND"), ("SVR 3L", 20, "USD")]

    con = client.get(f"/api/unit-daily/day?kind=consumption&as_of={iso}",
                     headers=h).json()["entries"][unit]["fields"]
    assert con["sales_ccy"] == "USD" and con["sales"][0]["invoice_date"] == iso
    assert con["sales"][0]["warehouse_date"] == iso
    # Số HĐ/PL đi theo TỪNG DÒNG (cả bảng tiêu thụ lẫn khối tồn kho đã ký HĐ).
    assert con["sales"][0]["code"] == "HĐ-01/2026"
    # Loại tiền phải xuống TỪNG DÒNG, không thì web tính lại doanh thu ra số khác.
    assert con["sales"][0]["ccy"] == "USD"
    # Dòng "Mủ khai thác" vào bảng riêng, không lẫn với mủ thu mua.
    assert len(con["sales"]) == 1 and len(con["sales_own"]) == 1
    own = con["sales_own"][0]
    assert own["qty"] == 10 and own["contract"] == "spot" and own["channel"] == "domestic"
    assert own["code"] == "PL-02/2026"
    assert con["stock_ccy"] == "USD"
    assert con["stock_not_warehoused"][0]["qty"] == 33
    assert con["stock_warehoused"][0]["qty"] == 66
    assert con["stock_material"] == 21
    # Nhóm "Đã ký HĐ" nhập từ Excel thành HỢP ĐỒNG có vòng đời (bảng CŨ unit_stock_contract) —
    # báo cáo ngày tự hiển thị lại, đánh dấu "legacy": True (phân biệt nguồn sales_contract mới).
    assert con["stock_signed_undelivered"]["qty"] == 12
    assert con["stock_signed_undelivered"]["items"][0]["code"] == "HĐ-03/2026"
    assert con["stock_signed_undelivered"]["items"][0]["legacy"] is True
    ctr = client.get("/api/unit-daily/stock-contracts", headers=h).json()["contracts"]
    mine = [c for c in ctr if c["company"] == unit]
    assert len(mine) == 1 and mine[0]["start_date"] == iso and mine[0]["delivered_date"] is None

    with session_scope() as db:
        db.execute(text("DELETE FROM unit_stock_contract WHERE company = :u"), {"u": unit})
        db.execute(text("DELETE FROM unit_daily_report WHERE company = :u"), {"u": unit})
        db.execute(text("DELETE FROM fact_price WHERE source='vrg' AND grade = :u"), {"u": unit})
    client.delete(f"/api/member-units/{unit}", headers=h)


@_excel_on
def test_sales_import_old_file_without_source_column() -> None:
    """File mẫu CŨ (chưa có cột 'Nguồn mủ' và 'Số HĐ/PL') vẫn nhập được — mọi dòng vào bảng mủ thu mua.

    Bộ đọc dò cột theo TIÊU ĐỀ nên thiếu cột không phải là lỗi; giữ đúng cách hiểu trước đây
    để đơn vị đang dùng file cũ không bị gãy.
    """
    from datetime import timedelta
    h = _admin()
    unit = "_zz_xl_old"
    client.post("/api/member-units", json={"name": unit}, headers=h)
    d = date.today() - timedelta(days=3)
    dmy, iso = d.strftime("%d/%m/%Y"), d.isoformat()

    tpl = client.get("/api/unit-daily/import/template?kind=sales", headers=h)
    wb = load_workbook(io.BytesIO(tpl.content))
    ws = wb.active
    # Dựng lại đúng file mẫu đời trước: bỏ các cột thêm sau này (dò lại vị trí theo TIÊU ĐỀ
    # sau mỗi lần xoá, vì xoá cột làm các cột sau dịch chỗ).
    for title in ("Nguồn mủ", "Số HĐ/PL"):
        col = next(i for i in range(1, ws.max_column + 1)
                   if str(ws.cell(row=5, column=i).value or "").strip() == title)
        ws.delete_cols(col)
    for j, v in enumerate((unit, dmy, "Dài hạn", "XK / UTXK", "SVR CV 50", 25, 40, "VND"), start=1):
        ws.cell(row=7, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)

    prev = client.post("/api/unit-daily/import/preview?kind=sales", headers=h,
                       files={"file": ("f.xlsx", buf.getvalue())}).json()
    assert prev["summary"]["error"] == 0, prev["rows"]
    assert client.post("/api/unit-daily/import/commit", headers=h,
                       json={"kind": "sales", "rows": prev["rows"]}).json()["saved"] == 1

    con = client.get(f"/api/unit-daily/day?kind=consumption&as_of={iso}",
                     headers=h).json()["entries"][unit]["fields"]
    assert len(con["sales"]) == 1 and con["sales"][0]["qty"] == 25
    assert con["sales"][0]["code"] is None                # file cũ không có cột mã HĐ → để trống
    assert con["sales_own"] == []
    assert con["revenue"] == 25 * 40 * 1_000_000          # 25 tấn × 40 triệu đ/tấn

    with session_scope() as db:
        db.execute(text("DELETE FROM unit_daily_report WHERE company = :u"), {"u": unit})
    client.delete(f"/api/member-units/{unit}", headers=h)


@_excel_on
def test_sales_import_old_header_ma_hd_pl_alias() -> None:
    """File mẫu cũ còn tiêu đề cột 'Mã HĐ/PL' (trước khi đổi nhãn thành 'Số HĐ/PL') vẫn nhập được.

    Đổi nhãn hiển thị không được làm gãy file mẫu người dùng đang giữ — bộ đọc dò cột theo
    tiêu đề nên phải nhận cả tiêu đề cũ lẫn tiêu đề mới cho cùng 1 cột `code`.
    """
    from datetime import timedelta
    h = _admin()
    unit = "_zz_xl_alias"
    client.post("/api/member-units", json={"name": unit}, headers=h)
    d = date.today() - timedelta(days=5)
    dmy, iso = d.strftime("%d/%m/%Y"), d.isoformat()

    tpl = client.get("/api/unit-daily/import/template?kind=sales", headers=h)
    wb = load_workbook(io.BytesIO(tpl.content))
    ws = wb.active
    col = next(i for i in range(1, ws.max_column + 1)
               if str(ws.cell(row=5, column=i).value or "").strip() == "Số HĐ/PL")
    ws.cell(row=5, column=col, value="Mã HĐ/PL")     # giả lập tiêu đề CŨ trong file người dùng
    row = (unit, dmy, "HĐ-CU/2026", "Mủ thu mua", "Dài hạn", "XK / UTXK", "SVR CV 50",
           25, 40, "VND", dmy, dmy)
    for j, v in enumerate(row, start=1):
        ws.cell(row=7, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)

    prev = client.post("/api/unit-daily/import/preview?kind=sales", headers=h,
                       files={"file": ("f.xlsx", buf.getvalue())}).json()
    assert prev["summary"]["error"] == 0, prev["rows"]
    assert client.post("/api/unit-daily/import/commit", headers=h,
                       json={"kind": "sales", "rows": prev["rows"]}).json()["saved"] == 1

    con = client.get(f"/api/unit-daily/day?kind=consumption&as_of={iso}",
                     headers=h).json()["entries"][unit]["fields"]
    assert con["sales"][0]["code"] == "HĐ-CU/2026"     # cột dò được dù tiêu đề vẫn là bản CŨ

    with session_scope() as db:
        db.execute(text("DELETE FROM unit_daily_report WHERE company = :u"), {"u": unit})
    client.delete(f"/api/member-units/{unit}", headers=h)


@_excel_on
def test_sales_import_giu_loai_tien_theo_TUNG_DONG() -> None:
    """Cột 'Giá bán bằng' là cột CỦA TỪNG DÒNG — một ngày có thể vừa bán USD vừa bán VNĐ.

    Lỗi đã gặp: bộ ghi lấy loại tiền của dòng ĐẦU rồi áp cho mọi dòng, nên dòng VNĐ bị biến
    thành USD và doanh thu tính sai theo loại tiền không phải của nó.
    """
    from datetime import timedelta
    h = _admin()
    unit = "_zz_xl_ccy"
    client.post("/api/member-units", json={"name": unit}, headers=h)
    d = date.today() - timedelta(days=4)
    dmy, iso = d.strftime("%d/%m/%Y"), d.isoformat()

    tpl = client.get("/api/unit-daily/import/template?kind=sales", headers=h)
    wb = load_workbook(io.BytesIO(tpl.content))
    ws = wb.active
    rows = [
        (unit, dmy, "HĐ-USD", "Mủ thu mua", "Dài hạn", "XK / UTXK", "SVR CV 50",
         25, 1800, "USD", dmy, dmy),
        (unit, dmy, "HĐ-VND", "Mủ khai thác", "Chuyến", "Nội tiêu", "SVR 3L",
         10, 40, "VND", dmy, dmy),
    ]
    for i, r in enumerate(rows, start=7):
        for j, v in enumerate(r, start=1):
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)

    prev = client.post("/api/unit-daily/import/preview?kind=sales", headers=h,
                       files={"file": ("f.xlsx", buf.getvalue())}).json()
    assert prev["summary"]["error"] == 0, prev["rows"]
    assert client.post("/api/unit-daily/import/commit", headers=h,
                       json={"kind": "sales", "rows": prev["rows"]}).json()["saved"] == 1

    con = client.get(f"/api/unit-daily/day?kind=consumption&as_of={iso}",
                     headers=h).json()["entries"][unit]["fields"]
    assert con["sales"][0]["ccy"] == "USD"        # dòng mủ thu mua giữ USD
    assert con["sales_own"][0]["ccy"] == "VND"    # dòng mủ khai thác giữ VNĐ, KHÔNG bị ép sang USD
    # Doanh thu chỉ cộng dòng tính được: 10 tấn × 40 triệu = 400 triệu (dòng USD thiếu tỷ giá).
    assert con["revenue"] == 400_000_000

    with session_scope() as db:
        db.execute(text("DELETE FROM unit_daily_report WHERE company = :u"), {"u": unit})
    client.delete(f"/api/member-units/{unit}", headers=h)


def test_plan_template_puts_exploit_first_and_blank_cell_clears_it() -> None:
    """Biểu Kế hoạch năm: cột "Kế hoạch khai thác" đứng ĐẦU (ngay sau Đơn vị · Năm, trước "Kế hoạch
    thu mua"); điền vào đúng mẫu → ghi được → lần nhập sau để trống ô đó = XOÁ chỉ tiêu.

    Gọi thẳng tầng service (không qua endpoint) để test chạy được cả khi cờ nhập Excel đang tắt.
    """
    from app.services import unit_daily_excel_io as xio
    from app.services import unit_daily_repo

    h = _admin()
    unit = "_zz_xl_plan"
    year = date.today().year
    client.post("/api/member-units", json={"name": unit}, headers=h)

    ws = load_workbook(io.BytesIO(xio.build_template("plan")))["Kế hoạch năm"]
    head = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
    assert head[:4] == ["Đơn vị", "Năm", "Kế hoạch khai thác", "Kế hoạch thu mua"]

    def import_plan(exploit: float | None) -> int:
        cells = {"Đơn vị": unit, "Năm": year, "Kế hoạch khai thác": exploit, "Kế hoạch thu mua": 800}
        data = _rows_to_xlsx("Kế hoạch năm", head, [tuple(cells.get(t) for t in head)])
        prev = xio.parse_upload("plan", data)
        assert prev["summary"]["error"] == 0, prev["rows"]
        return xio.commit_rows("plan", prev["rows"], "admin")["saved"]

    try:
        assert import_plan(2500) == 1
        assert unit_daily_repo.year_plan(year, [unit])[unit]["plan_exploit_tonnes"] == 2500
        assert import_plan(None) == 1
        got = unit_daily_repo.year_plan(year, [unit])[unit]
        assert got["plan_exploit_tonnes"] is None and got["plan_tonnes"] == 800
    finally:
        with session_scope() as db:
            db.execute(text("DELETE FROM unit_purchase_plan WHERE company = :u"), {"u": unit})
        client.delete(f"/api/member-units/{unit}", headers=h)


@pytest.mark.skipif(EXCEL_IMPORT_ENABLED, reason="Nhập liệu bằng Excel đang bật")
def test_import_endpoints_dong_khi_tat_co() -> None:
    """Khi cờ tắt: mọi endpoint nhập Excel trả 503 — không tải được mẫu, không xem trước, không ghi.

    Frontend đã ẩn thanh nút, nhưng người biết đường dẫn vẫn có thể gọi thẳng API nên chặn ở server.
    """
    h = _admin()
    assert client.get("/api/unit-daily/import/template?kind=purchase",
                      headers=h).status_code == 503
    assert client.post("/api/unit-daily/import/preview?kind=purchase", headers=h,
                       files={"file": ("f.xlsx", b"khong-doc-den")}).status_code == 503
    assert client.post("/api/unit-daily/import/commit", headers=h,
                       json={"kind": "purchase", "rows": []}).status_code == 503


# ── Hồi quy review 24/09/2026 (gọi thẳng tầng xử lý — chạy được cả khi cờ nhập Excel đang tắt) ──
_PLAN_SHEET = "Kế hoạch năm"


def _plan_unit(name: str) -> dict[str, str]:
    h = _admin()
    client.post("/api/member-units", json={"name": name}, headers=h)
    return h


def _drop_plan_unit(name: str, h: dict[str, str]) -> None:
    with session_scope() as db:
        db.execute(text("DELETE FROM unit_purchase_plan WHERE company = :u"), {"u": name})
    client.delete(f"/api/member-units/{name}", headers=h)


def test_excel_non_finite_or_negative_plan_numbers_are_row_errors() -> None:
    """Ô chữ "nan"/"inf" KHÔNG được đọc thành số (float() nhận cả hai) → lỗi dòng, không ghi.

    Client gửi lại dòng lúc ghi nên JSON NaN vẫn có thể tới tầng ghi — phải chặn lần nữa ở đó.
    """
    from app.services import unit_daily_excel_io as xio
    from app.services import unit_daily_repo

    assert [xio._as_num(v) for v in ("nan", "NaN", "inf", "-Infinity", float("nan"))] == [None] * 5
    assert xio._as_num("1.234,5") == 1234.5

    unit, year = "_zz_xl_plan_nan", date.today().year
    h = _plan_unit(unit)
    try:
        head = ["Đơn vị", "Năm", "Kế hoạch khai thác", "Kế hoạch thu mua"]
        data = _rows_to_xlsx(_PLAN_SHEET, head,
                             [(unit, year, "nan", 800), (unit, year, 100, "inf"), (unit, year, -5, 800)])
        prev = xio.parse_upload("plan", data)
        assert prev["summary"] == {"total": 3, "ok": 0, "error": 3}
        errs = [r["_errors"] for r in prev["rows"]]
        assert "không phải số" in errs[0][0] and "không phải số" in errs[1][0]
        assert "không được âm" in errs[2][0]

        hacked = [{"_row": 7, "_errors": [], "company": unit, "year": year,
                   "plan_exploit_tonnes": float("nan"), "plan_tonnes": 800}]
        assert xio.commit_rows("plan", hacked, "admin")["saved"] == 0
        assert unit not in unit_daily_repo.year_plan(year, [unit])
    finally:
        _drop_plan_unit(unit, h)


def test_plan_import_old_template_keeps_columns_missing_from_file() -> None:
    """File mẫu CŨ (chưa có cột Kế hoạch khai thác / doanh thu): cột VẮNG MẶT giữ nguyên số đang lưu;
    ô TRỐNG của cột CÓ trong file vẫn là xoá (luật cũ)."""
    from app.services import unit_daily_excel_io as xio
    from app.services import unit_daily_repo

    unit, year = "_zz_xl_plan_old", date.today().year
    h = _plan_unit(unit)
    try:
        unit_daily_repo.set_year_plan(year, unit, 800, None, None, None, 300, 12.5, "admin",
                                      plan_exploit_tonnes=2500)
        old_head = ["Đơn vị", "Năm", "Kế hoạch thu mua", "Kế hoạch tiêu thụ (HĐ chuyến)",
                    "HĐ dài hạn đã ký", "HĐ dài hạn năm trước chuyển sang",
                    "HĐ chuyến năm trước chuyển sang"]
        data = _rows_to_xlsx(_PLAN_SHEET, old_head, [(unit, year, 900, None, 50, None, None)])
        prev = xio.parse_upload("plan", data)
        assert prev["summary"]["error"] == 0, prev["rows"]
        row = prev["rows"][0]
        assert "plan_exploit_tonnes" not in row and "plan_revenue_ty" not in row
        assert row["plan_sales_spot_tonnes"] is None           # cột có mặt, ô trống
        shown = [c["key"] for c in prev["columns"]]
        assert "plan_exploit_tonnes" not in shown and "plan_tonnes" in shown

        assert xio.commit_rows("plan", prev["rows"], "admin")["saved"] == 1
        got = unit_daily_repo.year_plan(year, [unit])[unit]
        assert got["plan_exploit_tonnes"] == 2500 and got["plan_revenue_ty"] == pytest.approx(12.5)
        assert got["plan_tonnes"] == 900 and got["signed_lt_tonnes"] == 50
        assert got["plan_sales_spot_tonnes"] is None
    finally:
        _drop_plan_unit(unit, h)


def test_period_excel_total_row_does_not_sum_prices_or_percents() -> None:
    """Dòng Tổng cộng của file báo cáo kỳ KHÔNG cộng giá BQ / % (40% + 70% ≠ 110%)."""
    from app.services import unit_period_excel

    def total_row(kind: str, rows: list[dict]) -> dict[str, object]:
        rep = {"kind": kind, "date_from": "2026-09-01", "date_to": "2026-09-07",
               "grades": [], "rows": rows}
        ws = load_workbook(io.BytesIO(unit_period_excel.build_period_xlsx(rep))).active
        head = [ws.cell(row=5, column=c).value for c in range(1, ws.max_column + 1)]
        last = next(r for r in range(8, ws.max_row + 1) if ws.cell(row=r, column=2).value == "Tổng cộng")
        return {t: ws.cell(row=last, column=i).value for i, t in enumerate(head, start=1)}

    pur = total_row("purchase", [
        {"company": "A", "latex_wet": 10, "price_lace_avg": 30000, "pct_plan": 40},
        {"company": "B", "latex_wet": 5, "price_lace_avg": 32000, "pct_plan": 70}])
    assert pur["Sản lượng thu mua mủ nước"] == 15
    assert pur["Giá thu mua mủ dây BQ"] is None and pur["% thực hiện kế hoạch"] is None

    con = total_row("consumption", [
        {"company": "A", "total_consumption": 10, "pct_plan_sales_spot": 40},
        {"company": "B", "total_consumption": 5, "pct_plan_sales_spot": 70}])
    assert con["Tổng tiêu thụ"] == 15
    assert con["% thực hiện KH tiêu thụ (HĐ chuyến)"] is None


def test_excel_import_respects_edit_window(monkeypatch) -> None:
    """Nhập Excel qua CÙNG hàng rào cửa sổ nhập liệu như form web (sổ tay đã hứa).

    Đơn vị (MEMBER window) và chuyên viên (EDITOR window): dòng có ngày quá hạn không ghi, báo
    lỗi từng dòng — cả ở bước xem trước lẫn lúc ghi. Admin không giới hạn. Endpoint đang tắt bằng
    cờ nên test mở tạm dependency cờ để đi đúng đường router thật.
    """
    from app.core import edit_window
    from app.core.feature_flags import require_excel_import
    from app.services import unit_daily_repo

    unit, mem, ed = "_zz_xl_win", "xl_win_mem", "xl_win_ed"
    today = edit_window.today()
    old = (today - timedelta(days=30)).isoformat()
    h = _admin()
    for u in (mem, ed):
        client.delete(f"/api/users/{u}", headers=h)
    client.post("/api/member-units", json={"name": unit}, headers=h)
    client.post("/api/users", json={"username": mem, "password": "pass123", "role": "member",
                                    "member_units": [unit]}, headers=h)
    client.post("/api/users", json={"username": ed, "password": "pass123", "role": "editor",
                                    "permissions": ["unit_daily"]}, headers=h)

    def login(u: str) -> dict[str, str]:
        tok = client.post("/api/auth/login", json={"username": u, "password": "pass123"}).json()
        return {"Authorization": f"Bearer {tok['access_token']}"}

    def rows(*days: str) -> list[dict]:
        return [{"_row": 7 + i, "_errors": [], "company": unit, "as_of": d, "latex_wet": 5.0}
                for i, d in enumerate(days)]

    def has(d: str) -> bool:
        return unit_daily_repo.has_entry("purchase", d, unit)

    # N = 1: hôm nay luôn còn hạn ở mọi giờ chạy; 30 ngày trước thì chắc chắn đã khoá.
    monkeypatch.setattr(edit_window, "member_window", lambda: 1)
    monkeypatch.setattr(edit_window, "editor_window", lambda: 1)
    app.dependency_overrides[require_excel_import] = lambda: None
    try:
        mh, eh = login(mem), login(ed)
        # Xem trước (mẫu 1 đơn vị, không có cột Đơn vị): dòng quá hạn bị đánh lỗi sẵn.
        data = _rows_to_xlsx("Thu mua", ["Ngày", "SL thu mua mủ nước"],
                             [(old, 5), (today.isoformat(), 6)])
        prev = client.post("/api/member/import/preview?kind=purchase", headers=mh,
                           files={"file": ("f.xlsx", data)}).json()
        assert prev["summary"] == {"total": 2, "ok": 1, "error": 1}
        assert "chỉ xem" in prev["rows"][0]["_errors"][0]

        # Ghi (client xoá cờ lỗi): đơn vị & chuyên viên đều bị chặn dòng quá hạn, dòng hôm nay vẫn ghi.
        for hdr, path in ((mh, "/api/member/import/commit"), (eh, "/api/unit-daily/import/commit")):
            res = client.post(path, headers=hdr,
                              json={"kind": "purchase", "rows": rows(old, today.isoformat())}).json()
            assert res["saved"] == 1 and res["skipped"] == 1, (path, res)
            assert "Dòng 7" in res["warnings"][0] and "chỉ xem" in res["warnings"][0]
            assert not has(old) and has(today.isoformat())

        # Admin: không giới hạn cửa sổ (như form web).
        res = client.post("/api/unit-daily/import/commit", headers=h,
                          json={"kind": "purchase", "rows": rows(old)}).json()
        assert res["saved"] == 1 and has(old)
    finally:
        app.dependency_overrides.pop(require_excel_import, None)
        for u in (mem, ed):
            client.delete(f"/api/users/{u}", headers=h)
        with session_scope() as db:
            db.execute(text("DELETE FROM unit_daily_report WHERE company = :u"), {"u": unit})
        client.delete(f"/api/member-units/{unit}", headers=h)
