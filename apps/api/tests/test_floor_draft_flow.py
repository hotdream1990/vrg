"""Test quy trình bản nháp giá sàn Nháp → Dự thảo → Tờ trình → Áp dụng: chuyển bước, khoá theo bước,
nội dung tờ trình mẫu mới, xuất Word, tải hình/PDF (Chromium giả) và phân quyền."""

from __future__ import annotations

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.core.db import db_healthy
from app.main import app
from app.services import doc_render, user_repo
from tests import floor_proposal_env as env

pytestmark = pytest.mark.skipif(not db_healthy(), reason="DB không sẵn sàng")

client = TestClient(app)
SVR10 = "SVR 10 / CSR 10"
URL = "/api/floor-proposal/drafts"


def _bearer(u: str, p: str) -> dict[str, str]:
    r = client.post("/api/auth/login", json={"username": u, "password": p})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture
def h():
    env.seed()
    user_repo.seed_admin()
    admin = _bearer("admin", "admin")
    client.delete("/api/users/fdf_exec", headers=admin)
    assert client.post("/api/users", headers=admin, json={
        "username": "fdf_exec", "password": "pass123", "full_name": "x", "role": "executive",
        "permissions": []}).status_code == 200
    yield {"admin": admin, "exec": _bearer("fdf_exec", "pass123")}
    client.delete("/api/users/fdf_exec", headers=admin)
    env.clear()


@pytest.fixture
def draft(h):
    r = client.post(URL, headers=h["admin"], json={"proposal": env.proposal(), "source": "manual"})
    assert r.status_code == 201, r.text
    d = r.json()
    yield d
    client.delete(f"{URL}/{d['id']}", headers=h["admin"])


def _move(h, d: dict, to: str, who: str = "admin"):
    return client.post(f"{URL}/{d['id']}/stage", headers=h[who], json={"to": to})


def _put(h, d: dict, **kw):
    body = {"title": d["title"], "proposal": d["proposal"], **kw}
    return client.put(f"{URL}/{d['id']}", headers=h["admin"], json=body)


def _bump(h, prop: dict) -> dict:
    r = client.post("/api/floor-proposal/apply", headers=h["admin"], json={
        "proposal": prop, "changes": [{"grades": [SVR10], "op": "step", "value": 1}]})
    return r.json()["proposal"]


def test_stage_flow_locks_each_part_to_its_step(h, draft) -> None:
    assert draft["stage"] == "nhap" and draft["sheet"] is None and draft["memo"] is None
    assert _move(h, draft, "to_trinh").status_code == 400                    # chỉ một nấc
    d = _move(h, draft, "du_thao").json()
    assert d["stage"] == "du_thao" and d["sheet"]["vcb_date"] == env.AS_OF and len(d["history"]) == 1
    # Dự thảo: số đã chốt, chỉ sửa tỷ giá chú thích.
    bad = _put(h, d, proposal=_bump(h, d["proposal"]))
    assert bad.status_code == 400 and "bước Nháp" in bad.json()["detail"]
    d = _put(h, d, sheet={"vcb_rate": 25790, "vcb_time": "8g30"}).json()
    assert d["sheet"]["vcb_rate"] == 25790
    assert _put(h, d, sheet={"vcb_rate": 999}).status_code == 400              # tỷ giá phi lý
    assert _put(h, d, memo={"so": "1/TTr-TTKD"}).status_code == 400            # chưa tới bước Tờ trình
    # Tờ trình: nội dung mặc định dựng từ số, sửa được; tỷ giá khoá.
    d = _move(h, d, "to_trinh").json()
    memo = d["memo"]
    assert memo["intro"].startswith("Ban TTKD kính trình") and memo["signers"]["approver_name"]
    assert memo["inventory"]["lead"].startswith("Tồn kho Tập đoàn")
    d = _put(h, d, memo={"so": "54/TTr-TTKD", "futures": [{"lead": "SGX – RSS3:", "text": "<b>tăng</b> nhẹ"}]}).json()
    assert d["memo"]["so"] == "54/TTr-TTKD" and d["memo"]["intro"] == memo["intro"]   # khoá không gửi giữ nguyên
    html = client.get(f"{URL}/{d['id']}/html", headers=h["admin"]).text
    assert "54/TTr-TTKD" in html and "SGX – RSS3:" in html and "&lt;b&gt;tăng" in html and "<b>tăng" not in html
    assert "TỜ TRÌNH" in html and "SÀN<br>FUTURES" in html
    assert _put(h, d, sheet={"vcb_rate": 25800}).status_code == 400
    # Áp dụng: có trong quy trình nhưng chưa mở.
    r = _move(h, d, "ap_dung")
    assert r.status_code == 400 and "chưa triển khai" in r.json()["detail"]
    # Lui về Nháp: nội dung tờ trình giữ lại, số sửa được.
    d = _move(h, _move(h, d, "du_thao").json(), "nhap").json()
    assert d["stage"] == "nhap" and d["memo"]["so"] == "54/TTr-TTKD" and len(d["history"]) == 4
    assert _put(h, d, proposal=_bump(h, d["proposal"])).status_code == 200


def test_conflict_permissions_and_exports(h, draft, monkeypatch) -> None:
    d = _move(h, draft, "du_thao").json()
    stale = "2000-01-01T00:00:00+00:00"           # mốc mở bản cũ hơn lần lưu gần nhất (độ phân giải 1 giây)
    r = client.post(f"{URL}/{d['id']}/stage", headers=h["admin"], json={"to": "to_trinh", "base_updated_at": stale})
    assert r.status_code == 409
    assert _move(h, d, "to_trinh", who="exec").status_code == 403             # Lãnh đạo chỉ xem
    d = _move(h, d, "to_trinh").json()
    assert _put(h, d, memo={"so": "54/TTr-TTKD"}).status_code == 200
    # Word: Lãnh đạo tải được; nội dung đúng mẫu.
    r = client.get(f"{URL}/{d['id']}/to-trinh.docx", headers=h["exec"])
    assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
    assert "To-trinh-gia-san-lan-" in r.headers["content-disposition"]
    word = Document(io.BytesIO(r.content))
    cells = [c.text for t in word.tables for row in t.rows for c in row.cells]
    text = "\n".join([p.text for p in word.paragraphs] + cells)
    assert "TỜ TRÌNH" in text and "Số: 54/TTr-TTKD" in text and "kính trình Tổng giám đốc./." in text
    # PNG/PDF: Chromium giả (máy test không có trình duyệt) — kiểm tra nối dây + báo lỗi gọn.
    monkeypatch.setattr(doc_render, "html_to_png", lambda html, sel, scale=2.0: b"PNG" + html.encode()[:0])
    monkeypatch.setattr(doc_render, "html_to_pdf", lambda html: b"%PDF")
    r = client.get(f"{URL}/{d['id']}/du-thao.png?inline=true", headers=h["admin"])
    assert r.status_code == 200 and r.headers["content-type"] == "image/png" and r.content == b"PNG"
    assert r.headers["content-disposition"].startswith("inline")
    assert client.get(f"{URL}/{d['id']}/to-trinh.pdf", headers=h["admin"]).content == b"%PDF"

    def boom(*_a, **_k):
        raise doc_render.RenderError("Không dựng được tệp (Chromium): thiếu")
    monkeypatch.setattr(doc_render, "html_to_pdf", boom)
    r = client.get(f"{URL}/{d['id']}/to-trinh.pdf", headers=h["admin"])
    assert r.status_code == 503 and "Chromium" in r.json()["detail"]
    lst = client.get(f"{URL}?page=1&page_size=5", headers=h["admin"]).json()
    assert next(i for i in lst["items"] if i["id"] == d["id"])["stage"] == "to_trinh"


def test_memo_ai_endpoint_returns_draft_without_saving(h, draft, monkeypatch) -> None:
    from app.services import llm, to_trinh_ai

    assert client.post(f"{URL}/{draft['id']}/memo/ai", headers=h["admin"], json={}).status_code == 400  # Nháp
    d = _move(h, _move(h, draft, "du_thao").json(), "to_trinh").json()
    monkeypatch.setattr(to_trinh_ai.ctx, "news_block", lambda doc: [])     # không gọi mạng
    monkeypatch.setattr(to_trinh_ai.llm, "complete", lambda s, u, max_tokens: (
        "### SAN\nSHANGHAI (SHFE) – RSS3: Giá đi ngang.\n### CUNGCAU\nCơ sở để điều chỉnh tăng giá sàn."))
    r = client.post(f"{URL}/{d['id']}/memo/ai", headers=h["admin"], json={"memo": {"so": "9/TTr-TTKD"}})
    assert r.status_code == 200, r.text
    m = r.json()["memo"]
    assert m["so"] == "9/TTr-TTKD" and m["outlook"][0]["text"].startswith("Cơ sở") and m["ai"]["by"] == "admin"
    assert client.get(f"{URL}/{d['id']}", headers=h["admin"]).json()["memo"]["so"] == ""   # chưa lưu

    def no_key(*_a, **_k):
        raise llm.LLMNotConfigured("Chưa có OpenAI API Key.")
    monkeypatch.setattr(to_trinh_ai.llm, "complete", no_key)
    r = client.post(f"{URL}/{d['id']}/memo/ai", headers=h["admin"], json={})
    assert r.status_code == 400 and "API Key" in r.json()["detail"]
    assert client.post(f"{URL}/{d['id']}/memo/ai", headers=h["exec"], json={}).status_code == 403


def test_return_from_to_trinh_straight_to_nhap_keeps_memo_and_flags_review(h, draft) -> None:
    """Lãnh đạo không duyệt tờ trình → trả về thẳng Nháp kèm lý do, sửa số, đi lại tới Tờ trình."""
    d = _move(h, _move(h, draft, "du_thao").json(), "to_trinh").json()
    d = _put(h, d, memo={"so": "54/TTr-TTKD"}).json()
    sig_at_memo = d["memo"]["sig"]
    assert sig_at_memo == d["sig"]
    r = client.post(f"{URL}/{d['id']}/stage", headers=h["admin"],
                    json={"to": "nhap", "note": "TGĐ chưa duyệt mức tăng, đề nghị giảm bớt"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["stage"] == "nhap" and d["history"][-1]["from"] == "to_trinh"
    assert d["history"][-1]["note"] == "TGĐ chưa duyệt mức tăng, đề nghị giảm bớt"
    d = _put(h, d, proposal=_bump(h, d["proposal"])).json()          # sửa số ở Nháp
    d = _move(h, _move(h, d, "du_thao").json(), "to_trinh").json()
    assert d["memo"]["so"] == "54/TTr-TTKD"                            # nội dung tờ trình giữ nguyên
    assert d["memo"]["sig"] == sig_at_memo != d["sig"]                 # → giao diện nhắc soát lại
    d = _put(h, d, memo={"sig": d["sig"]}).json()                      # "Đã soát, khớp số mới"
    assert d["memo"]["sig"] == d["sig"]
