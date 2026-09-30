"""Test file Excel chỉ số theo ngày: nhãn số ngày đúng nghĩa, màu hôm nay khác màu thiếu số, cột suất
tiêu hao chỉ hiện khi nhà máy có đủ chỉ số (thuần hàm)."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from app.services import scada_meters_excel as xls
from tests.test_scada_daily_meters import D1, D2, F4, at, hourly, raw, report

HEAD, FIRST = 5, 7  # dòng tiêu đề nhóm · dòng số liệu đầu tiên (xem build_xlsx)
RATIO_LABELS = ("kWh / bành", "m³ / bành")


def _ws(rep: dict):
    return load_workbook(io.BytesIO(xls.build_xlsx(rep))).active


def _header(ws) -> list:
    return [c.value for c in ws[HEAD]]


def _rep(factory: dict = F4) -> dict:
    """Ngày 1 đủ số trừ nước thiếu mốc 00:00 (partial); ngày 2 = hôm nay (in_progress)."""
    rows = ([raw(at(D1, 0), 1000, None, 100)] + hourly(D1, 1000, 50, 100)[1:]
            + hourly(D2, 1240, 74, 148, hours=16))
    return report(rows, D1, D2, now=at(D2, 15, 30), factory=factory)


def test_ratio_columns_follow_factory_metrics() -> None:
    assert [h for h in _header(_ws(_rep())) if h in RATIO_LABELS] == list(RATIO_LABELS)
    no_water = _header(_ws(_rep({**F4, "water_tag": None})))
    assert [h for h in no_water if h in RATIO_LABELS] == ["kWh / bành"]
    ws = _ws(_rep({**F4, "bales_tag": ""}))
    assert not any(h in RATIO_LABELS for h in _header(ws))
    assert ws.max_column == 1 + 4 * 2  # Ngày + 2 chỉ số × 4 cột, không có cột suất tiêu hao


def test_footer_day_count_labels() -> None:
    ws = _ws(_rep())
    total_row, avg_row = FIRST + 2, FIRST + 3
    assert ws.cell(total_row, 1).value == "Tổng cộng" and ws.cell(avg_row, 1).value.startswith("Bình")
    # Điện: cả 2 ngày có số (tính cả hôm nay) nhưng chỉ ngày 1 đủ số. Nước: ngày 1 thiếu mốc.
    assert ws.cell(total_row, 5).value == "2 ngày có số"
    assert ws.cell(avg_row, 5).value == "1 ngày đủ số"
    assert ws.cell(avg_row, 9).value == "0 ngày đủ số" and ws.cell(avg_row, 8).value is None
    assert ws.cell(total_row, 14).value == 5.0  # kWh/bành cả kỳ — chỉ ngày 1


def test_today_fill_differs_from_missing_data_fill() -> None:
    ws = _ws(_rep())
    today_energy, day1_water, day1_energy = ws.cell(FIRST + 1, 2), ws.cell(FIRST, 6), ws.cell(FIRST, 2)
    assert today_energy.fill.fgColor.rgb.endswith("E3EEF9")
    assert day1_water.fill.fgColor.rgb.endswith("FFF4D6")
    assert day1_energy.fill.fill_type is None
    assert "ô tô xanh = hôm nay" in ws["A3"].value and "ô tô vàng = thiếu số" in ws["A3"].value.lower()
    assert ws.cell(FIRST, 9).value.startswith("Thiếu mốc 00:00")
    assert ws.cell(FIRST + 1, 5).value.startswith("Đang trong ngày")
